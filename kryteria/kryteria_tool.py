#!/usr/bin/env python3
"""kryteria_tool.py - symulacja punktacji ofert (cena i kryteria pozacenowe) dla przetargow MiD.

Specyfikacja (JSON): kryteria z SWZ + oferty (znane z otwarcia albo hipotetyczne). Wzor: `wzor`.

Komendy:
  wzor [--out kryteria.json]                    przykladowa specyfikacja do edycji
  swz TXT_DIR                                   fragmenty SWZ o kryteriach: wagi, wzory, progi, dokumenty,
                                                "nie podlega uzupelnieniu" (TXT_DIR = wynik swz_tool.py teksty)
  otwarcie PLIK.txt [--out oferty.json]         oferty z informacji z otwarcia (wykonawca, cena, kwota na sfinansowanie)
  ocena SPEC.json [--konkurencja max|min|jak_my]  punkty i ranking (wartosci null u konkurencji = zalozenie)
  prog SPEC.json [--konkurencja ...]            maksymalna cena, przy ktorej oferta "my" wygrywa, i ceny remisu z kazda oferta
  scenariusze SPEC.json [--konkurencja ...]     nasze punkty pozacenowe (kazdy poziom) x maksymalna cena wygrywajaca
        [--budzet B --udzialy 0.6,0.7,0.8,0.9]  przed otwarciem: hipotetyczny najtanszy konkurent = udzial x budzet
  xlsx SPEC.json --out Symulacja_punktacji.xlsx arkusz z formulami (ceny i deklaracje do zmiany recznie)

Typy kryteriow: cena (Cmin/C x W), cena_liniowa ((Cmax-C)/(Cmax-Cmin) x W), min (Xmin/X x W, np. termin),
max (X/Xmax x W, opcjonalnie "limit"), liniowy (od/do -> pkt_od/pkt_do), progi ([[wartosc, pkt], ...] rosnaco,
"ponizej": pkt ponizej pierwszego progu), tak_nie (W albo 0), punkty (wartosc = punkty).
Pole "zaokraglenie" (domyslnie 2) - zaokraglenie punktow kryterium jak u zamawiajacego.
"""
import argparse, json, re, statistics, sys
from pathlib import Path

PRZYKLAD = {
    "postepowanie": "nr 34 ZDW Katowice, zadanie 1 (przyklad)",
    "zrodlo": "SWZ pkt 17 s. 9-10",
    "remis": "cena",
    "kryteria": [
        {"id": "C", "nazwa": "Cena oferty brutto", "typ": "cena", "waga": 60},
        {"id": "B", "nazwa": "Doswiadczenie projektanta mostowego - liczba dokumentacji (zal. 1A)", "typ": "progi", "waga": 40,
         "progi": [[2, 0], [3, 10], [4, 20], [5, 30], [6, 40]], "ponizej": 0,
         "dokument": "zal. 1A", "uzupelnienie": False}
    ],
    "oferty": [
        {"wykonawca": "MiD", "my": True, "C": 350000, "B": 6},
        {"wykonawca": "Konkurent A", "C": 307377, "B": None},
        {"wykonawca": "Konkurent B", "C": 377763.75, "B": 6}
    ]
}


# ----------------------------------------------------------------------------- pomocnicze
def zl(v, waluta=True):
    if v is None:
        return "—"
    s = f"{v:,.2f}".replace(",", " ").replace(".", ",")
    return s + (" zł" if waluta else "")


def pkt(v):
    return "—" if v is None else f"{v:.2f}".replace(".", ",")


def liczba(s):
    s = str(s).replace(" ", " ").strip()
    s = re.sub(r"[^\d,.\- ]", "", s).replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def wczytaj(p):
    spec = json.loads(Path(p).read_text(encoding="utf-8"))
    for k in spec["kryteria"]:
        k.setdefault("zaokraglenie", 2)
    if not any(o.get("my") for o in spec["oferty"]):
        print("Uwaga: zadna oferta nie ma \"my\": true - prog i scenariusze potrzebuja oferty MiD.", file=sys.stderr)
    return spec


def kryt_ceny(spec):
    k = [k for k in spec["kryteria"] if k["typ"] in ("cena", "cena_liniowa")]
    if len(k) != 1:
        sys.exit("Specyfikacja musi miec dokladnie jedno kryterium typu cena albo cena_liniowa.")
    return k[0]


def najlepsza(k, znane):
    """Wartosc dajaca maksimum punktow w kryterium (do zalozenia 'max' dla nieznanej deklaracji)."""
    t = k["typ"]
    if "najlepsza" in k:
        return k["najlepsza"]
    if t == "progi":
        return max(p for p, _ in k["progi"])
    if t == "liniowy":
        return k["do"] if k["pkt_do"] >= k["pkt_od"] else k["od"]
    if t == "tak_nie":
        return 1
    if t == "punkty":
        return k["waga"]
    if t == "min":
        return min(znane) if znane else k.get("dolna", 0)
    if t == "max":
        return k.get("limit") or (max(znane) if znane else None)
    return None


def najgorsza(k, znane):
    t = k["typ"]
    if t == "progi":
        return min(p for p, _ in k["progi"]) - 1
    if t == "liniowy":
        return k["od"] if k["pkt_do"] >= k["pkt_od"] else k["do"]
    if t in ("tak_nie", "punkty"):
        return 0
    if t == "min":
        return max(znane) if znane else k.get("gorna")
    if t == "max":
        return min(znane) if znane else 0
    return None


def uzupelnij(spec, konkurencja="max"):
    """Zwraca oferty z wartosciami; null u konkurencji zastepuje zalozeniem. Lista zalozen do raportu."""
    oferty = [dict(o) for o in spec["oferty"]]
    my = next((o for o in oferty if o.get("my")), None)
    zal = []
    for k in spec["kryteria"]:
        if k["typ"] in ("cena", "cena_liniowa"):
            continue
        znane = [o[k["id"]] for o in oferty if o.get(k["id"]) is not None]
        for o in oferty:
            if o.get(k["id"]) is None:
                if o.get("my"):
                    sys.exit(f"Brak wartości kryterium {k['id']} w ofercie MiD.")
                v = {"max": najlepsza(k, znane), "min": najgorsza(k, znane), "jak_my": my.get(k["id"]) if my else None}[konkurencja]
                o[k["id"]] = v
                zal.append(f"{o['wykonawca']}: {k['id']} = {v} (założenie '{konkurencja}')")
    return oferty, zal


def punkty_kryterium(k, x, wszystkie):
    """Punkty za wartosc x w kryterium k; wszystkie = wartosci wszystkich ofert (kryteria wzgledne)."""
    t, W = k["typ"], k["waga"]
    if x is None:
        return 0.0
    if t == "cena":
        p = W * min(wszystkie) / x
    elif t == "cena_liniowa":
        lo, hi = min(wszystkie), max(wszystkie)
        p = W if hi == lo else W * (hi - x) / (hi - lo)
    elif t == "min":
        dol = k.get("dolna")
        xs = [max(v, dol) if dol is not None else v for v in wszystkie]
        xx = max(x, dol) if dol is not None else x
        p = W * min(xs) / xx
    elif t == "max":
        lim = k.get("limit")
        xs = [min(v, lim) if lim else v for v in wszystkie]
        xx = min(x, lim) if lim else x
        p = W * xx / max(xs) if max(xs) else 0
    elif t == "liniowy":
        od, do, po, pd = k["od"], k["do"], k["pkt_od"], k["pkt_do"]
        p = po + (x - od) * (pd - po) / (do - od)
        p = max(min(p, max(po, pd)), min(po, pd))
    elif t == "progi":
        p = k.get("ponizej", 0)
        for prog, pp in sorted(k["progi"]):
            if x >= prog:
                p = pp
    elif t == "tak_nie":
        p = W if x else 0
    elif t == "punkty":
        p = min(float(x), W)
    else:
        sys.exit(f"Nieznany typ kryterium: {t}")
    return round(p, k.get("zaokraglenie", 2))


def ocen(spec, oferty):
    kc = kryt_ceny(spec)
    wyniki = []
    for o in oferty:
        r = {"wykonawca": o["wykonawca"], "my": bool(o.get("my")), "cena": o[kc["id"]], "pkt": {}}
        for k in spec["kryteria"]:
            wsz = [x[k["id"]] for x in oferty if x.get(k["id"]) is not None]
            r["pkt"][k["id"]] = punkty_kryterium(k, o.get(k["id"]), wsz)
        r["suma"] = round(sum(r["pkt"].values()), 2)
        wyniki.append(r)
    # ranking: suma malejaco, remis -> nizsza cena (typowe w SWZ; "remis": "cena")
    # przy pelnym remisie (te same punkty i cena) MiD liczymy za konkurentem - ostroznie (losowanie / oferty dodatkowe)
    wyniki.sort(key=lambda r: (-r["suma"], r["cena"] if spec.get("remis", "cena") == "cena" else 0, 1 if r["my"] else 0))
    for i, r in enumerate(wyniki, 1):
        r["miejsce"] = i
    return wyniki


def czy_wygrywa(spec, oferty, cena):
    kc = kryt_ceny(spec)
    of = [dict(o) for o in oferty]
    for o in of:
        if o.get("my"):
            o[kc["id"]] = cena
    w = ocen(spec, of)
    return w[0]["my"], w


def max_cena_wygrywajaca(spec, oferty, gora=None):
    """Najwyzsza cena (z dokladnoscia do 1 gr), przy ktorej 'my' jest na 1. miejscu. Przewaga 'my' maleje z cena -> bisekcja."""
    kc = kryt_ceny(spec)
    ceny = [o[kc["id"]] for o in oferty if not o.get("my")]
    lo, hi = 0.01, gora or max(ceny + [1.0]) * 20
    if not czy_wygrywa(spec, oferty, lo)[0]:
        return None
    if czy_wygrywa(spec, oferty, hi)[0]:
        return hi
    for _ in range(80):
        mid = (lo + hi) / 2
        if czy_wygrywa(spec, oferty, mid)[0]:
            lo = mid
        else:
            hi = mid
        if hi - lo < 0.005:
            break
    return _do_grosza(lo, lambda c: czy_wygrywa(spec, oferty, c)[0])


def _do_grosza(lo, warunek):
    """Zaokraglenie w dol do pelnego grosza z kontrola warunku (bisekcja konczy sie tuz pod progiem)."""
    import math
    c = math.floor(lo * 100 + 1e-6) / 100
    for _ in range(3):
        if c <= 0.01 or warunek(c):
            break
        c = round(c - 0.01, 2)
    return c


def cena_wyprzedzenia(spec, oferty, wykonawca):
    """Najwyzsza cena MiD, przy ktorej MiD jest w rankingu przed dana oferta (punkty zaokraglone jak u zamawiajacego,
    remis rozstrzyga nizsza cena). Pozostale oferty wplywaja na Cmin."""
    kc = kryt_ceny(spec)

    def przed(c):
        of = [dict(o) for o in oferty]
        for o in of:
            if o.get("my"):
                o[kc["id"]] = c
        kolej = [r["wykonawca"] if not r["my"] else None for r in ocen(spec, of)]
        return kolej.index(None) < kolej.index(wykonawca)

    lo, hi = 0.01, max(o[kc["id"]] for o in oferty) * 20
    if not przed(lo):
        return None
    if przed(hi):
        return hi
    for _ in range(80):
        mid = (lo + hi) / 2
        if przed(mid):
            lo = mid
        else:
            hi = mid
        if hi - lo < 0.005:
            break
    return _do_grosza(lo, przed)


def uwagi_cenowe(spec, oferty, budzet=None):
    """Art. 224 ust. 2 pkt 1 Pzp: cena nizsza o >=30% od sredniej wszystkich ofert (albo od wartosci zamowienia z VAT)
    = obowiazkowe wezwanie do wyjasnien RNC. Oferty powyzej kwoty na sfinansowanie - ryzyko art. 255 pkt 3."""
    kc = kryt_ceny(spec)
    ceny = [o[kc["id"]] for o in oferty]
    sr = sum(ceny) / len(ceny)
    out = [f"Średnia cen wszystkich ofert: {zl(sr)}; próg wyjaśnień RNC (−30%, art. 224 ust. 2 pkt 1 Pzp): {zl(0.7 * sr)}" +
           (f"; od kwoty na sfinansowanie (zastępczo za wartość zamówienia z VAT): {zl(0.7 * budzet)}" if budzet else "") + "."]
    for o in sorted(oferty, key=lambda o: o[kc["id"]]):
        c = o[kc["id"]]
        if c < 0.7 * sr or (budzet and c < 0.7 * budzet):
            out.append(f"- {o['wykonawca']}: {zl(c)} = {proc(c / sr)} średniej" + (f", {proc(c / budzet)} budżetu" if budzet else "") +
                       " → wezwanie do wyjaśnień rażąco niskiej ceny")
        if budzet and c > budzet:
            out.append(f"- {o['wykonawca']}: {zl(c)} powyżej kwoty na sfinansowanie (+{proc(c / budzet - 1)}) - zamawiający może odrzucić (art. 255 pkt 3), jeśli nie zwiększy kwoty")
    return "\n".join(out)


def proc(x):
    return f"{x * 100:.1f}%".replace(".", ",")


def opis_krytrium(k):
    t = k["typ"]
    if t == "cena":
        return f"Cmin/C × {k['waga']}"
    if t == "cena_liniowa":
        return f"(Cmax−C)/(Cmax−Cmin) × {k['waga']}"
    if t == "min":
        return f"Xmin/X × {k['waga']}" + (f" (wartości < {k['dolna']} liczone jak {k['dolna']})" if k.get("dolna") is not None else "")
    if t == "max":
        return f"X/Xmax × {k['waga']}" + (f" (limit {k['limit']})" if k.get("limit") else "")
    if t == "liniowy":
        return f"liniowo {k['od']}→{k['pkt_od']} pkt … {k['do']}→{k['pkt_do']} pkt"
    if t == "progi":
        return "progi: " + ", ".join(f"≥{p}: {pp}" for p, pp in sorted(k["progi"])) + f"; poniżej: {k.get('ponizej', 0)}"
    if t == "tak_nie":
        return f"tak = {k['waga']} pkt, nie = 0"
    return f"punkty wprost (max {k['waga']})"


def tabela_ocen(spec, wyniki):
    ids = [k["id"] for k in spec["kryteria"]]
    kc = kryt_ceny(spec)
    linie = ["| miejsce | wykonawca | cena brutto | " + " | ".join(f"{i} pkt" for i in ids) + " | suma |",
             "|---|---|---|" + "---|" * len(ids) + "---|"]
    for r in wyniki:
        nazwa = f"**{r['wykonawca']}**" if r["my"] else r["wykonawca"]
        linie.append(f"| {r['miejsce']} | {nazwa} | {zl(r['cena'])} | " + " | ".join(pkt(r["pkt"][i]) for i in ids) + f" | {pkt(r['suma'])} |")
    return "\n".join(linie)


def naglowek(spec, zal, konkurencja):
    out = [f"# Symulacja punktacji: {spec.get('postepowanie', '')}", ""]
    if spec.get("zrodlo"):
        out.append(f"Źródło kryteriów: {spec['zrodlo']}")
    out.append("")
    out.append("| kryterium | waga | sposób oceny | dokument | uzupełnienie |\n|---|---|---|---|---|")
    for k in spec["kryteria"]:
        uz = "" if "uzupelnienie" not in k else ("**NIE** (błąd = 0 pkt)" if not k["uzupelnienie"] else "tak")
        out.append(f"| {k['id']}: {k['nazwa']} | {k['waga']} | {opis_krytrium(k)} | {k.get('dokument', '')} | {uz} |")
    if zal:
        out.append(f"\nZałożenia dla nieznanych deklaracji konkurencji ({konkurencja}): " + "; ".join(zal[:12]) + (" …" if len(zal) > 12 else ""))
    return "\n".join(out)


# ----------------------------------------------------------------------------- komendy
def cmd_wzor(a):
    Path(a.out).write_text(json.dumps(PRZYKLAD, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Zapisano {a.out}. Uzupełnij kryteria wg SWZ i oferty (null = nieznana deklaracja konkurenta).")


def cmd_ocena(a):
    spec = wczytaj(a.spec)
    oferty, zal = uzupelnij(spec, a.konkurencja)
    w = ocen(spec, oferty)
    print(naglowek(spec, zal, a.konkurencja))
    print("\n## Ranking\n")
    print(tabela_ocen(spec, w))
    kc = kryt_ceny(spec)
    ceny = [o[kc["id"]] for o in oferty if not o.get("my")]
    if ceny:
        print(f"\nCeny konkurencji (liczba ofert: {len(ceny)}): min {zl(min(ceny))}, mediana {zl(statistics.median(ceny))}, max {zl(max(ceny))}.")
    print("\n## Cena a rażąco niska cena / budżet\n")
    print(uwagi_cenowe(spec, oferty, a.budzet))


def cmd_prog(a):
    spec = wczytaj(a.spec)
    oferty, zal = uzupelnij(spec, a.konkurencja)
    kc = kryt_ceny(spec)
    my = next(o for o in oferty if o.get("my"))
    print(naglowek(spec, zal, a.konkurencja))
    w = ocen(spec, oferty)
    print("\n## Ranking przy obecnej cenie MiD\n")
    print(tabela_ocen(spec, w))
    mx = max_cena_wygrywajaca(spec, oferty)
    ceny = [o[kc["id"]] for o in oferty if not o.get("my")]
    print("\n## Próg ceny MiD\n")
    if mx is None:
        print("MiD nie wygrywa przy żadnej cenie — konkurencja ma przewagę w kryteriach pozacenowych nie do odrobienia ceną.")
    else:
        med = statistics.median(ceny) if ceny else None
        print(f"Najwyższa cena brutto, przy której MiD jest na 1. miejscu: **{zl(mx)}**" +
              (f" ({proc(mx / min(ceny))} najniższej ceny konkurencji, {proc(mx / med)} mediany)" if ceny else "") + ".")
        print(f"Obecna cena MiD {zl(my[kc['id']])}: {'wygrywa' if my[kc['id']] <= mx else 'przegrywa'}" +
              (f", zapas {zl(mx - my[kc['id']])}" if my[kc['id']] <= mx else f", trzeba obniżyć o {zl(my[kc['id']] - mx)}") + ".")
        of = [dict(o) for o in oferty]
        for o in of:
            if o.get("my"):
                o[kc["id"]] = mx
        sr = sum(o[kc["id"]] for o in of) / len(of)
        if mx < 0.7 * sr or (a.budzet and mx < 0.7 * a.budzet):
            print(f"⚠ Przy cenie progowej MiD będzie wezwane do wyjaśnień RNC (cena < 70% średniej {zl(sr)}" +
                  (f" lub budżetu {zl(a.budzet)}" if a.budzet else "") + "). Przygotuj kalkulację (skill wycena-oferty).")
        if a.budzet and mx > a.budzet:
            print(f"Próg jest powyżej kwoty na sfinansowanie ({zl(a.budzet)}) - realnym limitem jest budżet.")
    print("\n| konkurent | cena | pkt pozacenowe | maks. cena MiD, by go wyprzedzić | MiD / konkurent |\n|---|---|---|---|---|")
    for r in sorted((r for r in w if not r["my"]), key=lambda r: r["cena"]):
        cr = cena_wyprzedzenia(spec, oferty, r["wykonawca"])
        poz = round(r["suma"] - r["pkt"][kc["id"]], 2)
        print(f"| {r['wykonawca']} | {zl(r['cena'])} | {pkt(poz)} | {zl(cr) if cr else 'nieosiągalna'} | " +
              (f"{proc(cr / r['cena'])} |" if cr else "— |"))
    print("\nRemis rozstrzyga " + ("niższa cena." if spec.get("remis", "cena") == "cena" else f"{spec.get('remis')}."))


def poziomy_pozacenowe(spec):
    """Kombinacje naszych deklaracji w kryteriach pozacenowych (dla progi/tak_nie wszystkie poziomy; inne - obecna wartosc)."""
    import itertools
    osie = []
    for k in spec["kryteria"]:
        if k["typ"] == "progi":
            osie.append([(k["id"], p) for p, _ in sorted(k["progi"])])
        elif k["typ"] == "tak_nie":
            osie.append([(k["id"], 0), (k["id"], 1)])
        elif k["typ"] == "liniowy":
            n = 4
            osie.append([(k["id"], k["od"] + (k["do"] - k["od"]) * i / n) for i in range(n + 1)])
    return [dict(c) for c in itertools.product(*osie)] if osie else [{}]


def cmd_scenariusze(a):
    spec = wczytaj(a.spec)
    kc = kryt_ceny(spec)
    if a.budzet:
        # przed otwarciem: jeden hipotetyczny najtanszy konkurent z cena = udzial x budzet
        udzialy = [float(x) for x in a.udzialy.split(",")]
        my = next((o for o in spec["oferty"] if o.get("my")), {"wykonawca": "MiD", "my": True})
        print(f"# Scenariusze przed otwarciem: {spec.get('postepowanie', '')}\n")
        print(f"Budżet (kwota na sfinansowanie): {zl(a.budzet)}. Najtańszy konkurent: cena = udział × budżet, "
              f"pozacenowe wg założenia '{a.konkurencja}'.\n")
        poz = poziomy_pozacenowe(spec)
        print("| pozacenowe MiD | " + " | ".join(f"konkurent {u:.0%} budżetu" for u in udzialy) + " |")
        print("|---|" + "---|" * len(udzialy))
        for p in poz:
            komorki = []
            for u in udzialy:
                kon = {"wykonawca": "konkurent", kc["id"]: u * a.budzet}
                kon.update({k["id"]: None for k in spec["kryteria"] if k is not kc})
                m = dict(my)
                m.update(p)
                m[kc["id"]] = a.budzet
                sp = dict(spec)
                sp["oferty"] = [m, kon]
                of, _ = uzupelnij(sp, a.konkurencja)
                mx = max_cena_wygrywajaca(sp, of)
                komorki.append(f"{zl(mx)} ({mx / a.budzet:.0%})" if mx else "nie wygrywa")
            print(f"| {', '.join(f'{k}={v:g}' for k, v in p.items()) or '—'} | " + " | ".join(komorki) + " |")
        print("\nW komórce: najwyższa cena MiD brutto, przy której MiD wygrywa (w nawiasie: udział w budżecie).")
        return
    oferty, zal = uzupelnij(spec, a.konkurencja)
    ceny = [o[kc["id"]] for o in oferty if not o.get("my")]
    med = statistics.median(ceny) if ceny else None
    print(naglowek(spec, zal, a.konkurencja))
    print("\n## Ile wolno zaoferować przy danym poziomie punktów pozacenowych\n")
    print("| pozacenowe MiD | pkt | najwyższa cena wygrywająca | % najniższej ceny konk. | % mediany |\n|---|---|---|---|---|")
    for p in poziomy_pozacenowe(spec):
        of = [dict(o) for o in oferty]
        for o in of:
            if o.get("my"):
                o.update(p)
        mx = max_cena_wygrywajaca(spec, of)
        myo = next(o for o in of if o.get("my"))
        pp = sum(punkty_kryterium(k, myo.get(k["id"]), [x[k["id"]] for x in of if x.get(k["id"]) is not None])
                 for k in spec["kryteria"] if k is not kc)
        opis = ", ".join(f"{k}={v:g}" for k, v in p.items()) or "—"
        print(f"| {opis} | {pkt(pp)} | {zl(mx) if mx else 'nie wygrywa'} | " +
              (f"{proc(mx / min(ceny))} | {proc(mx / med)} |" if mx and ceny else "— | — |"))
    if kc["typ"] == "cena":
        W = kc["waga"]
        print(f"\nReguła dla Cmin/C × {W}: przewaga ΔP pkt pozacenowych pozwala być droższym od konkurenta o ΔP/({W}−ΔP) "
              f"(10 pkt: +{proc(10 / (W - 10))}); strata ΔP pkt wymaga ceny niższej o ΔP/{W} (10 pkt: −{proc(10 / W)}). "
              "Zaokrąglenie punktów do 0,01 przesuwa progi o kilkadziesiąt złotych.")


# ----------------------------------------------------------------------------- SWZ i otwarcie
WZORCE_SWZ = [
    ("kryteria", r"kryteri\w* oceny|kryterium|kryteria"),
    ("waga", r"\bwag[aię]\b|\d{1,3} ?%|\d{1,3} ?pkt|punkt"),
    ("nie_uzupelnia", r"nie podleg\w* (?:uzupe|wyja|popraw|zmian)|nie bed\w* podleg\w* uzupe|nie zostanie uzupe"),
    ("zero", r"otrzyma\w* 0|0 ?pkt|zero punkt|nie otrzyma\w* punkt"),
    ("remis", r"taki sam bilans|jednakow\w* liczb\w* punkt|tak\w* sam\w* liczb\w* punkt|dogrywk|oferty dodatkow"),
    ("dokument", r"(?:za[łl][aą]cznik\w*|formularz\w*)[^.]{0,80}kryteri|kryteri[^.]{0,80}(?:za[łl][aą]cznik|formularz)"),
]


def cmd_swz(a):
    pliki = sorted(p for p in Path(a.txt).glob("*.txt") if not p.name.startswith("_"))
    if not pliki:
        sys.exit("Brak plików .txt - najpierw swz_tool.py teksty.")
    wyniki = {k: [] for k, _ in WZORCE_SWZ}
    strony_kryt = []
    for p in pliki:
        strona = 1
        linie = p.read_text(encoding="utf-8", errors="replace").splitlines()
        tekst_stron = {}
        for ln in linie:
            m = re.match(r"\[\[s\. (\d+)\]\]", ln)
            if m:
                strona = int(m.group(1))
                continue
            tekst_stron.setdefault(strona, []).append(ln)
        for s, ll in tekst_stron.items():
            t = " ".join(ll)
            tl = t.lower()
            if re.search(r"kryteri", tl) and re.search(r"\bwag|pkt|punkt", tl):
                strony_kryt.append((p.stem, s, len(re.findall(r"kryteri", tl))))
            for i, ln in enumerate(ll):
                l2 = " ".join(ll[max(0, i - 1):i + 2]).lower()
                for k, rx in WZORCE_SWZ[2:]:
                    if re.search(rx, ln.lower()) and (k != "zero" or "kryteri" in " ".join(ll).lower()):
                        wyniki[k].append((p.stem, s, " ".join(" ".join(ll[max(0, i - 1):i + 2]).split())[:260]))
    print("## Strony z opisem kryteriów (kryteria + wagi/punkty)\n")
    for f, s, n in sorted(strony_kryt, key=lambda x: -x[2])[:12]:
        print(f"- [{f}, s. {s}] ({n}× „kryteri”)")
    nazwy = {"nie_uzupelnia": "Nie podlega uzupełnieniu / wyjaśnieniom / zmianie", "zero": "0 pkt / brak punktów",
             "remis": "Rozstrzyganie remisu", "dokument": "Dokumenty do kryteriów"}
    for k, tytul in nazwy.items():
        print(f"\n## {tytul}\n")
        widziane = set()
        for f, s, t in wyniki[k]:
            if (f, s, t[:80]) in widziane:
                continue
            widziane.add((f, s, t[:80]))
            print(f"- [{f}, s. {s}] {t}")
        if not widziane:
            print("- nie znaleziono")
    print("\nPrzeczytaj wskazane strony (swz_tool.py strony) i przepisz kryteria do specyfikacji JSON (komenda wzor).")


RX_CENA = re.compile(r"(?<![\d.,])(\d{1,3}(?:[ \u00a0.]\d{3})+|\d{3,9}),(\d{2})(?![\d.,])")  # polski zapis: przecinek dziesietny
RX_CZESC = re.compile(r"^\s*(?:ZADANIE|CZ[ĘE][ŚS][ĆC]|Zadanie|Cz[ęe][śs][ćc])\s*(?:nr\s*)?(\d+)\b", re.I)
POMIN = re.compile(r"^(?:numer|cena|wykonawca|oferty|\(brutto\)|nazwa|adres|lp|l\.p|strona|nip|regon|e-mail|tel)", re.I)
ADRES = re.compile(r"^(?:ul\.|al\.|os\.|pl\.|\d{2}-\d{3}|\d+\.?\s*\|?\s*\d{2}-\d{3})", re.I)


def cmd_otwarcie(a):
    tekst = Path(a.plik).read_text(encoding="utf-8", errors="replace")
    linie = [l for l in tekst.splitlines() if l.strip() and not l.startswith("###") and not l.startswith("[[s.")]
    budzet = {}
    for m in re.finditer(r"kwot\w*[^.]{0,120}?sfinansowani\w*[^.]{0,200}?(\d{1,3}(?:[  .]\d{3})+[,.]\d{2})", tekst, re.I | re.S):
        budzet.setdefault("ogolem", liczba(m.group(1)))
    czesci, biez = {}, "1"
    blok = []
    for ln in linie:
        mc = RX_CZESC.match(ln)
        if mc and not RX_CENA.search(ln):
            biez = mc.group(1)
            blok = []
            continue
        if re.search(r"e-mail|tel\.|NIP", ln, re.I) and not RX_CENA.search(ln):
            blok.append(ln)
            continue
        m = RX_CENA.search(ln)
        if m and not re.search(r"NIP|REGON|KRS|tel", ln, re.I):
            cena = liczba(m.group(1) + "," + m.group(2))
            kand = blok + [ln]
            nazwa = None
            for x in kand:
                x2 = re.sub(r"^\s*\d+\.\s*\|?\s*", "", x).split("|")[0].strip()
                if len(re.findall(r"[A-Za-zĄĆĘŁŃÓŚŹŻąćęłńóśźż]", x2)) >= 3 and not POMIN.match(x2) and not ADRES.match(x2) \
                        and not re.search(r"e-mail|NIP", x2, re.I):
                    nazwa = x2
                    break
            czesci.setdefault(biez, []).append({"wykonawca": nazwa or "[nazwa?]", "C": cena})
            blok = []
            continue
        # nowa oferta zaczyna sie po linii z e-mail/NIP poprzedniej
        if blok and re.search(r"e-mail|NIP", blok[-1], re.I):
            blok = []
        blok.append(ln)
    if not czesci:
        sys.exit("Nie znalazłem cen w pliku - przepisz oferty ręcznie.")
    for nr, of in czesci.items():
        ceny = [o["C"] for o in of]
        print(f"## Część/zadanie {nr}: {len(of)} ofert, min {zl(min(ceny))}, mediana {zl(statistics.median(ceny))}, max {zl(max(ceny))}")
        for o in sorted(of, key=lambda o: o["C"]):
            print(f"- {o['wykonawca']}: {zl(o['C'])}")
    if budzet:
        print(f"\nKwota na sfinansowanie (pierwsze trafienie): {zl(budzet['ogolem'])} - sprawdź, której części dotyczy.")
    if a.out:
        Path(a.out).write_text(json.dumps({"czesci": czesci, "budzet": budzet}, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"\nZapisano {a.out}. Sprawdź nazwy z oryginałem (rozbiór heurystyczny), wklej oferty do specyfikacji.")


# ----------------------------------------------------------------------------- XLSX
def cmd_xlsx(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    spec = wczytaj(a.spec)
    oferty, zal = uzupelnij(spec, a.konkurencja)
    kc = kryt_ceny(spec)
    wb = Workbook()
    ws = wb.active
    ws.title = "Symulacja"
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    hdr = PatternFill("solid", fgColor="1F4E79")
    inp = PatternFill("solid", fgColor="FFF2CC")
    cienka = Side(style="thin", color="A6A6A6")
    ramka = Border(left=cienka, right=cienka, top=cienka, bottom=cienka)
    ws["A1"] = f"Symulacja punktacji: {spec.get('postepowanie', '')}"
    ws["A1"].font = F(bold=True, size=13, color="1F4E79")
    ws["A2"] = "Dokument roboczy - nie do publikacji. Żółte komórki można zmieniać (ceny, deklaracje); punkty i miejsca liczą się formułami."
    ws["A2"].font = F(italic=True, size=9, color="7F7F7F")
    # tabela kryteriow
    r0 = 4
    for j, h in enumerate(["Kryterium", "Nazwa", "Waga", "Sposób oceny"], 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill = F(bold=True, color="FFFFFF"), hdr
    for i, k in enumerate(spec["kryteria"], 1):
        for j, v in enumerate([k["id"], k["nazwa"], k["waga"], opis_krytrium(k)], 1):
            c = ws.cell(r0 + i, j, v)
            c.font, c.border = F(), ramka
    # tabela ofert
    t0 = r0 + len(spec["kryteria"]) + 3
    ids = [k["id"] for k in spec["kryteria"]]
    naglowki = ["Wykonawca"] + [f"{i}: wartość" for i in ids] + [f"{i}: pkt" for i in ids] + ["Suma pkt", "Miejsce"]
    for j, h in enumerate(naglowki, 1):
        c = ws.cell(t0, j, h)
        c.font, c.fill, c.alignment = F(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True, horizontal="center")
    n = len(oferty)
    pierwszy, ostatni = t0 + 1, t0 + n
    for i, o in enumerate(oferty):
        r = t0 + 1 + i
        c = ws.cell(r, 1, o["wykonawca"] + (" (MiD)" if o.get("my") and "mid" not in o["wykonawca"].lower() else ""))
        c.font, c.border = F(bold=bool(o.get("my"))), ramka
        for j, k in enumerate(spec["kryteria"]):
            col = 2 + j
            v = o.get(k["id"])
            if k["typ"] == "tak_nie":
                v = 1 if v else 0
            c = ws.cell(r, col, v)
            c.font, c.fill, c.border = F(color="0000FF"), inp, ramka
            if k["typ"] in ("cena", "cena_liniowa"):
                c.number_format = '#,##0.00 "zł"'
        for j, k in enumerate(spec["kryteria"]):
            vc = L(2 + j)
            x = f"{vc}{r}"
            rng = f"${vc}${pierwszy}:${vc}${ostatni}"
            W = k["waga"]
            t = k["typ"]
            if t == "cena":
                f = f"={W}*MIN({rng})/{x}"
            elif t == "cena_liniowa":
                f = f"=IF(MAX({rng})=MIN({rng}),{W},{W}*(MAX({rng})-{x})/(MAX({rng})-MIN({rng})))"
            elif t == "min":
                d = k.get("dolna")
                f = f"={W}*MIN({rng})/{x}" if d is None else f"={W}*MAX(MIN({rng}),{d})/MAX({x},{d})"
            elif t == "max":
                lim = k.get("limit")
                f = f"={W}*{x}/MAX({rng})" if not lim else f"={W}*MIN({x},{lim})/MIN(MAX({rng}),{lim})"
            elif t == "liniowy":
                lo, hi = min(k["pkt_od"], k["pkt_do"]), max(k["pkt_od"], k["pkt_do"])
                f = f"=MAX({lo},MIN({hi},{k['pkt_od']}+({x}-{k['od']})*({k['pkt_do']}-{k['pkt_od']})/({k['do']}-{k['od']})))"
            elif t == "progi":
                pr = sorted(k["progi"])
                f = (f"=IF({x}<{pr[0][0]},{k.get('ponizej', 0)},LOOKUP({x},{{{','.join(str(p) for p, _ in pr)}}},"
                     f"{{{','.join(str(pp) for _, pp in pr)}}}))")
            elif t == "tak_nie":
                f = f"=IF({x}=1,{W},0)"
            else:
                f = f"=MIN({x},{W})"
            f = f"=ROUND({f[1:]},{k.get('zaokraglenie', 2)})"
            c = ws.cell(r, 2 + len(ids) + j, f)
            c.font, c.border, c.number_format = F(), ramka, "0.00"
        s_col = 2 + 2 * len(ids)
        c = ws.cell(r, s_col, f"=SUM({L(2 + len(ids))}{r}:{L(1 + 2 * len(ids))}{r})")
        c.font, c.border, c.number_format = F(bold=True), ramka, "0.00"
        # miejsce: suma malejaco, remis -> nizsza cena
        sc = L(s_col)
        cc = L(2 + ids.index(kc["id"]))
        c = ws.cell(r, s_col + 1, f"=COUNTIF(${sc}${pierwszy}:${sc}${ostatni},\">\"&{sc}{r})+"
                                   f"COUNTIFS(${sc}${pierwszy}:${sc}${ostatni},{sc}{r},${cc}${pierwszy}:${cc}${ostatni},\"<\"&{cc}{r})+1")
        c.font, c.border = F(bold=True), ramka
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    for j in range(3, 4 + 2 * len(ids)):
        ws.column_dimensions[L(j)].width = 14
    ws.column_dimensions["D"].width = max(ws.column_dimensions["D"].width or 0, 14)
    ws.row_dimensions[t0].height = 30
    if zal:
        ws.cell(ostatni + 2, 1, "Założenia dla nieznanych deklaracji: " + "; ".join(zal)).font = F(italic=True, size=9)
    # arkusz progow (wyliczony narzedziem)
    w2 = wb.create_sheet("Progi")
    w2["A1"] = "Próg ceny MiD (wyliczony przez kryteria_tool.py przy deklaracjach z arkusza Symulacja w chwili eksportu)"
    w2["A1"].font = F(bold=True, color="1F4E79")
    for j, h in enumerate(["Pozacenowe MiD", "Pkt pozacenowe", "Najwyższa cena wygrywająca", "% najniższej ceny konk."], 1):
        c = w2.cell(3, j, h)
        c.font, c.fill = F(bold=True, color="FFFFFF"), hdr
    ceny = [o[kc["id"]] for o in oferty if not o.get("my")]
    for i, p in enumerate(poziomy_pozacenowe(spec), 4):
        of = [dict(o) for o in oferty]
        for o in of:
            if o.get("my"):
                o.update(p)
        mx = max_cena_wygrywajaca(spec, of)
        myo = next(o for o in of if o.get("my"))
        pp = sum(punkty_kryterium(k, myo.get(k["id"]), [x[k["id"]] for x in of if x.get(k["id"]) is not None])
                 for k in spec["kryteria"] if k is not kc)
        vals = [", ".join(f"{k}={v:g}" for k, v in p.items()) or "—", pp, mx or "nie wygrywa",
                (mx / min(ceny)) if mx and ceny else None]
        for j, v in enumerate(vals, 1):
            c = w2.cell(i, j, v)
            c.font = F()
            if j == 3 and isinstance(v, float):
                c.number_format = '#,##0.00 "zł"'
            if j == 4 and v:
                c.number_format = "0.0%"
    for col, wdt in zip("ABCD", (24, 16, 26, 22)):
        w2.column_dimensions[col].width = wdt
    wb.save(a.out)
    print(f"Zapisano {a.out} ({n} ofert, {len(ids)} kryteriów). Formuły przelicza Excel/LibreOffice przy otwarciu.")


def main():
    import signal
    if hasattr(signal, "SIGPIPE"):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("wzor"); s.add_argument("--out", default="kryteria.json")
    s = sub.add_parser("swz"); s.add_argument("txt")
    s = sub.add_parser("otwarcie"); s.add_argument("plik"); s.add_argument("--out")
    for nazwa in ("ocena", "prog", "scenariusze", "xlsx"):
        s = sub.add_parser(nazwa)
        s.add_argument("spec")
        s.add_argument("--konkurencja", choices=["max", "min", "jak_my"], default="max",
                       help="nieznane deklaracje konkurencji: max (ostrożnie, domyślnie), min, jak_my")
        s.add_argument("--budzet", type=float, help="kwota przeznaczona na sfinansowanie (brutto)")
        if nazwa == "scenariusze":
            s.add_argument("--udzialy", default="0.6,0.7,0.8,0.9,1.0")
        if nazwa == "xlsx":
            s.add_argument("--out", default="Symulacja_punktacji.xlsx")
    a = ap.parse_args()
    {"wzor": cmd_wzor, "swz": cmd_swz, "otwarcie": cmd_otwarcie, "ocena": cmd_ocena, "prog": cmd_prog,
     "scenariusze": cmd_scenariusze, "xlsx": cmd_xlsx}[a.cmd](a)


if __name__ == "__main__":
    main()
