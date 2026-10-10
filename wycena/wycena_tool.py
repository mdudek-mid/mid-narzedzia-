#!/usr/bin/env python3
"""wycena_tool.py - kalkulacja wyceny prac projektowych MiD i dokumenty wyceny (XLSX + MD).

Parametry metody (stawki, narzut, rezerwa, jnp, P&B) sa prywatne: KATALOG/parametry_wyceny.json
(--baza KATALOG albo MID_BAZA; repo mid-przetargi/baza_mid). Pole "parametry" w specyfikacji je nadpisuje.
  praca wlasna = rbg x stawka; pozycje jnp = jnp x jnp_wsp x jnp_stawka
  koszty zewnetrzne = kwota (+ narzut dla branz zlecanych podwykonawcom, "podwykonawca": true)
  A - pelna kalkulacja: koszt kalkulacyjny x (1 + rezerwa)
  B - rekomendowany: koszt kalkulacyjny; z --cel NETTO: pozycje sztywne bez zmian, miekkie skalowane do celu
  C - prog bolu: rbg x stawka_c + koszty zewnetrzne po kosztach wlasnych, bez rezerwy
  formula "pib" (MiD projektant u wykonawcy robot): wszystkie scenariusze x wsp_pib

Komendy:
  wzor [--out wycena.json]                      przykladowa specyfikacja
  licz SPEC.json [--cel NETTO | --cel-brutto B] kalkulacja, scenariusze, formularz zamawiajacego, kontrole
  xlsx SPEC.json --out Wycena.xlsx [--cel ...]  arkusz w stylu MiD (Formularz cenowy + Kalkulacja, formuly)
  md SPEC.json --out Wycena.md [--cel ...]      opis wyceny w stylu MiD (dane, kalkulacja, warianty, rynek, ryzyka)
  wzorce [--baza KATALOG] [--szukaj REGEX] [--lista]   pozycje z wczesniejszych wycen MiD (prywatna baza)
"""
import argparse, copy, json, math, os, re, statistics, sys
from datetime import date
from pathlib import Path

DOMYSLNE = {"vat": 0.23, "wsp_pib": 1.0, "zaokraglenie": 100.0, "max_ciecie": 0.40}
WYMAGANE = ("stawka", "stawka_c", "rezerwa", "narzut")

SZTYWNE = r"most|wiadukt|obiekt|kładk|kladk|przepust|tunel|konstrukc|geolog|geotech|odwiert|wierc|geodez|mapa|map[ay] do cel|podzia|decyzj|zrid|pozwoleni|operat wodno|raport o oddz|środowisk|srodowisk"
MIEKKIE = r"kompletac|kompilac|egzemplarz|wydruk|archiw|organizac\w* ruchu|\bsor\b|\btor\b|nadz[oó]r|pobyt|koordynac|spotkan|narad|rad[ay] techn|kierowanie|zarządzanie|zarzadzanie|wizualizac|prezentac|materiały przetarg|materialy przetarg|odpowiedzi na pytania"

PRZYKLAD = {
    "tytul": "Rozbiórka i budowa wiaduktu (przykład)",
    "zamawiajacy": "Zarząd Dróg Wojewódzkich w …",
    "znak": "…", "platforma": "platformazakupowa.pl/transakcja/…",
    "termin_ofert": "RRRR-MM-DD GG:MM", "wadium": 0, "termin_realizacji": "… mies.", "zwiazanie": "…",
    "kryteria": "cena 60 / doświadczenie 40", "formula": "projektowa",
    "warunki_sprawdzone": False,
    "parametry": {},
    "praca": [
        {"id": "A", "zakres": "PB + PT branża mostowa (obiekt, rozbiórka)", "rbg": 400, "poz": "11", "sztywna": True},
        {"id": "B", "zakres": "Koordynacja, narady, kompletacja", "rbg": 80, "poz": {"11": 0.5, "12": 0.5}},
        {"id": "N", "zakres": "Nadzór autorski (metoda jnp)", "jnp": 1200, "poz": "29"}
    ],
    "zewnetrzne": [
        {"zakres": "Badania geotechniczne (odwierty, laboratorium)", "kwota": 25000, "poz": "1", "sztywna": True},
        {"zakres": "Branża drogowa (podwykonawca)", "kwota": 30000, "podwykonawca": True, "poz": "11"}
    ],
    "formularz": [
        {"poz": "1", "nazwa": "Rozpoznanie podłoża gruntowego", "jm": "komplet", "ilosc": 1},
        {"poz": "11", "nazwa": "Projekt budowlany", "jm": "komplet", "ilosc": 1},
        {"poz": "12", "nazwa": "Projekt techniczny", "jm": "komplet", "ilosc": 1},
        {"poz": "29", "nazwa": "Nadzór autorski", "jm": "komplet", "ilosc": 1},
        {"poz": "30", "nazwa": "Kwota tymczasowa 5%", "procent": 0.05}
    ],
    "rynek": {"budzet_brutto": None, "oferty_brutto": [], "opis": ["…"]},
    "dane": ["…"], "uwagi": ["…"], "ryzyka": ["…"], "do_zrobienia": ["…"]
}


# ----------------------------------------------------------------------------- formatowanie
def zl(v, waluta=True):
    if v is None:
        return "—"
    s = f"{v:,.2f}".replace(",", " ").replace(".", ",")
    return s + (" zł" if waluta else "")


def zl0(v):
    return "—" if v is None else f"{v:,.0f}".replace(",", " ")


def proc(x, d=1):
    return "—" if x is None else f"{x * 100:.{d}f}%".replace(".", ",")


def lp(v):
    return f"{v:g}".replace(".", ",")


# ----------------------------------------------------------------------------- obliczenia
def parametry_firmy(baza=None):
    f = Path(baza or os.environ.get("MID_BAZA") or "baza_mid") / "parametry_wyceny.json"
    return json.loads(f.read_text(encoding="utf-8")).get("parametry", {}) if f.exists() else {}


def wczytaj(p, baza=None):
    s = json.loads(Path(p).read_text(encoding="utf-8"))
    s["parametry"] = {**DOMYSLNE, **parametry_firmy(baza), **(s.get("parametry") or {})}
    brak = [k for k in WYMAGANE if s["parametry"].get(k) is None]
    if any("jnp" in x for x in s.get("praca") or []):
        brak += [k for k in ("jnp_wsp", "jnp_stawka") if s["parametry"].get(k) is None]
    if brak:
        sys.exit(f"Brak parametrów wyceny: {', '.join(brak)}. Dołącz prywatne repo mid-przetargi (baza_mid/parametry_wyceny.json, "
                 "--baza albo MID_BAZA) albo podaj je w \"parametry\" specyfikacji.")
    s.setdefault("praca", [])
    s.setdefault("zewnetrzne", [])
    s.setdefault("formularz", [])
    for i, it in enumerate(s["praca"]):
        it.setdefault("id", f"P{i + 1}")
    for i, it in enumerate(s["zewnetrzne"]):
        it.setdefault("id", f"Z{i + 1}")
    return s


def czy_sztywna(it):
    if "sztywna" in it:
        return bool(it["sztywna"]), "spec"
    t = it.get("zakres", "").lower()
    if re.search(MIEKKIE, t):
        return False, "auto"
    if re.search(SZTYWNE, t):
        return True, "auto"
    return False, "auto"


def udzialy(it):
    p = it.get("poz")
    if p is None:
        return {}
    if isinstance(p, dict):
        s = sum(p.values())
        return {str(k): v / s for k, v in p.items()}
    return {str(p): 1.0}


def oblicz(s, cel=None, cel_brutto=None):
    P = s["parametry"]
    pozycje = []
    for it in s["praca"]:
        if "jnp" in it:
            koszt = it["jnp"] * it.get("jnp_wsp", P["jnp_wsp"]) * it.get("jnp_stawka", P["jnp_stawka"])
            koszt_c = koszt
            rbg = None
        else:
            rbg = float(it.get("rbg") or 0)
            koszt = rbg * P["stawka"]
            koszt_c = rbg * P["stawka_c"]
        szt, zr = czy_sztywna(it)
        pozycje.append({**it, "rodzaj": "praca", "rbg": rbg, "koszt": koszt, "koszt_c": koszt_c, "sztywna": szt, "sztywna_zrodlo": zr})
    for it in s["zewnetrzne"]:
        narz = it.get("narzut", P["narzut"] if it.get("podwykonawca") else 0.0)
        koszt = float(it["kwota"]) * (1 + narz)
        szt, zr = czy_sztywna(it)
        pozycje.append({**it, "rodzaj": "zewn", "narzut_eff": narz, "koszt": koszt, "koszt_c": float(it["kwota"]),
                        "sztywna": szt, "sztywna_zrodlo": zr})
    kalk = sum(x["koszt"] for x in pozycje)
    wsp = P["wsp_pib"] if s.get("formula") == "pib" else 1.0
    procent = sum(f.get("procent", 0) for f in s["formularz"])
    A = kalk * (1 + P["rezerwa"]) * wsp
    C = sum(x["koszt_c"] for x in pozycje) * wsp
    # cel dotyczy ceny oferty netto (z pozycjami procentowymi, np. kwota tymczasowa)
    if cel_brutto:
        cel = cel_brutto / (1 + P["vat"])
    uwagi = []
    if cel:
        baza_cel = cel / (1 + procent)
        R = sum(x["koszt"] for x in pozycje if x["sztywna"]) * wsp
        S = sum(x["koszt"] for x in pozycje if not x["sztywna"]) * wsp
        if baza_cel >= R + S:
            f_m = f_s = baza_cel / (R + S)
        else:
            f_m = (baza_cel - R) / S if S else 0
            f_s = 1.0
            if f_m < 1 - P["max_ciecie"]:
                f_m = 1 - P["max_ciecie"]
                f_s = (baza_cel - f_m * S) / R if R else 0
                uwagi.append(f"Cel {zl(cel)} wymaga cięcia pozycji miękkich o więcej niż {proc(P['max_ciecie'], 0)} — przyjęto "
                             f"−{proc(P['max_ciecie'], 0)} na miękkich i {proc(f_s - 1).replace("-", "−")} na sztywnych (mostowe PB, geodezja, decyzje, geologia).")
        for x in pozycje:
            x["B"] = x["koszt"] * wsp * (f_s if x["sztywna"] else f_m)
        wsp_m, wsp_s = f_m, f_s
    else:
        for x in pozycje:
            x["B"] = x["koszt"] * wsp
        wsp_m = wsp_s = 1.0
    B_baza = sum(x["B"] for x in pozycje)
    form = formularz(s, pozycje, cel, procent, uwagi)
    return {"pozycje": pozycje, "kalk": kalk, "wsp": wsp, "A": A * (1 + procent), "C": C * (1 + procent),
            "A_baza": A, "C_baza": C, "B_baza": B_baza, "procent": procent, "cel": cel, "wsp_miekkie": wsp_m,
            "wsp_sztywne": wsp_s, "form": form, "B": form["netto"] if form["wiersze"] else B_baza * (1 + procent),
            "uwagi": uwagi, "rbg": sum(x["rbg"] or 0 for x in pozycje if x["rodzaj"] == "praca")}


def zaokr(v, krok):
    return round(v / krok) * krok if krok else round(v, 2)


def formularz(s, pozycje, cel, procent, uwagi):
    P = s["parametry"]
    F = s["formularz"]
    if not F:
        return {"wiersze": [], "netto": None}
    wart = {}
    for x in pozycje:
        for poz, u in udzialy(x).items():
            wart[poz] = wart.get(poz, 0) + x["B"] * u
    znane = {str(f["poz"]) for f in F} | {str(f.get("grupa")) for f in F if f.get("grupa")}
    for poz in wart:
        if poz not in znane:
            uwagi.append(f"Pozycje kalkulacji przypisane do „{poz}”, której nie ma w formularzu — ich wartość nie trafi do ceny.")
    for x in pozycje:
        if not udzialy(x):
            uwagi.append(f"{x['id']} „{x.get('zakres', '')[:60]}” nie ma przypisanej pozycji formularza (pole poz).")
    grupy = {}
    for f in F:
        if f.get("grupa") and f.get("udzial") is not None:
            grupy.setdefault(str(f["grupa"]), []).append(f)
    for g, fs in grupy.items():
        su = sum(f["udzial"] for f in fs)
        if abs(su - 1) > 1e-6:
            uwagi.append(f"Udziały w grupie „{g}” sumują się do {proc(su)}, a nie 100%.")
    wiersze = []
    for f in F:
        poz = str(f["poz"])
        w = {"poz": poz, "nazwa": f.get("nazwa", ""), "jm": f.get("jm", ""), "ilosc": f.get("ilosc", 1),
             "opcja": bool(f.get("opcja")), "procent": f.get("procent"), "uwagi": f.get("uwagi", "")}
        if f.get("procent") is not None:
            w["wartosc"] = None
        elif f.get("grupa") and f.get("udzial") is not None:
            w["wartosc"] = wart.get(str(f["grupa"]), 0) * f["udzial"]
        else:
            w["wartosc"] = wart.get(poz, 0)
        wiersze.append(w)
    # zaokraglenie: ryczalty do kroku, ceny jednostkowe do 10 zl przy ilosci > 1
    krok = P["zaokraglenie"]
    for w in wiersze:
        if w["wartosc"] is None:
            continue
        il = w["ilosc"] or 1
        if il != 1:
            w["cena_jedn"] = zaokr(w["wartosc"] / il, 10)
        else:
            w["cena_jedn"] = zaokr(w["wartosc"], krok if w["wartosc"] >= 10 * krok else 10)
        w["wartosc"] = round(w["cena_jedn"] * il, 2)
    suma = lambda: sum(w["wartosc"] for w in wiersze if w["wartosc"] is not None)
    if cel:
        # roznica po zaokragleniu na najwieksza pozycje ryczaltowa
        roznica = cel / (1 + procent) - suma()
        ryczalt = [w for w in wiersze if w["wartosc"] is not None and (w["ilosc"] or 1) == 1]
        if ryczalt and abs(roznica) >= 0.01:
            w = max(ryczalt, key=lambda w: w["wartosc"])
            w["cena_jedn"] = round(w["cena_jedn"] + roznica, 2)
            w["wartosc"] = w["cena_jedn"]
    baza = suma()
    for w in wiersze:
        if w["procent"] is not None:
            w["cena_jedn"] = w["wartosc"] = round(baza * w["procent"], 2)
    for w in wiersze:
        if not w["wartosc"]:
            uwagi.append(f"Pozycja formularza {w['poz']} „{w['nazwa'][:50]}” ma wartość 0 — sprawdź SWZ (często cena 0 = odrzucenie oferty). "
                         "Przypisz jej pozycję kalkulacji albo część pozycji miękkiej.")
    netto = round(sum(w["wartosc"] for w in wiersze), 2)
    podst = round(sum(w["wartosc"] for w in wiersze if not w["opcja"]), 2)
    return {"wiersze": wiersze, "netto": netto, "podstawowe": podst, "opcje": round(netto - podst, 2), "baza": baza}


def kontrole(s, R):
    P = s["parametry"]
    out = list(R["uwagi"])
    if not s.get("warunki_sprawdzone"):
        out.insert(0, "Warunki udziału nie są oznaczone jako sprawdzone (\"warunki_sprawdzone\": true). Przed wyceną potwierdź "
                      "referencje i kadrę (skill referencje-i-kadra) — wycena bez spełnienia warunków to strata czasu.")
    wad = s.get("wadium") or 0
    if wad:
        out.append(f"Wadium {zl(wad)}: " + ("przelew gotówkowy na rachunek zamawiającego (< 10 tys. zł)." if wad < 10000
                                           else "gwarancja wadialna (≥ 10 tys. zł) — zamówić z wyprzedzeniem."))
    elif "wadium" in s:
        out.append("Wadium: nie jest wymagane.")
    else:
        out.append("Wadium: brak danych w specyfikacji — sprawdź SWZ.")
    war = [x for x in R["pozycje"] if x.get("warunkowa")]
    if war:
        out.append("Pozycje warunkowe wliczone w cenę (zasada MiD: nie wykazujemy pozycji warunkowych): " +
                   ", ".join(x["id"] for x in war) + ".")
    auto = [x for x in R["pozycje"] if x["sztywna_zrodlo"] == "auto"]
    if auto and R["cel"]:
        out.append("Klasyfikacja sztywna/miękka nadana automatycznie dla: " + ", ".join(
            f"{x['id']}={'S' if x['sztywna'] else 'M'}" for x in auto) + " — sprawdź i ustaw \"sztywna\" w specyfikacji.")
    if R["B"] < R["C"]:
        out.append(f"Scenariusz B ({zl(R['B'])}) jest poniżej progu bólu C ({zl(R['C'])}) — oferta poniżej kosztów.")
    ryn = s.get("rynek") or {}
    bud = ryn.get("budzet_brutto")
    Bb = R["B"] * (1 + P["vat"])
    if bud:
        if Bb < 0.7 * bud:
            out.append(f"Cena B brutto {zl(Bb)} to {proc(Bb / bud)} budżetu — możliwe wezwanie do wyjaśnień RNC (art. 224 ust. 2 pkt 1 Pzp); "
                       "przygotuj kalkulację do wyjaśnień (arkusz Kalkulacja).")
        if Bb > bud:
            out.append(f"Cena B brutto {zl(Bb)} przekracza kwotę na sfinansowanie ({zl(bud)}, {proc(Bb / bud)}).")
    if s.get("formula") == "pib" and s.get("wartosc_robot_netto"):
        wr = s["wartosc_robot_netto"]
        ref = P.get("pib_procent_robot")
        out.append(f"P&B: dokumentacja = {proc(R['B'] / wr, 2)} wartości robót netto" +
                   (f" (punkt odniesienia MiD: {proc(ref[0], 1)}–{proc(ref[1], 1)})." if ref else "."))
    return out


def rynek_tabela(s, R):
    P = s["parametry"]
    ryn = s.get("rynek") or {}
    of = sorted(ryn.get("oferty_brutto") or [])
    linie = []
    ref = []
    if ryn.get("budzet_brutto"):
        ref.append(("budżet", ryn["budzet_brutto"]))
    if of:
        ref += [("min ofert", of[0]), ("mediana ofert", statistics.median(of))]
    if ryn.get("min_brutto"):
        ref.append(("min (rynek)", ryn["min_brutto"]))
    if ryn.get("mediana_brutto"):
        ref.append(("mediana (rynek)", ryn["mediana_brutto"]))
    if not ref:
        return ""
    linie.append("| scenariusz | brutto | " + " | ".join(f"% {n}" for n, _ in ref) + (" | miejsce cenowo |" if of else " |"))
    linie.append("|---|---|" + "---|" * len(ref) + ("---|" if of else ""))
    for n, v in (("A", R["A"]), ("B", R["B"]), ("C", R["C"])):
        b = v * (1 + P["vat"])
        m = (f" {1 + sum(1 for o in of if o < b)}/{len(of) + 1} |" if of else "")
        linie.append(f"| {n} | {zl(b)} | " + " | ".join(proc(b / r) for _, r in ref) + " |" + m)
    return "\n".join(linie)


# ----------------------------------------------------------------------------- komendy tekstowe
def cmd_wzor(a):
    Path(a.out).write_text(json.dumps(PRZYKLAD, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Zapisano {a.out}. Pozycje formularza przepisz z formularza cenowego zamawiającego (poz = numer pozycji). "
          "Parametry metody (stawki, narzut, rezerwa) bierze z baza_mid/parametry_wyceny.json.")


def cmd_licz(a):
    s = wczytaj(a.spec, a.baza)
    R = oblicz(s, a.cel, a.cel_brutto)
    P = s["parametry"]
    print(f"# Wycena: {s.get('tytul', '')}\n")
    print(f"Praca własna {zl0(R['rbg'])} rbg × {zl0(P['stawka'])} zł = {zl(sum(x['koszt'] for x in R['pozycje'] if x['rodzaj'] == 'praca'))}; "
          f"koszty zewnętrzne {zl(sum(x['koszt'] for x in R['pozycje'] if x['rodzaj'] == 'zewn'))}; "
          f"koszt kalkulacyjny {zl(R['kalk'])}" + (f"; × {lp(R['wsp'])} (P&B)" if R["wsp"] != 1 else "") +
          (f"; pozycje procentowe formularza +{proc(R['procent'])}" if R["procent"] else "") + ".\n")
    print("| scenariusz | netto | brutto | opis |\n|---|---|---|---|")
    vat = 1 + P["vat"]
    print(f"| A — pełna kalkulacja | {zl(R['A'])} | {zl(R['A'] * vat)} | koszt + rezerwa {proc(P['rezerwa'], 0)} |")
    opisB = (f"cel; miękkie ×{lp(round(R['wsp_miekkie'], 3))}, sztywne ×{lp(round(R['wsp_sztywne'], 3))}" if R["cel"]
             else "koszt kalkulacyjny bez rezerwy — do kalibracji rynkowej i decyzji Marcina")
    print(f"| **B — rekomendowany** | **{zl(R['B'])}** | **{zl(R['B'] * vat)}** | {opisB} |")
    print(f"| C — próg bólu | {zl(R['C'])} | {zl(R['C'] * vat)} | {zl0(P['stawka_c'])} zł/rbg, koszty zewn. po kosztach własnych, bez rezerwy |")
    rt = rynek_tabela(s, R)
    if rt:
        print("\n## Na tle rynku\n\n" + rt)
    if R["form"]["wiersze"]:
        print("\n## Formularz (scenariusz B)\n")
        print("| poz. | przedmiot | j.m. | ilość | cena jedn. netto | wartość netto |\n|---|---|---|---|---|---|")
        for w in R["form"]["wiersze"]:
            print(f"| {w['poz']} | {w['nazwa'][:70]}{' (opcja)' if w['opcja'] else ''} | {w['jm']} | {lp(w['ilosc'] or 1)} | "
                  f"{zl(w['cena_jedn'])} | {zl(w['wartosc'])} |")
        f = R["form"]
        print(f"| | RAZEM netto | | | | {zl(f['netto'])} |" + (f"\n| | w tym opcje | | | | {zl(f['opcje'])} |" if f["opcje"] else ""))
        print(f"| | VAT {proc(P['vat'], 0)} | | | | {zl(f['netto'] * P['vat'])} |\n| | RAZEM brutto | | | | {zl(f['netto'] * vat)} |")
    print("\n## Kontrole\n")
    for u in kontrole(s, R):
        print(f"- {u}")


# ----------------------------------------------------------------------------- XLSX
def cmd_xlsx(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    s = wczytaj(a.spec, a.baza)
    R = oblicz(s, a.cel, a.cel_brutto)
    P = s["parametry"]
    wb = Workbook()
    F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    HDR = PatternFill("solid", fgColor="1F4E79")
    SUB = PatternFill("solid", fgColor="D9E1F2")
    cienka = Side(style="thin", color="A6A6A6")
    RAM = Border(left=cienka, right=cienka, top=cienka, bottom=cienka)
    ZL = '#,##0.00 "zł"'
    WRAP = Alignment(wrap_text=True, vertical="top")

    def naglowek_tabeli(ws, r, nazwy):
        for j, h in enumerate(nazwy, 1):
            c = ws.cell(r, j, h)
            c.font, c.fill, c.alignment, c.border = F(bold=True, color="FFFFFF"), HDR, Alignment(wrap_text=True, horizontal="center", vertical="center"), RAM

    def kom(ws, r, c, v, fmt=None, bold=False, color=None, fill=None, wrap=False):
        x = ws.cell(r, c, v)
        x.font = F(bold=bold, color=color) if color else F(bold=bold)
        x.border = RAM
        if fmt:
            x.number_format = fmt
        if fill:
            x.fill = fill
        if wrap:
            x.alignment = WRAP
        return x

    # ---------------- Kalkulacja
    wk = wb.active
    wk.title = "Kalkulacja"
    wk["A1"] = f"KALKULACJA — {s.get('tytul', '')}"
    wk["A1"].font = F(bold=True, size=12, color="1F4E79")
    par = [("stawka zł/rbg", P["stawka"], "0"), ("stawka próg bólu zł/rbg", P["stawka_c"], "0"), ("rezerwa ryzyka", P["rezerwa"], "0%"),
           ("narzut na podwykonawców", P["narzut"], "0%"), ("VAT", P["vat"], "0%"), ("współczynnik P&B", R["wsp"], "0.00"),
           ("pozycje procentowe formularza", R["procent"], "0.0%")]
    PAR = {}
    for i, (n, v, fm) in enumerate(par):
        r = 3 + i
        wk.cell(r, 1, n).font = F()
        c = wk.cell(r, 2, v)
        c.font, c.number_format, c.fill = F(color="0000FF"), fm, PatternFill("solid", fgColor="FFF2CC")
        PAR[n] = f"$B${r}"
    st, stc, rez, nar, vat_c, wsp_c, proc_c = (PAR[n] for n, _, _ in par)
    r = 3 + len(par) + 1
    wk.cell(r, 1, "Współczynnik pozycji miękkich (scenariusz B)").font = F()
    c = wk.cell(r, 2, round(R["wsp_miekkie"], 6)); c.font, c.number_format = F(color="0000FF"), "0.000"; FM = f"$B${r}"
    wk.cell(r + 1, 1, "Współczynnik pozycji sztywnych (scenariusz B)").font = F()
    c = wk.cell(r + 1, 2, round(R["wsp_sztywne"], 6)); c.font, c.number_format = F(color="0000FF"), "0.000"; FS = f"$B${r + 1}"
    r += 3
    kol = ["Lp.", "Zakres prac", "Poz. formularza", "rbg / jnp", "Koszt netto", "Sztywna", "Scenariusz B netto", "Uwagi / założenia"]
    wk.cell(r, 1, "PRACA WŁASNA").font = F(bold=True, color="1F4E79")
    r += 1
    naglowek_tabeli(wk, r, kol)
    r0 = r + 1
    praca = [x for x in R["pozycje"] if x["rodzaj"] == "praca"]
    for i, x in enumerate(praca):
        rr = r0 + i
        kom(wk, rr, 1, x["id"])
        kom(wk, rr, 2, x.get("zakres", ""), wrap=True)
        kom(wk, rr, 3, ", ".join(f"{k}" + (f" ({proc(v, 0)})" if v != 1 else "") for k, v in udzialy(x).items()))
        if x.get("jnp") is not None:
            kom(wk, rr, 4, x["jnp"], "#,##0", color="0000FF")
            kom(wk, rr, 5, f"=D{rr}*{x.get('jnp_wsp', P['jnp_wsp'])}*{x.get('jnp_stawka', P['jnp_stawka'])}*{wsp_c}", ZL)
            uw = f"metoda jnp: Σjnp × {lp(x.get('jnp_wsp', P['jnp_wsp']))} × {lp(x.get('jnp_stawka', P['jnp_stawka']))} zł. " + x.get("uwagi", "")
        else:
            kom(wk, rr, 4, x["rbg"], "#,##0", color="0000FF")
            kom(wk, rr, 5, f"=D{rr}*{st}*{wsp_c}", ZL)
            uw = x.get("uwagi", "")
        kom(wk, rr, 6, "tak" if x["sztywna"] else "nie")
        kom(wk, rr, 7, f'=E{rr}*IF(F{rr}="tak",{FS},{FM})', ZL)
        kom(wk, rr, 8, (uw + (" [pozycja warunkowa — wliczona w cenę]" if x.get("warunkowa") else "")).strip(), wrap=True)
    r1 = r0 + len(praca) - 1
    rs_p = r1 + 1
    kom(wk, rs_p, 2, "RAZEM praca własna", bold=True, fill=SUB)
    jnp_rows = [r0 + i for i, x in enumerate(praca) if x.get("jnp") is not None]
    kom(wk, rs_p, 4, f"=SUM(D{r0}:D{r1})" + "".join(f"-D{j}" for j in jnp_rows), "#,##0", bold=True, fill=SUB)
    kom(wk, rs_p, 8, "rbg bez pozycji jnp", wrap=True)
    kom(wk, rs_p, 5, f"=SUM(E{r0}:E{r1})", ZL, bold=True, fill=SUB)
    kom(wk, rs_p, 7, f"=SUM(G{r0}:G{r1})", ZL, bold=True, fill=SUB)
    r = rs_p + 2
    wk.cell(r, 1, "KOSZTY ZEWNĘTRZNE (podwykonawcy, badania, sprzęt, opłaty)").font = F(bold=True, color="1F4E79")
    r += 1
    naglowek_tabeli(wk, r, ["Lp.", "Zakres", "Poz. formularza", "Kwota netto", "Narzut", "Koszt netto", "Sztywna", "Scenariusz B netto", "Uwagi"])
    z0 = r + 1
    zew = [x for x in R["pozycje"] if x["rodzaj"] == "zewn"]
    for i, x in enumerate(zew):
        rr = z0 + i
        kom(wk, rr, 1, x["id"])
        kom(wk, rr, 2, x.get("zakres", ""), wrap=True)
        kom(wk, rr, 3, ", ".join(f"{k}" + (f" ({proc(v, 0)})" if v != 1 else "") for k, v in udzialy(x).items()))
        kom(wk, rr, 4, float(x["kwota"]), ZL, color="0000FF")
        kom(wk, rr, 5, f"={nar}" if x.get("podwykonawca") and "narzut" not in x else x["narzut_eff"], "0%")
        kom(wk, rr, 6, f"=D{rr}*(1+E{rr})*{wsp_c}", ZL)
        kom(wk, rr, 7, "tak" if x["sztywna"] else "nie")
        kom(wk, rr, 8, f'=F{rr}*IF(G{rr}="tak",{FS},{FM})', ZL)
        kom(wk, rr, 9, x.get("uwagi", ""), wrap=True)
    z1 = z0 + max(len(zew), 1) - 1
    rs_z = z1 + 1
    kom(wk, rs_z, 2, "RAZEM koszty zewnętrzne", bold=True, fill=SUB)
    kom(wk, rs_z, 4, f"=SUM(D{z0}:D{z1})", ZL, bold=True, fill=SUB)
    kom(wk, rs_z, 6, f"=SUM(F{z0}:F{z1})", ZL, bold=True, fill=SUB)
    kom(wk, rs_z, 8, f"=SUM(H{z0}:H{z1})", ZL, bold=True, fill=SUB)
    r = rs_z + 2
    wiersze_sum = [
        ("Koszt kalkulacyjny (praca własna + koszty zewnętrzne)", f"=E{rs_p}+F{rs_z}"),
        ("A — pełna kalkulacja: koszt z rezerwą ryzyka (z pozycjami procentowymi formularza)", f"=(E{rs_p}+F{rs_z})*(1+{rez})*(1+{proc_c})"),
        ("B — rekomendowany: kalkulacja wg współczynników (z pozycjami procentowymi)", f"=(G{rs_p}+H{rs_z})*(1+{proc_c})"),
    ]
    rbg_c = f"D{rs_p}"
    jnp_val = "+".join(f"E{j}" for j in jnp_rows) or "0"
    wiersze_sum.append(("C — próg bólu: rbg × stawka C + koszty zewnętrzne po kosztach własnych, bez rezerwy",
                        f"=({rbg_c}*{stc}*{wsp_c}+{jnp_val}+D{rs_z}*{wsp_c})*(1+{proc_c})"))
    SC = {}
    for n, f in wiersze_sum:
        kom(wk, r, 2, n, bold=True)
        kom(wk, r, 5, f, ZL, bold=True)
        SC[n[0]] = f"Kalkulacja!$E${r}"
        r += 1
    wk.cell(r + 1, 1, "Opracowanie: Pracownia Projektowa MiD | Data: " + date.today().strftime("%d.%m.%Y") +
            " | Dokument roboczy – wewnętrzny, nie do publikacji").font = F(italic=True, size=9, color="7F7F7F")
    for col, wd in zip("ABCDEFGHI", (6, 62, 14, 12, 16, 10, 18, 40, 30)):
        wk.column_dimensions[col].width = wd

    # ---------------- Formularz cenowy
    wf = wb.create_sheet("Formularz cenowy", 0)
    wf["A1"] = "WYCENA PRAC PROJEKTOWYCH"
    wf["A1"].font = F(bold=True, size=13, color="1F4E79")
    wf["A2"] = s.get("tytul", "")
    wf["A2"].font = F(bold=True)
    wf["A3"] = "  |  ".join(x for x in [f"Zamawiający: {s.get('zamawiajacy', '')}", f"Postępowanie znak: {s.get('znak', '')}", s.get("platforma", "")] if x)
    wad = s.get("wadium") or 0
    wf["A4"] = "  |  ".join([f"Termin składania ofert: {s.get('termin_ofert', '')}",
                             f"Wadium: {zl(wad) + (' (gwarancja wadialna)' if wad >= 10000 else ' (przelew)') if wad else 'brak'}",
                             f"Termin realizacji: {s.get('termin_realizacji', '')}", f"Związanie ofertą do: {s.get('zwiazanie', '')}"])
    for c in ("A3", "A4"):
        wf[c].font = F(size=9)
    wf["A6"] = "Tabela zgodna z formularzem cenowym zamawiającego — scenariusz B. Ceny jednostkowe (niebieskie) do kalibracji przed złożeniem oferty."
    wf["A6"].font = F(italic=True, size=9)
    r = 8
    naglowek_tabeli(wf, r, ["Lp.", "Przedmiot zamówienia", "J.m.", "Ilość", "Cena jedn. netto", "Wartość netto", "Uwagi / założenia"])
    f0 = r + 1
    wiersze = R["form"]["wiersze"]
    proc_rows = []
    for i, w in enumerate(wiersze):
        rr = f0 + i
        kom(wf, rr, 1, w["poz"])
        kom(wf, rr, 2, w["nazwa"] + (" (opcja)" if w["opcja"] else ""), wrap=True)
        kom(wf, rr, 3, w["jm"])
        kom(wf, rr, 4, w["ilosc"] or 1, "#,##0.##")
        if w["procent"] is not None:
            proc_rows.append((rr, w["procent"]))
            kom(wf, rr, 5, None, ZL)
        else:
            kom(wf, rr, 5, w["cena_jedn"], ZL, color="0000FF")
            kom(wf, rr, 6, f"=D{rr}*E{rr}", ZL)
        kom(wf, rr, 7, w.get("uwagi", ""), wrap=True)
    f1 = f0 + len(wiersze) - 1
    nproc = [rr for rr, _ in proc_rows]
    # pozycja procentowa liczona od sumy pozycji nieprocentowych (zakresy bez wierszy procentowych - bez odwolan cyklicznych)
    zakresy, start = [], None
    for x in range(f0, f1 + 2):
        if x <= f1 and x not in nproc:
            start = x if start is None else start
        elif start is not None:
            zakresy.append(f"F{start}:F{x - 1}")
            start = None
    suma_np = "SUM(" + ",".join(zakresy) + ")" if zakresy else "0"
    for rr, p in proc_rows:
        wf.cell(rr, 6).value = f"=ROUND({suma_np}*{p},2)"
        wf.cell(rr, 6).number_format = ZL
        wf.cell(rr, 6).font = F()
        wf.cell(rr, 6).border = RAM
        wf.cell(rr, 7).value = (wf.cell(rr, 7).value or "") + f" {proc(p)} sumy pozycji nieprocentowych".strip()
    rn = f1 + 1
    kom(wf, rn, 2, "RAZEM NETTO", bold=True, fill=SUB)
    kom(wf, rn, 6, f"=SUM(F{f0}:F{f1})", ZL, bold=True, fill=SUB)
    kom(wf, rn + 1, 2, f"Podatek VAT {proc(P['vat'], 0)}")
    kom(wf, rn + 1, 6, f"=ROUND(F{rn}*Kalkulacja!{vat_c},2)", ZL)
    kom(wf, rn + 2, 2, "OGÓŁEM BRUTTO", bold=True, fill=SUB)
    kom(wf, rn + 2, 6, f"=F{rn}+F{rn + 1}", ZL, bold=True, fill=SUB)
    r = rn + 3
    opcje = [f0 + i for i, w in enumerate(wiersze) if w["opcja"]]
    if opcje:
        kom(wf, r, 2, "w tym zakres podstawowy netto")
        kom(wf, r, 6, f"=F{rn}-" + "-".join(f"F{x}" for x in opcje), ZL)
        kom(wf, r + 1, 2, "w tym opcje netto")
        kom(wf, r + 1, 6, "=" + "+".join(f"F{x}" for x in opcje), ZL)
        r += 2
    kom(wf, r, 2, "Kontrola: formularz − kalkulacja B (różnica z zaokrągleń)")
    kom(wf, r, 6, f"=F{rn}-{SC['B']}", ZL)
    r += 3
    wf.cell(r, 1, "SCENARIUSZE CENOWE").font = F(bold=True, color="1F4E79")
    r += 1
    ryn = s.get("rynek") or {}
    of = sorted(ryn.get("oferty_brutto") or [])
    ref = []
    if ryn.get("budzet_brutto"):
        ref.append(("% budżetu", ryn["budzet_brutto"]))
    if of:
        ref += [("% min ofert", of[0]), ("% mediany ofert", statistics.median(of))]
    naglowek_tabeli(wf, r, ["", "Scenariusz", "", "", "Netto", "Brutto", "Komentarz"] + [n for n, _ in ref])
    kom_A = f"Pełny koszt kalkulacyjny ({zl0(R['rbg'])} rbg × stawka + koszty zewnętrzne) z rezerwą ryzyka {proc(P['rezerwa'], 0)}."
    kom_B = ((f"Cel cenowy: pozycje miękkie ×{lp(round(R['wsp_miekkie'], 3))}, sztywne (mostowe, geodezja, decyzje, geologia) "
              + ("bez cięć." if R["wsp_sztywne"] >= 1 else f"×{lp(round(R['wsp_sztywne'], 3))} — CEL PONIŻEJ KOSZTÓW SZTYWNYCH.")) if R["cel"]
             else "Koszt kalkulacyjny bez rezerwy — skalibrować rynkowo (WIEDZA_Rynek, wyniki) i uzgodnić z Marcinem.")
    kom_C = f"Stawka {zl0(P['stawka_c'])} zł/rbg, koszty zewnętrzne po kosztach własnych, bez rezerwy. Nie schodzić poniżej bez dodatkowej analizy."
    for n, f, k in (("A — bezpieczny", f"={SC['A']}", kom_A), ("B — rekomendowany (ta oferta)", f"=F{rn}", kom_B),
                    ("C — dolna granica (próg bólu)", f"={SC['C']}", kom_C)):
        r += 1
        kom(wf, r, 2, n, bold=n.startswith("B"))
        kom(wf, r, 5, f, ZL, bold=n.startswith("B"))
        kom(wf, r, 6, f"=ROUND(E{r}*(1+Kalkulacja!{vat_c}),2)", ZL)
        kom(wf, r, 7, k, wrap=True)
        for j, (_, v) in enumerate(ref):
            kom(wf, r, 8 + j, f"=F{r}/{v}", "0.0%")
    r += 2
    teksty = [("UWAGI I RYZYKA DO WYCENY", (s.get("uwagi") or []) + (s.get("ryzyka") or []))]
    if ryn.get("opis"):
        teksty.append(("RYNEK", ryn["opis"]))
    for tyt, lst in teksty:
        wf.cell(r, 1, tyt).font = F(bold=True, color="1F4E79")
        r += 1
        for i, t in enumerate(lst, 1):
            c = wf.cell(r, 1, f"{i}. {t}")
            c.font, c.alignment = F(), Alignment(wrap_text=True, vertical="top")
            wf.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            wf.row_dimensions[r].height = max(15, 13 * math.ceil(len(t) / 150))
            r += 1
        r += 1
    wf.cell(r, 1, "Legenda: pola w kolorze niebieskim (cena jednostkowa netto, rbg, kwoty) są przeznaczone do edycji / kalibracji przed złożeniem oferty.").font = F(italic=True, size=9)
    wf.cell(r + 1, 1, "Opracowanie: Pracownia Projektowa MiD | Data: " + date.today().strftime("%d.%m.%Y") +
            " | Dokument roboczy – wewnętrzny, nie do publikacji").font = F(italic=True, size=9, color="7F7F7F")
    for col, wd in zip("ABCDEFGHIJ", (6, 58, 10, 8, 18, 18, 60, 12, 12, 12)):
        wf.column_dimensions[col].width = wd
    for ws in (wf, wk):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(a.out)
    print(f"Zapisano {a.out}: formularz {len(wiersze)} poz., kalkulacja {len(praca)} + {len(zew)} pozycji; "
          f"B netto {zl(R['B'])}. Sprawdź przeliczenie formuł (LibreOffice) i porównaj z `licz`.")


# ----------------------------------------------------------------------------- MD
def cmd_md(a):
    s = wczytaj(a.spec, a.baza)
    R = oblicz(s, a.cel, a.cel_brutto)
    P = s["parametry"]
    vat = 1 + P["vat"]
    L = []
    A = L.append
    A(f"# Wycena: {s.get('zamawiajacy', '')} — „{s.get('tytul', '')}”" + (f" ({s['znak']})" if s.get("znak") else ""))
    A("")
    A(f"Status: {date.today().strftime('%d.%m.%Y')} — **wycena wstępna Claude, do decyzji Marcina.** Kalkulacja bazowa: **{zl(R['kalk'])} netto** "
      f"({zl0(R['rbg'])} rbg × {zl0(P['stawka'])} zł + {zl(sum(x['koszt'] for x in R['pozycje'] if x['rodzaj'] == 'zewn'))} kosztów zewnętrznych). "
      f"Wariant rekomendowany **{zl(R['B'])} netto / {zl(R['B'] * vat)} brutto**, próg bólu {zl(R['C'])} netto.")
    A("")
    A("## Dane postępowania\n")
    for k, n in (("zamawiajacy", "Zamawiający"), ("znak", "Znak"), ("platforma", "Platforma"), ("termin_ofert", "Termin ofert"),
                 ("termin_realizacji", "Termin realizacji"), ("zwiazanie", "Związanie ofertą"), ("kryteria", "Kryteria")):
        if s.get(k):
            A(f"- {n}: {s[k]}")
    wad = s.get("wadium") or 0
    A(f"- Wadium: " + (f"{zl(wad)} → " + ("przelew na rachunek zamawiającego (< 10 tys. zł)" if wad < 10000 else "gwarancja wadialna (≥ 10 tys. zł)") if wad else "brak"))
    for d in s.get("dane") or []:
        A(f"- {d}")
    A(f"\n## Kalkulacja (stawka {zl0(P['stawka'])} zł/rbg)\n")
    A("### Praca własna\n")
    A("| Poz. | Zakres | rbg | Netto [zł] | Formularz |\n|---|---|---:|---:|---|")
    for x in [x for x in R["pozycje"] if x["rodzaj"] == "praca"]:
        il = f"{zl0(x['jnp'])} jnp" if x.get("jnp") is not None else zl0(x["rbg"])
        A(f"| {x['id']} | {x.get('zakres', '')}{' **[sztywna]**' if x['sztywna'] else ''} | {il} | {zl0(x['koszt'])} | {', '.join(udzialy(x))} |")
    A(f"| | **Razem praca własna** | **{zl0(R['rbg'])}** | **{zl0(sum(x['koszt'] for x in R['pozycje'] if x['rodzaj'] == 'praca'))}** | |")
    zew = [x for x in R["pozycje"] if x["rodzaj"] == "zewn"]
    if zew:
        A("\n### Koszty zewnętrzne\n")
        A("| Pozycja | Kwota [zł] | Narzut | Netto [zł] | Formularz |\n|---|---:|---:|---:|---|")
        for x in zew:
            A(f"| {x.get('zakres', '')} | {zl0(x['kwota'])} | {proc(x['narzut_eff'], 0)} | {zl0(x['koszt'])} | {', '.join(udzialy(x))} |")
        A(f"| **Razem** | | | **{zl0(sum(x['koszt'] for x in zew))}** | |")
    A(f"\n**Koszt kalkulacyjny: {zl(R['kalk'])} netto ({zl(R['kalk'] * vat)} brutto).**" +
      (f" Formuła P&B: × {lp(R['wsp'])}." if R["wsp"] != 1 else "") +
      (f" Pozycje procentowe formularza (np. kwota tymczasowa): +{proc(R['procent'])}." if R["procent"] else ""))
    A("\n### Warianty\n")
    A("| Wariant | Netto | Brutto (" + proc(P["vat"], 0) + ") | Komentarz |\n|---|---:|---:|---|")
    A(f"| A — pełna kalkulacja | {zl0(R['A'])} | {zl0(R['A'] * vat)} | koszt + rezerwa {proc(P['rezerwa'], 0)} |")
    A(f"| **B — rekomendowany** | **{zl0(R['B'])}** | **{zl0(R['B'] * vat)}** | " +
      (f"cel cenowy; pozycje sztywne ×{lp(round(R['wsp_sztywne'], 3))}, miękkie ×{lp(round(R['wsp_miekkie'], 3))}" if R["cel"]
       else "koszt kalkulacyjny bez rezerwy; do kalibracji rynkowej") + " |")
    A(f"| C — próg bólu | {zl0(R['C'])} | {zl0(R['C'] * vat)} | {zl0(P['stawka_c'])} zł/rbg, koszty zewnętrzne po kosztach własnych, bez rezerwy |")
    if R["form"]["wiersze"]:
        A("\n### Rozbicie wariantu B na formularz zamawiającego\n")
        A("| Poz. | Wyszczególnienie | J.m. | Ilość | Cena jedn. [zł] | Wartość netto [zł] |\n|---|---|---|---:|---:|---:|")
        for w in R["form"]["wiersze"]:
            A(f"| {w['poz']} | {w['nazwa']}{' (opcja)' if w['opcja'] else ''} | {w['jm']} | {lp(w['ilosc'] or 1)} | {zl(w['cena_jedn'], False)} | {zl(w['wartosc'], False)} |")
        f = R["form"]
        A(f"| | **RAZEM netto** | | | | **{zl(f['netto'], False)}** |")
        A(f"| | VAT {proc(P['vat'], 0)} | | | | {zl(f['netto'] * P['vat'], False)} |")
        A(f"| | **RAZEM brutto** | | | | **{zl(f['netto'] * vat, False)}** |")
    ryn = s.get("rynek") or {}
    rt = rynek_tabela(s, R)
    if rt or ryn.get("opis"):
        A("\n## Benchmark rynkowy\n")
        for o in ryn.get("opis") or []:
            A(f"- {o}")
        if rt:
            A("\n" + rt)
    if s.get("ryzyka"):
        A("\n## Ryzyka\n")
        for i, t in enumerate(s["ryzyka"], 1):
            A(f"{i}. {t}")
    if s.get("uwagi"):
        A("\n## Uwagi do wyceny\n")
        for i, t in enumerate(s["uwagi"], 1):
            A(f"{i}. {t}")
    k = kontrole(s, R)
    if k:
        A("\n## Kontrole\n")
        for t in k:
            A(f"- {t}")
    if s.get("do_zrobienia"):
        A("\n## Do zrobienia przed terminem ofert\n")
        for i, t in enumerate(s["do_zrobienia"], 1):
            A(f"{i}. {t}")
    A("\n---\nDokument roboczy – wewnętrzny, nie do publikacji.")
    Path(a.out).write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"Zapisano {a.out} ({len(L)} wierszy).")


# ----------------------------------------------------------------------------- wzorce
def cmd_wzorce(a):
    kat = Path(a.baza or os.environ.get("MID_BAZA") or "baza_mid")
    p = kat / "wzorce_wycen.json"
    if not p.exists():
        sys.exit(f"Brak {p} - wzorce są w prywatnym repo mid-przetargi (baza_mid/). Ustaw --baza albo MID_BAZA.")
    W = json.loads(p.read_text(encoding="utf-8"))["wyceny"]
    if a.lista or not a.szukaj:
        print("| id | data | zamawiający | przedmiot | MiD netto | wynik |\n|---|---|---|---|---|---|")
        for w in W:
            print(f"| {w['id']} | {w.get('data', '')} | {w.get('zamawiajacy', '')} | {w.get('przedmiot', '')[:80]} | "
                  f"{zl0(w.get('oferta_mid_netto') or w.get('kalkulacja_netto'))} | {(w.get('wynik') or '')[:90]} |")
        return
    rx = re.compile(a.szukaj, re.I)
    print("| wycena | pozycja | rbg | kwota netto | kontekst |\n|---|---|---:|---:|---|")
    n = 0
    for w in W:
        for x in w.get("pozycje") or []:
            if rx.search(x.get("zakres", "")):
                n += 1
                kw = x.get("kwota") or (x["rbg"] * w["stawka"] if x.get("rbg") and w.get("stawka") else None)
                print(f"| {w['id']} | {x['zakres'][:110]} | {zl0(x.get('rbg')) if x.get('rbg') else ''} | {zl0(kw)} | {w.get('kontekst', '')[:70]} |")
    print(f"\n{n} pozycji. Kontekst (parametry obiektów, wynik przetargu) — `wzorce --lista`; kwoty netto z kalkulacji MiD, nie ceny rynkowe.")


def main():
    import signal
    if hasattr(signal, "SIGPIPE"):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("wzor"); s.add_argument("--out", default="wycena.json")
    for n in ("licz", "xlsx", "md"):
        s = sub.add_parser(n)
        s.add_argument("spec")
        g = s.add_mutually_exclusive_group()
        g.add_argument("--cel", type=float, help="cena oferty netto (cel Marcina)")
        g.add_argument("--cel-brutto", type=float)
        s.add_argument("--baza", help="katalog prywatnej bazy (parametry_wyceny.json); domyślnie MID_BAZA")
        if n != "licz":
            s.add_argument("--out", required=True)
    s = sub.add_parser("wzorce"); s.add_argument("--baza"); s.add_argument("--szukaj"); s.add_argument("--lista", action="store_true")
    a = ap.parse_args()
    {"wzor": cmd_wzor, "licz": cmd_licz, "xlsx": cmd_xlsx, "md": cmd_md, "wzorce": cmd_wzorce}[a.cmd](a)


if __name__ == "__main__":
    main()
