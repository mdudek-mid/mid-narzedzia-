#!/usr/bin/env python3
"""protokoly_tool.py - protokoly Pracowni MiD: rzeczowo-finansowe, przekazania dokumentacji, narady i nadzor autorski.

Protokol rzeczowo-finansowy (rozliczenia czesciowe wg harmonogramu, np. umowy z GW w P&B):
  rf-init HARMONOGRAM(.xlsx/.pdf/.docx/.csv/.txt/.json) --out projekt_rf.json [--nazwa --umowa --zamawiajacy --projektant --suma] [--nadpisz]
  rf-nowy projekt_rf.json --postep "1.1.1=100,1.1.2=50" [--data RRRR-MM-DD] [--okres 10.2026] [--zatrzymanie 5] [--uwagi ..]
          [--korekta] [--xlsx out.xlsx] [--md out.md]
          postep = zaawansowanie NARASTAJACO w % (stan po okresie); 1.12=+10 przyrost o 10 pkt %, 1.12=+1/12 ulamek (miesiac z 12)
  rf-xlsx projekt_rf.json --nr N --out out.xlsx    ponowne wygenerowanie protokolu nr N
  rf-stan projekt_rf.json                          stan rozliczen: zafakturowano, pozostalo, pozycje otwarte

Protokol przekazania dokumentacji:
  wzor-przekazanie --out przekazanie.json
  manifest KATALOG --out manifest.csv               lista plikow z rozmiarem i SHA-256 (dowod zawartosci przekazania)
  przekazanie przekazanie.json --out Protokol_przekazania.docx [--katalog DIR --manifest manifest.csv]

Protokol / notatka ze spotkania (narada, rada techniczna, pobyt nadzoru autorskiego):
  wzor-notatka [--typ narada|rada|nadzor|spotkanie] --out notatka.json
  transkrypt PLIK(.vtt/.docx/.txt) --out tekst.md   transkrypt (np. Teams) jako wypowiedzi mowcow, do redakcji notatki
  notatka notatka.json --out Notatka.docx [--md Notatka.md]
  zadania notatka1.json [notatka2.json ...] [--md zadania.md] [--na-dzien RRRR-MM-DD]   zbiorcza lista zadan, przeterminowane
"""
import argparse, csv, hashlib, json, re, sys
from datetime import date, timedelta
from pathlib import Path

FIRMA = "Pracownia Projektowa MiD Sp. z o.o."
RX_KWOTA = r"\d{1,3}(?:,\d{3})+\.\d{2}(?!\d)|\d{1,3}(?:[\u00a0 .]\d{3})+(?:,\d{2})?(?![\d,])|\d+,\d{2}|\d{4,}(?![\d,])"
JEDNOSTKI = r"ryczałt|ryczalt|kpl\.?|komplet|szt\.?|m-c|mies\.?|miesi[ąa]c|pobyt\w*|egz\.?|opracowanie|r-?g|rbg|ha|km|m"


def kw(s):
    if re.fullmatch(r"\d{1,3}(?:,\d{3})+\.\d{2}", s):     # zapis angielski 67,500.00
        return float(s.replace(",", ""))
    return float(re.sub(r"[  .]", "", s).replace(",", "."))


def fmt(v, n=2):
    if v is None:
        return "—"
    return f"{v:,.{n}f}".replace(",", " ").replace(".", ",")


def dzis():
    return date.today().isoformat()


# ============================================================================ RF
RX_NR = r"\d{1,2}(?:\.\d{1,2}){0,3}"


def parsuj_harmonogram(p):
    """-> (pozycje, suma_z_dokumentu, naglowki {nr: nazwa})"""
    p = Path(p)
    suf = p.suffix.lower()
    poz, nagl = [], {}
    if suf == ".json":
        d = json.loads(p.read_text(encoding="utf-8"))
        return (d.get("pozycje", d) if isinstance(d, dict) else d), None, {}
    if suf in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
            for hi, r in enumerate(rows):
                t = [str(c).lower().strip() if c is not None else "" for c in r]
                if not (any(re.match(r"^(poz|lp)", x) for x in t) and any(re.search(r"kwot|warto|cena", x) for x in t)):
                    continue
                ip = next(i for i, x in enumerate(t) if re.match(r"^(poz|lp)", x))
                ij = next((i for i, x in enumerate(t) if re.search(r"jedn|j\.\s?m", x)), None)
                kand = [i for i, x in enumerate(t) if re.search(r"kwot|warto|cena", x)]
                # kolumna kwot: wartosci liczbowe > 1 (nie procenty udzialu)
                def ocen(i):
                    v = [r2[i] for r2 in rows[hi + 1:] if i < len(r2) and isinstance(r2[i], (int, float))]
                    return sum(1 for x in v if abs(x) > 1)
                iv = max(kand, key=lambda i: (ocen(i), "kwot" in t[i]))
                suma = None
                for r2 in rows[hi + 1:]:
                    if ip >= len(r2):
                        continue
                    if r2[ip] is None:
                        if any(isinstance(c, str) and re.match(r"^\s*(suma|razem|ogółem)", c, re.I) for c in r2) and isinstance(r2[iv], (int, float)):
                            suma = float(r2[iv])
                        continue
                    nr = str(r2[ip]).strip().rstrip(".")
                    if not re.fullmatch(r"\d+(\.\d+)*", nr):
                        continue
                    naz = next((" ".join(str(c).split()) for c in r2[ip + 1:] if isinstance(c, str) and len(c.strip()) > 2), "")
                    val = r2[iv] if iv < len(r2) and isinstance(r2[iv], (int, float)) else None
                    if val is None:
                        nagl[nr] = naz
                        continue
                    jm = str(r2[ij]).strip() if ij is not None and ij < len(r2) and r2[ij] else "ryczałt"
                    poz.append({"poz": nr, "nazwa": naz, "jm": jm, "ilosc": 1.0, "wartosc": float(val)})
                if poz:
                    if any(isinstance(r2[ip], float) for r2 in rows[hi + 1:] if ip < len(r2)):
                        print("UWAGA: numery pozycji zapisane jako liczby (1.10 = 1.1) – sprawdź numerację.", file=sys.stderr)
                    return poz, suma, nagl
        sys.exit("Nie znaleziono tabeli z kolumnami „poz.” i „kwota/wartość”.")
    if suf == ".pdf":
        import subprocess
        tekst = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout
    elif suf == ".docx":
        import docx
        dd = docx.Document(str(p))
        tekst = "\n".join(x.text for x in dd.paragraphs)
        for t in dd.tables:
            for row in t.rows:
                kom = []
                for c in row.cells:
                    if not kom or c.text != kom[-1]:
                        kom.append(c.text.strip())
                tekst += "\n" + "\t".join(kom)
    elif suf in (".csv", ".tsv"):
        surowy = p.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dial = csv.Sniffer().sniff(surowy[:4000], delimiters=";\t,")
        except csv.Error:
            dial = csv.excel
            dial.delimiter = ";"
        tekst = "\n".join("\t".join(r) for r in csv.reader(surowy.splitlines(), dial))
    else:
        tekst = p.read_text(encoding="utf-8", errors="replace")
    plaski = re.sub(r"\s+", " ", tekst)
    # nazwa nie moze zawierac kolejnego numeru pozycji (np. „1. Prace projektowe 1.1 Złożenie…”)
    rx = re.compile(rf"(?<![\d,.])({RX_NR})\.?\s+(?:((?:(?!\s\d{{1,2}}(?:\.\d{{1,2}})+\.?\s)[^\n]){{3,250}}?)\s+)?({JEDNOSTKI})\s+({RX_KWOTA})", re.I)
    koniec = None
    for m in rx.finditer(plaski):
        if m.group(2) and m.group(2)[0].isdigit():
            continue
        nazwa = " ".join((m.group(2) or "").split())
        if koniec is not None:
            # PDF: nazwa zawinieta w komorce - fragmenty miedzy pozycjami (mala litera = koniec poprzedniej, wielka = poczatek biezacej)
            luka = re.sub(r"\d+(?:,\d+)?\s*%|(?<!\S)[\d\u00a0 .,]+(?!\S)", " ", plaski[koniec:m.start()]).split()
            if luka and not re.search(rf"(?<!\S){RX_NR}\.?(?!\S)", plaski[koniec:m.start()]):
                i = 0
                while i < len(luka) and (luka[i][0].islower() or not luka[i][0].isalnum()):
                    i += 1
                if i and poz:
                    poz[-1]["nazwa"] += " " + " ".join(luka[:i])
                if luka[i:]:
                    nazwa = (" ".join(luka[i:]) + " " + nazwa).strip()
        koniec = m.end()
        poz.append({"poz": m.group(1), "nazwa": nazwa, "jm": m.group(3), "ilosc": 1.0, "wartosc": kw(m.group(4))})
    # naglowki grup: linia „1. Prace projektowe” bez kwoty
    for l in tekst.splitlines():
        m = re.match(rf"^\s*({RX_NR})\.?\s+([^\d\s][^\t]{{2,120}}?)\s*$", l)
        if m and not re.search(RX_KWOTA, l) and not re.search(rf"\s({JEDNOSTKI})\s+\d", l, re.I):
            nagl.setdefault(m.group(1), m.group(2).strip())
    suma = re.search(rf"(?:SUMA|RAZEM|OGÓŁEM)\s*(?:netto)?\s*:?\s*({RX_KWOTA})", plaski, re.I)
    return poz, (kw(suma.group(1)) if suma else None), nagl


def cmd_rf_init(a):
    if Path(a.out).exists() and not a.nadpisz:
        sys.exit(f"{a.out} już istnieje (zawiera historię protokołów). Podaj inną nazwę albo --nadpisz.")
    poz, suma, nagl = parsuj_harmonogram(a.harmonogram)
    if not poz:
        sys.exit("Nie rozpoznano pozycji harmonogramu (poz. | nazwa | jednostka | kwota).")
    nr_y = [p["poz"] for p in poz]
    dup = sorted({n for n in nr_y if nr_y.count(n) > 1})
    if dup:
        print(f"UWAGA: powtórzone numery pozycji {', '.join(dup)} – popraw numerację w pliku JSON.")
    # pozycja nadrzedna = suma podpozycji -> podsuma (nie rozliczana osobno)
    for p in poz:
        dzieci = [q for q in poz if q["poz"].startswith(p["poz"] + ".") and q["poz"].count(".") == p["poz"].count(".") + 1]
        if dzieci and p.get("wartosc") is not None and abs(sum(q["wartosc"] or 0 for q in dzieci) - p["wartosc"]) < 0.5:
            p["podsuma"] = True
    # naglowki grup (np. „1. Prace projektowe”) bez wartosci
    rodzice = set()
    for p in poz:
        cz = p["poz"].split(".")
        rodzice.update(".".join(cz[:i]) for i in range(1, len(cz)))
    for r in sorted(rodzice):
        if not any(p["poz"] == r for p in poz):
            poz.append({"poz": r, "nazwa": nagl.get(r, ""), "jm": "", "ilosc": None, "wartosc": None, "naglowek": True})
    poz.sort(key=lambda p: [int(x) for x in p["poz"].split(".")])
    razem = sum(p["wartosc"] or 0 for p in poz if not p.get("podsuma"))
    st = {"projekt": a.nazwa or "", "umowa": a.umowa or "", "zamawiajacy": a.zamawiajacy or "", "projektant": a.projektant or FIRMA,
          "pozycje": poz, "suma_umowna": a.suma or suma or razem, "protokoly": [], "zrodlo": Path(a.harmonogram).name}
    Path(a.out).write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
    bez = [p["poz"] for p in poz if not p.get("naglowek") and not p.get("nazwa")]
    if bez:
        print(f"UWAGA: pozycje bez nazwy {', '.join(bez)} – uzupełnij nazwy w pliku JSON (układ PDF).")
    n = sum(1 for p in poz if p.get("wartosc") is not None and not p.get("podsuma"))
    print(f"Zapisano {a.out}: {n} pozycji rozliczeniowych, razem {fmt(razem)} zł netto.")
    if st["suma_umowna"] and abs(st["suma_umowna"] - razem) > 0.5:
        print(f"UWAGA: suma pozycji {fmt(razem)} ≠ suma w harmonogramie {fmt(st['suma_umowna'])} – sprawdź, czy nie brakuje pozycji "
              "albo czy pozycje nadrzędne nie dublują podpozycji.")
    elif a.suma or suma:
        print(f"Suma zgodna z {'podaną kwotą (--suma)' if a.suma else 'sumą w harmonogramie'}: {fmt(st['suma_umowna'])} zł.")
    else:
        print("UWAGA: w harmonogramie nie znaleziono sumy (SUMA/RAZEM) – porównaj wynik z wynagrodzeniem z umowy albo podaj --suma.")
    for p in poz:
        if p.get("naglowek"):
            print(f"  {p['poz'] + '.':<8} {'':>14}  {p['nazwa'][:80] or '(nagłówek grupy)'}")
        else:
            print(f"  {p['poz']:<8} {fmt(p['wartosc']):>14}  {p['nazwa'][:80]}{'  [podsuma]' if p.get('podsuma') else ''}")


def rozliczeniowe(st):
    return [p for p in st["pozycje"] if p.get("wartosc") is not None and not p.get("naglowek") and not p.get("podsuma")]


def stan_przed(st, nr=None):
    """narastajace zaawansowanie [%] po protokole nr (None = po ostatnim)"""
    s = {}
    for pr in st["protokoly"]:
        if nr is not None and pr["nr"] > nr:
            break
        s.update(pr["narastajaco"])
    return s


def cmd_rf_nowy(a):
    st = json.loads(Path(a.projekt).read_text(encoding="utf-8"))
    pozy = {p["poz"]: p for p in rozliczeniowe(st)}
    podsumy = {p["poz"] for p in st["pozycje"] if p.get("podsuma")}
    poprz = stan_przed(st)
    nowy = dict(poprz)
    bledy = []
    for el in [x.strip() for x in a.postep.split(",") if x.strip()]:
        if "=" not in el:
            bledy.append(f"„{el}” – oczekiwano poz=procent, np. 1.2=100, 1.12=+10 (przyrost) albo 1.12=+1/12 (ułamek)")
            continue
        k, v = el.split("=", 1)
        k, v = k.strip().rstrip("."), v.strip().replace(",", ".").rstrip("%")
        przyrost = v.startswith("+")
        try:
            if "/" in v:                       # ulamek, np. +1/12 (miesiac nadzoru z 12)
                l_, m_ = v.lstrip("+").split("/")
                v = 100 * float(l_) / float(m_)
            else:
                v = float(v)
        except (ValueError, ZeroDivisionError):
            bledy.append(f"{k}: „{v}” nie jest liczbą")
            continue
        if k in podsumy:
            bledy.append(f"{k} to suma podpozycji – podaj postęp podpozycji {k}.x")
            continue
        if k not in pozy:
            bledy.append(f"brak pozycji {k} w harmonogramie")
            continue
        if przyrost:
            v = poprz.get(k, 0) + v
        if v > 100 + 1e-9:
            bledy.append(f"{k}: {v}% > 100%")
        if v < poprz.get(k, 0) - 1e-9 and not a.korekta:
            bledy.append(f"{k}: {v}% mniej niż w poprzednim protokole ({poprz.get(k, 0)}%) – użyj --korekta, jeśli to świadoma korekta")
        nowy[k] = v
    if bledy:
        sys.exit("Błędy:\n - " + "\n - ".join(bledy))
    if a.data:
        try:
            date.fromisoformat(a.data)
        except ValueError:
            sys.exit(f"--data {a.data}: oczekiwano RRRR-MM-DD")
        if st["protokoly"] and a.data < st["protokoly"][-1]["data"]:
            print(f"UWAGA: data {a.data} wcześniejsza niż protokołu nr {st['protokoly'][-1]['nr']} ({st['protokoly'][-1]['data']}).")
    nr = (st["protokoly"][-1]["nr"] + 1) if st["protokoly"] else 1
    okres = sum(pozy[k]["wartosc"] * (nowy.get(k, 0) - poprz.get(k, 0)) / 100 for k in pozy)
    pr = {"nr": nr, "data": a.data or dzis(), "okres": a.okres or "", "narastajaco": nowy, "uwagi": a.uwagi or "",
          "wartosc_okresu": round(okres, 2), "zatrzymanie_proc": a.zatrzymanie}
    st["protokoly"].append(pr)
    Path(a.projekt).write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
    out = a.xlsx or f"Protokol_RF_nr{nr}.xlsx"
    rf_xlsx(st, nr, out)
    md = rf_md(st, nr)
    if a.md:
        Path(a.md).write_text(md, encoding="utf-8")
    print(md)
    print(f"\nZapisano {out} i zaktualizowano {a.projekt}.")


def rf_wiersze(st, nr):
    pozy = [p for p in st["pozycje"]]
    rozl = {p["poz"] for p in rozliczeniowe(st)}
    po = stan_przed(st, nr)
    pr = stan_przed(st, nr - 1) if nr > 1 else {}
    w = []
    for p in pozy:
        if p["poz"] not in rozl:
            w.append({**p, "naglowek": True})
            continue
        k = p["poz"]
        w.append({**p, "pop": pr.get(k, 0) / 100, "nar": po.get(k, 0) / 100, "okr": (po.get(k, 0) - pr.get(k, 0)) / 100})
    return w


def rf_md(st, nr):
    pr = next(x for x in st["protokoly"] if x["nr"] == nr)
    w = rf_wiersze(st, nr)
    suma = sum(x["wartosc"] for x in w if not x.get("naglowek"))
    nar = sum(x["wartosc"] * x["nar"] for x in w if not x.get("naglowek"))
    okr = sum(x["wartosc"] * x["okr"] for x in w if not x.get("naglowek"))
    L = [f"# Protokół rzeczowo-finansowy nr {nr} — {st['projekt']}", "",
         f"Umowa: {st['umowa'] or '—'}; data: {pr['data']}; okres: {pr['okres'] or '—'}.", "",
         "| Poz. | Element | Wartość | Narastająco | Poprzednio | W okresie [zł] |", "|---|---|---|---|---|---|"]
    wczesniej = []
    for x in w:
        if x.get("naglowek"):
            continue
        if abs(x["okr"]) > 1e-12:
            L.append(f"| {x['poz']} | {x['nazwa'][:70]} | {fmt(x['wartosc'])} | {fmt(100 * x['nar'], 1)}% | {fmt(100 * x['pop'], 1)}% | {fmt(x['wartosc'] * x['okr'])} |")
        elif x["nar"]:
            wczesniej.append(f"{x['poz']} ({fmt(100 * x['nar'], 0)}%)")
    if wczesniej:
        L += ["", "Rozliczone wcześniej, bez zmian w okresie: " + ", ".join(wczesniej) + "."]
    L += ["", f"Wartość umowna: **{fmt(suma)} zł** netto. Zaawansowanie narastająco: {fmt(nar)} zł ({fmt(100 * nar / suma, 1)}%).",
          f"**Do faktury za okres: {fmt(okr)} zł netto**, VAT 23%: {fmt(okr * 0.23)} zł, brutto: {fmt(okr * 1.23)} zł."]
    if pr.get("zatrzymanie_proc"):
        z = okr * pr["zatrzymanie_proc"] / 100
        L.append(f"Zatrzymanie (kaucja) {fmt(pr['zatrzymanie_proc'], 1)}%: {fmt(z)} zł – do wypłaty netto {fmt(okr - z)} zł (wg umowy: od netto czy brutto – sprawdź).")
    if abs(okr) < 0.005:
        L.append("UWAGA: wartość okresu = 0 – brak postępu do rozliczenia.")
    return "\n".join(L)


def rf_xlsx(st, nr, out):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    pr = next(x for x in st["protokoly"] if x["nr"] == nr)
    w = rf_wiersze(st, nr)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"RF nr {nr}"
    A = Font(name="Arial", size=9)
    B = Font(name="Arial", size=9, bold=True)
    H = PatternFill("solid", fgColor="D9E1F2")
    cien = Side(style="thin", color="808080")
    ramka = Border(left=cien, right=cien, top=cien, bottom=cien)
    zawin = Alignment(wrap_text=True, vertical="center")
    szer = {"A": 7, "B": 44, "C": 9, "D": 8, "E": 13, "F": 14, "G": 9, "H": 14, "I": 9, "J": 14, "K": 9, "L": 14, "M": 9, "N": 14}
    for k, v in szer.items():
        ws.column_dimensions[k].width = v
    ws["A1"] = f"Protokół Rzeczowo-Finansowy nr {nr}"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"] = f"„{st['projekt']}”" + (f"  Umowa nr {st['umowa']}" if st.get("umowa") else "")
    ws["A3"] = f"Protokół sporządzono w dniu {pr['data']}" + (f"; okres rozliczeniowy: {pr['okres']}" if pr.get("okres") else "")
    ws["A4"] = f"Zamawiający: {st.get('zamawiajacy') or '……………'}"
    ws["H4"] = f"Projektant: {st.get('projektant') or FIRMA}"
    for c in ("A2", "A3", "A4", "H4"):
        ws[c].font = A
    ws["A6"] = "Strony stwierdzają następujące zaawansowanie prac:"
    ws["A6"].font = B
    nag1 = ["Poz.", "Element prac", "Jednostka", "Ilość umowna", "Cena jednostkowa", "Wartość wg umowy", "Wartość od początku realizacji", "",
            "Wartość wg poprzedniego protokołu", "", "Wartość w okresie rozliczeniowym", "", "Wartość prac pozostałych do wykonania", ""]
    nag2 = ["", "", "", "", "zł", "zł", "ilość / %", "zł", "ilość / %", "zł", "ilość / %", "zł", "ilość / %", "zł"]
    for j, (t1, t2) in enumerate(zip(nag1, nag2), 1):
        for rr, t in ((8, t1), (9, t2)):
            c = ws.cell(rr, j, t)
            c.font = B
            c.fill = H
            c.border = ramka
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for a_, b_ in (("G8", "H8"), ("I8", "J8"), ("K8", "L8"), ("M8", "N8")):
        ws.merge_cells(f"{a_}:{b_}")
    ws.row_dimensions[8].height = 38
    r0 = 10
    r = r0
    wiersze_poz = []          # (wiersz, poz, naglowek?)
    for x in w:
        wiersze_poz.append((r, x["poz"], bool(x.get("naglowek"))))
        if x.get("naglowek"):
            ws.cell(r, 1, x["poz"] + ("." if "." not in x["poz"] else "")).font = B
            opis = x.get("nazwa") or ""
            if x.get("podsuma"):
                opis += f" (suma podpozycji: {fmt(x['wartosc'])} zł)"
            ws.cell(r, 2, opis).font = B
            ws.cell(r, 2).alignment = zawin
            for j in range(1, 15):
                ws.cell(r, j).border = ramka
            r += 1
            continue
        vals = [x["poz"], x["nazwa"], x.get("jm") or "ryczałt", x.get("ilosc") or 1.0, x["wartosc"] / (x.get("ilosc") or 1.0)]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)
        ws.cell(r, 6, f"=D{r}*E{r}")
        ws.cell(r, 9, round(x["pop"] * (x.get("ilosc") or 1.0), 6))
        ws.cell(r, 11, round(x["okr"] * (x.get("ilosc") or 1.0), 6))
        ws.cell(r, 7, f"=I{r}+K{r}")
        ws.cell(r, 8, f"=G{r}*E{r}")
        ws.cell(r, 10, f"=I{r}*E{r}")
        ws.cell(r, 12, f"=K{r}*E{r}")
        ws.cell(r, 13, f"=D{r}-G{r}")
        ws.cell(r, 14, f"=F{r}-H{r}")
        for j in range(1, 15):
            c = ws.cell(r, j)
            c.font = A
            c.border = ramka
            if j == 2:
                c.alignment = zawin
            if j in (5, 6, 8, 10, 12, 14):
                c.number_format = '#,##0.00'
            if j in (4, 7, 9, 11, 13):
                c.number_format = '0.00%' if (x.get("ilosc") or 1.0) == 1.0 else '#,##0.00'
        if (x.get("ilosc") or 1.0) == 1.0:
            for j in (7, 9, 11, 13):
                ws.cell(r, j).number_format = '0.0%'
            ws.cell(r, 4).number_format = '0.00'
        # wejscie uzytkownika: kolumna K (w okresie) i I (poprzednio) - niebieskie; pozycje rozliczane w okresie - tlo
        for j in (9, 11):
            ws.cell(r, j).font = Font(name="Arial", size=9, color="1F4E79", bold=True)
        if abs(x["okr"]) > 1e-12:
            for j in range(1, 15):
                ws.cell(r, j).fill = PatternFill("solid", fgColor="FFF2CC")
        r += 1
    # sumy grup w wierszach naglowkow (tylko pozycje rozliczeniowe: kolumna C niepusta)
    for i, (rr, pz, hd) in enumerate(wiersze_poz):
        if not hd:
            continue
        pot = [q for q in wiersze_poz[i + 1:] if q[1].startswith(pz + ".")]
        if not any(not q[2] for q in pot):
            continue
        a_, b_ = pot[0][0], pot[-1][0]
        for j, col in ((6, "F"), (8, "H"), (10, "J"), (12, "L"), (14, "N")):
            c = ws.cell(rr, j, f'=SUMIF($C${a_}:$C${b_},"<>",{col}{a_}:{col}{b_})')
            c.font = B
            c.number_format = '#,##0.00'
    ws.cell(r, 2, "Całkowita wartość prac:").font = B
    for j, col in ((6, "F"), (8, "H"), (10, "J"), (12, "L"), (14, "N")):
        c = ws.cell(r, j, f'=SUMIF($C${r0}:$C${r - 1},"<>",{col}{r0}:{col}{r - 1})')
        c.font = B
        c.number_format = '#,##0.00'
    for j, col in ((7, "H"), (9, "J"), (11, "L"), (13, "N")):
        c = ws.cell(r, j, f"=IF(F{r}=0,0,{col}{r}/F{r})")
        c.font = B
        c.number_format = '0.0%'
    for j in range(1, 15):
        ws.cell(r, j).border = ramka
        ws.cell(r, j).fill = H
    rs = r
    r += 2
    ws.cell(r, 2, "Do faktury za okres (netto):").font = B
    ws.cell(r, 6, f"=L{rs}").number_format = '#,##0.00'
    ws.cell(r + 1, 2, "VAT 23%:").font = A
    ws.cell(r + 1, 6, f"=ROUND(F{r}*0.23,2)").number_format = '#,##0.00'
    ws.cell(r + 2, 2, "Brutto:").font = B
    ws.cell(r + 2, 6, f"=F{r}+F{r + 1}").number_format = '#,##0.00'
    if pr.get("zatrzymanie_proc"):
        ws.cell(r + 3, 2, f"Zatrzymanie (kaucja) {pr['zatrzymanie_proc']}% od netto:").font = A
        ws.cell(r + 3, 6, f"=ROUND(F{r}*{pr['zatrzymanie_proc'] / 100},2)").number_format = '#,##0.00'
    r += 5
    ws.cell(r, 1, st.get("stopka") or "Niniejszy protokół jest częściowym rozliczeniem finansowym w oparciu o bieżące zaawansowanie prac "
            "i stanowi podstawę do wystawienia faktury. Protokół nie stanowi potwierdzenia odbioru technicznego i jakościowego prac.").font = Font(name="Arial", size=8, italic=True)
    if pr.get("uwagi"):
        ws.cell(r + 1, 1, f"Uwagi: {pr['uwagi']}").font = A
    r += 3
    ws.cell(r, 2, "Za Zamawiającego / Wykonawcę:").font = B
    ws.cell(r, 10, "Za Projektanta:").font = B
    for i in range(1, 3):
        ws.cell(r + i + 1, 2, f"{i}. ....................................  (data i podpis)").font = A
        ws.cell(r + i + 1, 10, f"{i}. ....................................  (data i podpis)").font = A
    ws.freeze_panes = "C10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.save(out)


def cmd_rf_xlsx(a):
    st = json.loads(Path(a.projekt).read_text(encoding="utf-8"))
    rf_xlsx(st, a.nr, a.out)
    print(rf_md(st, a.nr))
    print(f"\nZapisano {a.out}")


def cmd_rf_stan(a):
    st = json.loads(Path(a.projekt).read_text(encoding="utf-8"))
    s = stan_przed(st)
    pozy = rozliczeniowe(st)
    suma = sum(p["wartosc"] for p in pozy)
    zaf = sum(p["wartosc"] * s.get(p["poz"], 0) / 100 for p in pozy)
    print(f"# Stan rozliczeń — {st['projekt']}\n")
    print(f"Protokoły: {len(st['protokoly'])}; wartość umowna {fmt(suma)} zł; rozliczono {fmt(zaf)} zł ({fmt(100 * zaf / suma, 1)}%); pozostało {fmt(suma - zaf)} zł.\n")
    print("| Nr | Data | Okres | W okresie [zł] |\n|---|---|---|---|")
    for pr in st["protokoly"]:
        print(f"| {pr['nr']} | {pr['data']} | {pr['okres']} | {fmt(pr['wartosc_okresu'])} |")
    print("\n| Poz. | Element | Wartość | Narastająco | Pozostało [zł] |\n|---|---|---|---|---|")
    for p in pozy:
        v = s.get(p["poz"], 0)
        if v < 100:
            print(f"| {p['poz']} | {p['nazwa'][:60]} | {fmt(p['wartosc'])} | {fmt(v, 1)}% | {fmt(p['wartosc'] * (100 - v) / 100)} |")


# ============================================================================ DOCX - wspolne
def nowy_docx(tytul, podtytul=None, firma=None):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.1
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
        f = s.header.paragraphs[0]
        f.text = firma or FIRMA
        f.runs[0].font.size = Pt(8)
        f.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p = d.add_paragraph()
    r = p.add_run(tytul)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if podtytul:
        p = d.add_paragraph()
        p.add_run(podtytul).font.size = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return d


def tabela(d, dane, szer=None, naglowkowa=True, rozm=9):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    t = d.add_table(rows=len(dane), cols=len(dane[0]))
    t.style = "Table Grid"
    for i, w in enumerate(dane):
        for j, v in enumerate(w):
            c = t.cell(i, j)
            c.text = ""
            run = c.paragraphs[0].add_run("" if v is None else str(v))
            run.font.size = Pt(rozm)
            if naglowkowa and i == 0:
                run.bold = True
                sh = OxmlElement("w:shd")
                sh.set(qn("w:val"), "clear")
                sh.set(qn("w:fill"), "D9E1F2")
                c._tc.get_or_add_tcPr().append(sh)
    if szer:
        t.autofit = False
        for j, s in enumerate(szer):
            t.columns[j].width = Cm(s)
            for row in t.rows:
                row.cells[j].width = Cm(s)
    d.add_paragraph()
    return t


def naglowek(d, t):
    from docx.shared import Pt, RGBColor
    p = d.add_paragraph()
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    return p


def podpisy(d, lewa, prawa, n=2, osoby_l=None, osoby_p=None):
    """tabela podpisow bez ramek; osoby_* - nazwiska (puste/[..] = linia do uzupelnienia)"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    def linia(lst, i):
        o = (lst or [])[i - 1] if lst and i <= len(lst) else ""
        o = "" if not o or o.startswith("[") else o
        return f"{i}. {o}\n\n..................................................\n(data i podpis)"
    wiersze = [[lewa, prawa]] + [[linia(osoby_l, i), linia(osoby_p, i)] for i in range(1, n + 1)]
    t = tabela(d, wiersze, szer=[8.5, 8.5], naglowkowa=False, rozm=9.5)
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
    for c in t.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    # podpisy nie rozdzielane miedzy strony
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit")
        trPr.append(e)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    if len(d.paragraphs) >= 2:
        for p in d.paragraphs[-2:]:
            p.paragraph_format.keep_with_next = True
    # usun pusty akapit po tabeli (pusta ostatnia strona)
    ost = d.paragraphs[-1]
    if not ost.text.strip():
        ost._element.getparent().remove(ost._element)
    return t


def data_pl(s):
    m = ["stycznia", "lutego", "marca", "kwietnia", "maja", "czerwca", "lipca", "sierpnia", "września", "października", "listopada", "grudnia"]
    try:
        d = date.fromisoformat(s)
        return f"{d.day} {m[d.month - 1]} {d.year} r."
    except Exception:
        return s or "……………"


# ============================================================================ przekazanie
def manifest(katalog):
    k = Path(katalog)
    wyn = []
    for p in sorted(k.rglob("*")):
        if p.is_file() and not p.name.startswith("."):
            h = hashlib.sha256()
            with open(p, "rb") as f:
                for blok in iter(lambda: f.read(1 << 20), b""):
                    h.update(blok)
            wyn.append({"plik": str(p.relative_to(k)), "rozmiar_B": p.stat().st_size, "sha256": h.hexdigest()})
    return wyn


def cmd_manifest(a):
    m = manifest(a.katalog)
    with open(a.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["plik", "rozmiar_B", "sha256"], delimiter=";")
        w.writeheader()
        w.writerows(m)
    tot = sum(x["rozmiar_B"] for x in m)
    calosc = hashlib.sha256("".join(x["sha256"] for x in m).encode()).hexdigest()
    print(f"Zapisano {a.out}: {len(m)} plików, {tot / 1e6:.1f} MB; skrót zbiorczy SHA-256: {calosc}")


def cmd_wzor_przekazanie(a):
    spec = {
        "numer": "", "data": dzis(), "miejsce": "Gdańsk", "umowa": {"nr": "[nr umowy]", "data": "RRRR-MM-DD"},
        "temat": "[nazwa zadania z umowy]",
        "przekazujacy": {"nazwa": FIRMA, "adres": "", "osoby": ["[imię i nazwisko, funkcja]"]},
        "odbierajacy": {"nazwa": "[Zamawiający]", "adres": "", "osoby": ["[imię i nazwisko]"]},
        "cel": "celem sprawdzenia zgodności z zapisami umowy",
        "etap": "[np. Projekt budowlany – etap 2 harmonogramu]",
        "pozycje": [{"nazwa": "Tom I – Projekt zagospodarowania terenu", "branza": "mostowa", "papier_egz": 4, "elektronicznie": "PDF, DWG", "uwagi": ""},
                    {"nazwa": "Tom II – Projekt architektoniczno-budowlany", "branza": "mostowa", "papier_egz": 4, "elektronicznie": "PDF, DWG, DOCX", "uwagi": ""}],
        "nosnik": "[pendrive / płyta DVD / link do serwera – bez haseł w protokole]",
        "termin_sprawdzenia": "", "faktura": "Wykonawca wystawi fakturę po podpisaniu protokołu odbioru.",
        "oswiadczenie": "Wykonawca oświadcza, że przekazana dokumentacja została wykonana zgodnie z umową, obowiązującymi przepisami, Polskimi Normami i zasadami wiedzy technicznej oraz jest kompletna z punktu widzenia celu, któremu ma służyć.",
        "uwagi": "",
    }
    Path(a.out).write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Zapisano {a.out}")


def cmd_przekazanie(a):
    sp = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    u = sp.get("umowa", {})
    d = nowy_docx(f"PROTOKÓŁ PRZEKAZANIA DOKUMENTACJI{(' nr ' + sp['numer']) if sp.get('numer') else ''}",
                  f"{sp.get('miejsce') or '……………'}, dnia {data_pl(sp.get('data'))}", sp.get("firma"))
    tabela(d, [["Umowa:", f"nr {u.get('nr', '……')} z dnia {data_pl(u.get('data'))}"], ["Zadanie:", sp.get("temat", "")],
               ["Etap / część:", sp.get("etap", "—")],
               ["Strona przekazująca:", "; ".join(x for x in [sp["przekazujacy"].get("nazwa"), sp["przekazujacy"].get("adres")] if x)],
               ["Strona odbierająca:", "; ".join(x for x in [sp["odbierajacy"].get("nazwa"), sp["odbierajacy"].get("adres")] if x)]],
           szer=[4, 13], naglowkowa=False)
    d.add_paragraph(f"Strona przekazująca przekazuje, a strona odbierająca przyjmuje niżej wymienioną dokumentację {sp.get('cel', '')}.".replace("  ", " "))
    naglowek(d, "Przedmiot przekazania")
    w = [["Lp.", "Nazwa opracowania (tom / część)", "Branża", "Papier [egz.]", "Wersja elektroniczna", "Uwagi"]]
    egz = 0
    for i, x in enumerate(sp.get("pozycje", []), 1):
        w.append([i, x.get("nazwa", ""), x.get("branza", ""), x.get("papier_egz", "—"), x.get("elektronicznie", "—"), x.get("uwagi", "")])
        egz += int(x.get("papier_egz") or 0)
    tabela(d, w, szer=[1, 6.5, 2.2, 1.8, 3, 2.5])
    m = None
    if a.katalog:
        m = manifest(a.katalog)
        sp.setdefault("nosnik", "")
        calosc = hashlib.sha256("".join(x["sha256"] for x in m).encode()).hexdigest()
        naglowek(d, "Wersja elektroniczna – wykaz plików")
        d.add_paragraph(f"Nośnik / sposób przekazania: {sp.get('nosnik') or '……'}. Liczba plików: {len(m)}, łączny rozmiar: "
                        f"{fmt(sum(x['rozmiar_B'] for x in m) / 1e6, 1)} MB. Skrót zbiorczy SHA-256 (z sum kontrolnych plików w kolejności wykazu): {calosc}.")
        if len(m) <= 60:
            tabela(d, [["Lp.", "Plik", "Rozmiar [kB]", "SHA-256 (pierwsze 16 znaków)"]] +
                   [[i, x["plik"], f"{x['rozmiar_B'] / 1024:.0f}", x["sha256"][:16]] for i, x in enumerate(m, 1)], szer=[1, 9.5, 2, 4.5], rozm=8)
        else:
            d.add_paragraph("Pełny wykaz plików z sumami kontrolnymi stanowi załącznik (manifest CSV).")
    elif sp.get("nosnik"):
        d.add_paragraph(f"Nośnik / sposób przekazania wersji elektronicznej: {sp['nosnik']}.")
    naglowek(d, "Ustalenia")
    pkt = []
    if sp.get("oswiadczenie"):
        pkt.append(sp["oswiadczenie"])
    if sp.get("termin_sprawdzenia"):
        t = data_pl(sp["termin_sprawdzenia"])
        pkt.append(f"Zamawiający sprawdzi przekazaną dokumentację w terminie {'do dnia ' if t[:1].isdigit() else ''}{t}{'' if t.endswith('.') else '.'}")
    if sp.get("faktura"):
        pkt.append(sp["faktura"])
    if sp.get("uwagi"):
        pkt.append(sp["uwagi"])
    pkt.append("Protokół przekazania potwierdza wyłącznie fakt przekazania dokumentacji i nie stanowi jej odbioru.")
    for i, t in enumerate(pkt, 1):
        d.add_paragraph(f"{i}. {t}")
    d.add_paragraph(f"Protokół sporządzono w dwóch jednobrzmiących egzemplarzach, po jednym dla każdej ze stron. Łącznie przekazano {egz} egz. papierowych.")
    naglowek(d, "Podpisy")
    podpisy(d, "Strona przekazująca:", "Strona odbierająca:", n=max(len(sp["przekazujacy"].get("osoby", [])), len(sp["odbierajacy"].get("osoby", [])), 1),
            osoby_l=sp["przekazujacy"].get("osoby"), osoby_p=sp["odbierajacy"].get("osoby"))
    d.save(a.out)
    print(f"Zapisano {a.out}: {len(sp.get('pozycje', []))} pozycji" + (f", manifest {len(m)} plików" if m else ""))
    if a.katalog and a.manifest:
        with open(a.manifest, "w", newline="", encoding="utf-8") as f:
            wr = csv.DictWriter(f, fieldnames=["plik", "rozmiar_B", "sha256"], delimiter=";")
            wr.writeheader()
            wr.writerows(m)
        print(f"Zapisano manifest {a.manifest}")


# ============================================================================ notatki
TYPY = {"narada": "PROTOKÓŁ Z NARADY", "rada": "PROTOKÓŁ Z RADY TECHNICZNEJ", "nadzor": "NOTATKA Z POBYTU NADZORU AUTORSKIEGO",
        "spotkanie": "NOTATKA ZE SPOTKANIA"}


def cmd_wzor_notatka(a):
    spec = {"typ": a.typ, "numer": "", "temat": "[zadanie / obiekt]", "data": dzis(), "godzina": "10:00", "miejsce": "[miejsce / Teams]",
            "prowadzacy": "[imię i nazwisko]", "protokolant": "[imię i nazwisko]",
            "uczestnicy": [{"osoba": "[imię i nazwisko]", "firma": "[firma]", "funkcja": "[funkcja]"}],
            "porzadek": ["[punkt porządku]"],
            "ustalenia": [{"punkt": 1, "tresc": "[ustalenie]"}],
            "decyzje": ["[decyzja z uzasadnieniem]"],
            "zadania": [{"co": "[zadanie]", "kto": "[osoba / firma]", "termin": "RRRR-MM-DD", "status": "otwarte"}],
            "sprawy_otwarte": ["[sprawa wymagająca wyjaśnienia]"],
            "nastepne_spotkanie": "", "zalaczniki": [], "rozdzielnik": []}
    if a.typ == "nadzor":
        spec.update({"pobyt_nr": 1, "budowa": "[obiekt, adres budowy]", "inwestor": "", "wykonawca_robot": "", "kierownik_budowy": "",
                     "inspektor_nadzoru": "", "wpis_dziennik": "[nr wpisu w dzienniku budowy / brak]",
                     "zgodnosc_z_projektem": "[stwierdzenia w zakresie zgodności realizacji z projektem]",
                     "decyzja": "[pozwolenie na budowę / ZRID – nr i data]",
                     "odstapienia": [{"opis": "[odstąpienie / rozwiązanie zamienne zgłoszone przez kierownika budowy lub inspektora]",
                                      "projekt": "PB", "kwalifikacja": "nieistotne",
                                      "uzasadnienie": "[art. 36a ust. 5 – nie dotyczy żadnej z przesłanek pkt 1–7]",
                                      "zmiana_wymiaru_proc": None, "zmiana_pow_zabudowy_proc": None,
                                      "dokumentacja": "[rysunek zamienny nr … i opis do dokumentacji budowy]"}]})
    Path(a.out).write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Zapisano {a.out}")


def kontrola_notatki(sp, na_dzien=None):
    uw = []
    nd = date.fromisoformat(na_dzien) if na_dzien else date.today()
    for i, z in enumerate(sp.get("zadania", []), 1):
        if not z.get("kto") or z["kto"].startswith("["):
            uw.append(f"zadanie {i} bez osoby odpowiedzialnej: {z.get('co', '')[:60]}")
        if not z.get("termin") or z["termin"].startswith("R"):
            uw.append(f"zadanie {i} bez terminu: {z.get('co', '')[:60]}")
        else:
            try:
                if date.fromisoformat(z["termin"]) < nd and z.get("status", "otwarte") != "zamknięte":
                    uw.append(f"zadanie {i} po terminie ({z['termin']}): {z.get('co', '')[:60]}")
            except ValueError:
                uw.append(f"zadanie {i}: termin „{z['termin']}” nie jest datą RRRR-MM-DD")
    if sp.get("typ") == "nadzor":
        for o in sp.get("odstapienia", []):
            op = o.get("opis", "")[:60]
            if str(o.get("projekt", "PB")).upper() == "PT":
                uw.append(f"zmiana projektu technicznego (art. 36b PB): zmiana w PT przez projektanta, sprawdzenie przez sprawdzającego "
                          f"(jeśli wymagane), ponowne uzgodnienia, jeśli rozwiązanie podlegało uzgodnieniom: {op}")
                continue
            if o.get("kwalifikacja") not in ("istotne", "nieistotne"):
                uw.append(f"odstąpienie bez kwalifikacji (istotne/nieistotne, art. 36a ust. 6 PB): {op}")
            if o.get("kwalifikacja") == "nieistotne":
                if not o.get("dokumentacja"):
                    uw.append(f"nieistotne odstąpienie bez rysunku i opisu do dokumentacji budowy (art. 36a ust. 6 PB): {op}")
                wym, pz = o.get("zmiana_wymiaru_proc"), o.get("zmiana_pow_zabudowy_proc")
                if wym is not None and abs(float(wym)) > 2:
                    uw.append(f"zmiana wysokości/długości/szerokości {wym}% > 2% – przesłanka istotności (art. 36a ust. 5 pkt 2 lit. b PB): {op}")
                if pz is not None and abs(float(pz)) > 5:
                    uw.append(f"zmiana powierzchni zabudowy {pz}% > 5% – przesłanka istotności (art. 36a ust. 5 pkt 2 lit. a PB): {op}")
            if o.get("kwalifikacja") == "istotne":
                uw.append(f"istotne odstąpienie – dopuszczalne dopiero po decyzji o zmianie pozwolenia na budowę (art. 36a ust. 1 PB): {op}")
        if re.search(r"zrid", str(sp.get("decyzja", "")), re.I):
            uw.append("inwestycja na decyzji ZRID – tryb zmiany decyzji wg specustawy drogowej; kwalifikację odstąpień uzgodnij z inwestorem i organem")
    if not sp.get("uczestnicy"):
        uw.append("brak listy uczestników")
    return uw


def cmd_notatka(a):
    sp = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    typ = sp.get("typ", "narada")
    tyt = TYPY.get(typ, "NOTATKA") + (f" nr {sp['numer']}" if sp.get("numer") else "")
    if typ == "nadzor" and sp.get("pobyt_nr"):
        tyt += f" – pobyt nr {sp['pobyt_nr']}"
    d = nowy_docx(tyt, sp.get("temat"), sp.get("firma"))
    meta = [["Data i godzina:", f"{data_pl(sp.get('data'))}{', godz. ' + sp['godzina'] if sp.get('godzina') else ''}"],
            ["Miejsce / forma:", sp.get("miejsce", "")], ["Prowadzący:", sp.get("prowadzacy", "")], ["Protokolant:", sp.get("protokolant", "")]]
    if typ == "nadzor":
        meta += [["Budowa:", sp.get("budowa", "")], ["Inwestor:", sp.get("inwestor", "")], ["Wykonawca robót:", sp.get("wykonawca_robot", "")],
                 ["Kierownik budowy:", sp.get("kierownik_budowy", "")], ["Inspektor nadzoru:", sp.get("inspektor_nadzoru", "")],
                 ["Decyzja:", sp.get("decyzja", "")], ["Wpis do dziennika budowy:", sp.get("wpis_dziennik", "")]]
    tabela(d, [m for m in meta if m[1]], szer=[4.5, 12.5], naglowkowa=False)
    naglowek(d, "Uczestnicy")
    tabela(d, [["Lp.", "Imię i nazwisko", "Firma / instytucja", "Funkcja"]] +
           [[i, u.get("osoba", ""), u.get("firma", ""), u.get("funkcja", "")] for i, u in enumerate(sp.get("uczestnicy", []), 1)], szer=[1, 5, 6, 5])
    if sp.get("porzadek"):
        naglowek(d, "Porządek spotkania")
        for i, p in enumerate(sp["porzadek"], 1):
            d.add_paragraph(f"{i}. {p}")
    if typ == "nadzor":
        naglowek(d, "Zgodność realizacji z projektem (art. 20 ust. 1 pkt 4 lit. a PB)")
        d.add_paragraph(sp.get("zgodnosc_z_projektem", ""))
        if sp.get("odstapienia"):
            naglowek(d, "Rozwiązania zamienne i odstąpienia – kwalifikacja projektanta (art. 36a ust. 5–6 PB)")
            tabela(d, [["Lp.", "Opis", "PB / PT", "Kwalifikacja", "Uzasadnienie", "Dokumentacja"]] +
                   [[i, o.get("opis", ""), o.get("projekt", "PB"),
                     ("zmiana PT (art. 36b)" if str(o.get("projekt", "PB")).upper() == "PT" else o.get("kwalifikacja", "")),
                     o.get("uzasadnienie", ""), o.get("dokumentacja", "")]
                    for i, o in enumerate(sp["odstapienia"], 1)], szer=[0.9, 4.6, 1.6, 2.4, 4.5, 3.0], rozm=8.5)
    if sp.get("ustalenia"):
        naglowek(d, "Ustalenia")
        for i, u in enumerate(sp["ustalenia"], 1):
            t = u["tresc"] if isinstance(u, dict) else u
            pref = f"Ad {u['punkt']}." if isinstance(u, dict) and u.get("punkt") else f"{i}."
            d.add_paragraph(f"{pref} {t}")
    if sp.get("decyzje"):
        naglowek(d, "Decyzje")
        for i, t in enumerate(sp["decyzje"], 1):
            d.add_paragraph(f"{i}. {t}")
    if sp.get("zadania"):
        naglowek(d, "Zadania")
        tabela(d, [["Lp.", "Zadanie", "Odpowiedzialny", "Termin", "Status"]] +
               [[i, z.get("co", ""), z.get("kto", ""), z.get("termin", ""), z.get("status", "otwarte")] for i, z in enumerate(sp["zadania"], 1)],
               szer=[1, 8, 3.5, 2.3, 2.2])
    if sp.get("sprawy_otwarte"):
        naglowek(d, "Sprawy otwarte")
        for i, t in enumerate(sp["sprawy_otwarte"], 1):
            d.add_paragraph(f"{i}. {t}")
    if sp.get("nastepne_spotkanie"):
        d.add_paragraph(f"Następne spotkanie: {sp['nastepne_spotkanie']}.")
    if sp.get("zalaczniki"):
        naglowek(d, "Załączniki")
        for i, t in enumerate(sp["zalaczniki"], 1):
            d.add_paragraph(f"{i}. {t}")
    p = d.add_paragraph("Uwagi do protokołu uczestnicy zgłaszają w terminie 3 dni roboczych od jego otrzymania; po tym terminie protokół uznaje się za przyjęty.")
    p.runs[0].font.size = __import__("docx").shared.Pt(8.5)
    if sp.get("rozdzielnik"):
        d.add_paragraph("Rozdzielnik: " + ", ".join(sp["rozdzielnik"]))
    naglowek(d, "Podpisy")
    if typ == "nadzor":
        podpisy(d, "Projektant – nadzór autorski:", "Kierownik budowy / Inspektor nadzoru:", n=1)
    else:
        podpisy(d, "Sporządził:", "Zatwierdził:", n=1)
    d.save(a.out)
    print(f"Zapisano {a.out}")
    uw = kontrola_notatki(sp)
    if uw:
        print("Kontrola:\n - " + "\n - ".join(uw))
    if a.md:
        L = [f"# {tyt}", "", f"{sp.get('temat', '')} — {data_pl(sp.get('data'))}, {sp.get('miejsce', '')}", "",
             "**Uczestnicy:** " + "; ".join(f"{u.get('osoba')} ({u.get('firma')})" for u in sp.get("uczestnicy", [])), ""]
        if typ == "nadzor" and sp.get("odstapienia"):
            L += ["## Rozwiązania zamienne / odstąpienia", "| Opis | PB/PT | Kwalifikacja | Dokumentacja |", "|---|---|---|---|"]
            L += [f"| {o.get('opis', '')} | {o.get('projekt', 'PB')} | {o.get('kwalifikacja', '')} | {o.get('dokumentacja', '')} |" for o in sp["odstapienia"]] + [""]
        if sp.get("ustalenia"):
            L += ["## Ustalenia"]
            for i, u in enumerate(sp["ustalenia"], 1):
                L.append(f"- Ad {u['punkt']}. {u['tresc']}" if isinstance(u, dict) and u.get("punkt") else f"- {u['tresc'] if isinstance(u, dict) else u}")
            L.append("")
        for klucz, tyt_ in (("decyzje", "Decyzje"), ("sprawy_otwarte", "Sprawy otwarte")):
            if sp.get(klucz):
                L += [f"## {tyt_}"] + [f"{i}. {t}" for i, t in enumerate(sp[klucz], 1)] + [""]
        if sp.get("zadania"):
            L += ["## Zadania", "| Zadanie | Kto | Termin |", "|---|---|---|"] + [f"| {z.get('co')} | {z.get('kto') or '—'} | {z.get('termin') or '—'} |" for z in sp["zadania"]] + [""]
        if sp.get("nastepne_spotkanie"):
            L.append(f"Następne spotkanie: {sp['nastepne_spotkanie']}.")
        Path(a.md).write_text("\n".join(L) + "\n", encoding="utf-8")


def cmd_transkrypt(a):
    p = Path(a.plik)
    suf = p.suffix.lower()
    if suf == ".docx":
        import docx
        tekst = "\n".join(x.text for x in docx.Document(str(p)).paragraphs)
    else:
        tekst = p.read_text(encoding="utf-8", errors="replace")
    wyp = []
    if suf == ".vtt" or tekst.lstrip().startswith("WEBVTT"):
        bloki = re.split(r"\n\s*\n", tekst)
        for b in bloki:
            linie = [l for l in b.strip().splitlines() if l.strip()]
            linie = [l for l in linie if not re.match(r"^(WEBVTT|NOTE|\d+$|[\w-]+/\d+-\d+$)", l.strip())]
            czas = next((l for l in linie if "-->" in l), None)
            tresc = " ".join(l for l in linie if "-->" not in l)
            m = re.match(r"<v ([^>]+)>(.*?)(</v>)?$", tresc)
            mowca, t = (m.group(1), m.group(2)) if m else ("", tresc)
            if t.strip():
                wyp.append((czas.split("-->")[0].strip()[:8] if czas else "", mowca.strip(), re.sub(r"<[^>]+>", "", t).strip()))
    else:
        # format „Imię Nazwisko  0:01:23” / „Imię Nazwisko: tekst” (eksport Teams do DOCX/TXT)
        biez = None
        for l in tekst.splitlines():
            l = l.strip()
            if not l:
                continue
            m = re.match(r"^([A-ZŁŚŻŹĆŃÓĘĄ][\wąćęłńóśźż.\-]+(?: [A-ZŁŚŻŹĆŃÓĘĄ][\wąćęłńóśźż.\-]+){0,3})\s+(\d{1,2}:\d{2}(?::\d{2})?)$", l)
            m2 = re.match(r"^([A-ZŁŚŻŹĆŃÓĘĄ][\wąćęłńóśźż.\-]+(?: [A-ZŁŚŻŹĆŃÓĘĄ][\wąćęłńóśźż.\-]+){0,3}):\s+(.+)$", l)
            if m:
                biez = [m.group(2), m.group(1), ""]
                wyp.append(biez)
            elif m2:
                wyp.append(["", m2.group(1), m2.group(2)])
                biez = None
            elif biez is not None:
                biez[2] = (biez[2] + " " + l).strip()
            else:
                wyp.append(["", "", l])
        wyp = [tuple(x) for x in wyp if x[2]]
    # scalanie kolejnych wypowiedzi tego samego mowcy
    scal = []
    for c, m, t in wyp:
        if scal and scal[-1][1] == m:
            scal[-1] = (scal[-1][0], m, scal[-1][2] + " " + t)
        else:
            scal.append((c, m, t))
    L = [f"# Transkrypt: {p.name}", "", f"Wypowiedzi: {len(scal)}; mówcy: {', '.join(sorted({m for _, m, _ in scal if m})) or '—'}", ""]
    for c, m, t in scal:
        L += [f"**{m or '?'}**{' [' + c + ']' if c else ''}: {t}", ""]
    Path(a.out).write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"Zapisano {a.out}: {len(scal)} wypowiedzi.")


def cmd_zadania(a):
    wiersze = []
    for p in a.pliki:
        sp = json.loads(Path(p).read_text(encoding="utf-8"))
        for z in sp.get("zadania", []):
            wiersze.append({**z, "zrodlo": f"{TYPY.get(sp.get('typ'), 'notatka').lower()} {sp.get('data', '')}", "temat": sp.get("temat", "")})
    nd = date.fromisoformat(a.na_dzien) if a.na_dzien else date.today()

    def stan(z):
        if z.get("status") == "zamknięte":
            return "zamknięte"
        try:
            t = date.fromisoformat(z.get("termin", ""))
            return "PO TERMINIE" if t < nd else ("≤ 7 dni" if t <= nd + timedelta(days=7) else "otwarte")
        except ValueError:
            return "bez terminu"
    wiersze.sort(key=lambda z: ({"PO TERMINIE": 0, "≤ 7 dni": 1, "bez terminu": 2, "otwarte": 3, "zamknięte": 4}[stan(z)], z.get("termin", "")))
    L = [f"# Zadania z protokołów (stan na {nd.isoformat()})", "", "| Stan | Termin | Zadanie | Kto | Źródło |", "|---|---|---|---|---|"]
    L += [f"| {stan(z)} | {z.get('termin', '')} | {z.get('co', '')} | {z.get('kto', '')} | {z['zrodlo']} |" for z in wiersze]
    t = "\n".join(L)
    if a.md:
        Path(a.md).write_text(t + "\n", encoding="utf-8")
    print(t)


def main():
    import signal
    if hasattr(signal, "SIGPIPE"):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("rf-init"); s.add_argument("harmonogram"); s.add_argument("--out", required=True); s.add_argument("--nadpisz", action="store_true")
    for k in ("nazwa", "umowa", "zamawiajacy", "projektant"):
        s.add_argument("--" + k)
    s.add_argument("--suma", type=float)
    s = sub.add_parser("rf-nowy"); s.add_argument("projekt"); s.add_argument("--postep", required=True); s.add_argument("--data")
    s.add_argument("--okres"); s.add_argument("--uwagi"); s.add_argument("--zatrzymanie", type=float); s.add_argument("--korekta", action="store_true")
    s.add_argument("--xlsx"); s.add_argument("--md")
    s = sub.add_parser("rf-xlsx"); s.add_argument("projekt"); s.add_argument("--nr", type=int, required=True); s.add_argument("--out", required=True)
    s = sub.add_parser("rf-stan"); s.add_argument("projekt")
    s = sub.add_parser("wzor-przekazanie"); s.add_argument("--out", default="przekazanie.json")
    s = sub.add_parser("manifest"); s.add_argument("katalog"); s.add_argument("--out", default="manifest.csv")
    s = sub.add_parser("przekazanie"); s.add_argument("spec"); s.add_argument("--out", required=True); s.add_argument("--katalog"); s.add_argument("--manifest")
    s = sub.add_parser("wzor-notatka"); s.add_argument("--typ", default="narada", choices=list(TYPY)); s.add_argument("--out", default="notatka.json")
    s = sub.add_parser("transkrypt"); s.add_argument("plik"); s.add_argument("--out", required=True)
    s = sub.add_parser("notatka"); s.add_argument("spec"); s.add_argument("--out", required=True); s.add_argument("--md")
    s = sub.add_parser("zadania"); s.add_argument("pliki", nargs="+"); s.add_argument("--md"); s.add_argument("--na-dzien")
    a = ap.parse_args()
    {"rf-init": cmd_rf_init, "rf-nowy": cmd_rf_nowy, "rf-xlsx": cmd_rf_xlsx, "rf-stan": cmd_rf_stan, "wzor-przekazanie": cmd_wzor_przekazanie,
     "manifest": cmd_manifest, "przekazanie": cmd_przekazanie, "wzor-notatka": cmd_wzor_notatka, "transkrypt": cmd_transkrypt,
     "notatka": cmd_notatka, "zadania": cmd_zadania}[a.cmd](a)


if __name__ == "__main__":
    main()
