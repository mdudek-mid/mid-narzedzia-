#!/usr/bin/env python3
"""przedmiar_tool.py - ilosci do przedmiarow robot z rysunkow DWG/DXF i PDF, kontrola przedmiarow XLSX
i skoroszyt weryfikacji przedmiaru zamawiajacego (Pracownia Projektowa MiD).

Komendy:
  obmiar RYS.dwg|dxf [--mapa mapa.csv] [--lisp obmiar.csv] [--os UCHWYT] [--csv obiekty.csv] [--xlsx wynik.xlsx]
        obiekty przestrzeni modelu: dlugosc, pole (kreskowania z wyspami, zamkniete polilinie, okregi), liczba,
        bloki, kilometraz na osi; sumy wg warstw albo wg mapy warstw (warstwa -> pozycja przedmiaru)
  pdf-pola SKALA ARKUSZ1.pdf [ARKUSZ2.pdf ...] [--legenda legenda.pdf] [--json wynik.json]
        pola wg koloru wypelnienia (szrafury pelne z AutoCAD/Civil 3D); odrzuca znaczniki i logo,
        usuwa duplikaty miedzy arkuszami (dopasowanie tych samych szraf, transformacja sztywna)
  pdf-materialy RYS.pdf [--png podglad.png]
        tabele „Zastosowane materialy” z rysunkow zbrojenia (beton m3, stal t), takze strony obrocone
  stal ZESTAWIENIE.csv
        przeliczenie zestawienia stali: poz;szt;fi;dlugosc_m -> masa wg kg/m (kontrola zestawien projektanta)
  skaner PRZEDMIAR.xlsx [--arkusz NAZWA] [--pierwsza G] [--wiersz-nazw 1] [--od 10] [--opis C]
        K5: formula obiektu bierze parametr z kolumny innego obiektu; K6: wiersz liczony inaczej niz wiekszosc
  weryfikacja KONFIGURACJA.json WYNIK.xlsx
        skoroszyt weryfikacji przedmiaru zamawiajacego z wynikow branz (JSON): statusy, formuly roznic,
        „ilosc do wyceny”, PODSUMOWANIE, KONTROLE, ZRODLA; bez cen
  wzor-weryfikacji KATALOG    przykladowa konfiguracja i plik branzy
  test                        testy na danych syntetycznych (DXF, PDF, XLSX tworzone w locie)

Zaleznosci: ezdxf, shapely, openpyxl, pymupdf, numpy; do DWG LibreDWG (dwg2dxf, katalog ../libredwg).
Obiekty REGION, bryly ACIS i obiekty Civil 3D nie sa czytelne poza AutoCAD-em - wtedy MID_OBMIAR.lsp
(polecenie AutoCAD-a, eksport do CSV) i --lisp.
"""
import argparse, collections, csv, datetime, hashlib, json, math, os, re, shutil, subprocess, sys, tempfile

KG_M = {6: 0.222, 8: 0.395, 10: 0.617, 12: 0.888, 14: 1.208, 16: 1.578, 20: 2.466, 25: 3.853, 28: 4.834,
        32: 6.313, 40: 9.865}

# ---------------------------------------------------------------------------------------------- DWG/DXF
EXE = ".exe" if os.name == "nt" else ""
BIN = os.environ.get("LIBREDWG_BIN") or os.path.expanduser(os.path.join("~", ".local", "libredwg", "bin"))
KRZYWE = ("LINE", "LWPOLYLINE", "POLYLINE", "ARC", "SPLINE", "CIRCLE", "ELLIPSE")
TYPY_ACAD = {"AcDbHatch": "HATCH", "AcDbRegion": "REGION", "AcDbPolyline": "LWPOLYLINE", "AcDbLine": "LINE",
             "AcDb2dPolyline": "POLYLINE", "AcDb3dPolyline": "POLYLINE", "AcDbArc": "ARC", "AcDbCircle": "CIRCLE",
             "AcDbEllipse": "ELLIPSE", "AcDbSpline": "SPLINE", "AcDbBlockReference": "INSERT",
             "AcDbMText": "MTEXT", "AcDbText": "TEXT", "AcDb3dSolid": "3DSOLID"}


def _ezdxf():
    import ezdxf
    import ezdxf.entities.mtext as mt
    if not getattr(mt, "_mid_fix", False):           # LibreDWG zapisuje czasem MTEXT z blednymi kolumnami
        orig = mt.load_columns_from_embedded_object

        def safe(*a, **k):
            try:
                return orig(*a, **k)
            except Exception:
                return None
        mt.load_columns_from_embedded_object = safe; mt._mid_fix = True
    return ezdxf


def _u(s):
    from ezdxf.lldxf.encoding import decode_dxf_unicode, has_dxf_unicode
    s = s or ""
    return decode_dxf_unicode(s) if has_dxf_unicode(s) else s


def md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def do_dxf(path, workdir):
    if path.lower().endswith(".dxf"):
        return path
    os.makedirs(workdir, exist_ok=True)
    out = os.path.join(workdir, os.path.splitext(os.path.basename(path))[0] + ".dxf")
    if not (os.path.exists(out) and os.path.getmtime(out) >= os.path.getmtime(path)):
        with open(path, "rb") as f:
            if f.read(6).decode("ascii", "replace") == "AC1021":
                print(f"[uwaga] {os.path.basename(path)}: DWG 2007 (AC1021) - LibreDWG czyta go z bledami; "
                      "zapisz jako DWG 2018 albo DXF.", file=sys.stderr)
        exe = shutil.which("dwg2dxf") or os.path.join(BIN, "dwg2dxf" + EXE)
        r = subprocess.run([exe, "-y", "-o", out, path], capture_output=True, text=True)
        if not os.path.exists(out) or os.path.getsize(out) == 0:
            raise RuntimeError(f"dwg2dxf nie odczytal {path}: {r.stderr[-500:]}")
    return out


def wczytaj_rysunek(path, workdir=None):
    ezdxf = _ezdxf()
    dxf = do_dxf(path, workdir or os.path.join(tempfile.gettempdir(), "mid_obmiar"))
    try:
        return ezdxf.readfile(dxf)
    except Exception:
        from ezdxf import recover
        return recover.readfile(dxf)[0]


def _punkty(e):
    from ezdxf import path as P
    return [(v.x, v.y) for v in P.make_path(e).flattening(0.005)]


def pole_kreskowania(h):
    """pole netto: kontury zagniezdzone parzysto-nieparzyscie (wyspy odejmowane)"""
    from ezdxf import path as P
    from shapely.geometry import Polygon
    polys = []
    for p in h.paths:
        pts = [(v.x, v.y) for v in P.from_hatch_boundary_path(p).flattening(0.005)]
        if len(pts) >= 3:
            g = Polygon(pts).buffer(0)
            if g.area > 1e-6:
                polys.append(g)
    tot = 0.0
    for i, g in enumerate(polys):
        rp = g.representative_point()
        gl = sum(1 for j, o in enumerate(polys) if j != i and o.area > g.area and o.contains(rp))
        tot += g.area if gl % 2 == 0 else -g.area
    return tot


def obiekty(doc, klucz, os_odn=None):
    from shapely.geometry import Polygon, Point
    from ezdxf import bbox
    out = []
    for e in doc.modelspace():
        t = e.dxftype()
        rec = {"plik": klucz, "uchwyt": e.dxf.handle, "warstwa": _u(e.dxf.get("layer", "0")), "typ": t, "blok": "",
               "dlugosc": None, "pole": None, "zamknieta": "", "x": None, "y": None, "km": None, "uwaga": ""}
        try:
            if t in KRZYWE:
                pts = _punkty(e)
                rec["dlugosc"] = sum(math.dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1))
                zamk = t == "CIRCLE" or bool(getattr(e, "closed", False)) or (
                    len(pts) > 3 and math.dist(pts[0], pts[-1]) < 0.05)
                rec["zamknieta"] = "tak" if zamk else "nie"
                if zamk and len(pts) >= 3:
                    rec["pole"] = Polygon(pts).buffer(0).area
                xs, ys = [p[0] for p in pts], [p[1] for p in pts]
                rec["x"], rec["y"] = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
            elif t == "HATCH":
                rec["pole"] = pole_kreskowania(e)
                b = bbox.extents([e], fast=True)
                if b.has_data:
                    rec["x"], rec["y"] = b.center.x, b.center.y
            elif t == "INSERT":
                rec["blok"] = _u(e.dxf.name); rec["x"], rec["y"] = e.dxf.insert.x, e.dxf.insert.y
            elif t in ("REGION", "3DSOLID", "BODY", "SURFACE"):
                rec["uwaga"] = "dane ACIS - pole tylko z AutoCAD-a (MID_OBMIAR.lsp)"
        except Exception as ex:
            rec["uwaga"] = f"blad odczytu: {ex}"
        if os_odn is not None and rec["x"] is not None:
            rec["km"] = os_odn.project(Point(rec["x"], rec["y"]))
        out.append(rec)
    return out


def os_z_rysunku(doc, uchwyt):
    from shapely.geometry import LineString
    for e in doc.modelspace():
        if e.dxf.handle == uchwyt:
            return LineString(_punkty(e))
    raise KeyError(f"brak osi o uchwycie {uchwyt}")


def wczytaj_csv_lisp(path, klucz):
    """CSV z MID_OBMIAR.lsp (AutoCAD, kodowanie 1250)"""
    out = []
    with open(path, encoding="cp1250", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            n = lambda s: float(s) if s not in ("", None) else None
            out.append({"plik": klucz, "uchwyt": r["uchwyt"], "warstwa": r["warstwa"], "typ": TYPY_ACAD.get(r["typ"], r["typ"]),
                        "blok": r.get("blok", ""), "dlugosc": n(r.get("dlugosc_m")), "pole": n(r.get("pole_m2")),
                        "zamknieta": r.get("zamknieta", ""), "x": n(r.get("x")), "y": n(r.get("y")), "km": None,
                        "uwaga": "z AutoCAD-a (MID_OBMIAR.lsp)"})
    return out


def wczytaj_mape(path):
    """mapa warstw CSV (;): id;plik;warstwa;pomiar;typy;uchwyty;mnoznik;jm;pozycja;cel;uwagi
    pomiar: POLE | POLE_POLILINII | DLUGOSC | LICZBA | LICZBA_BLOKOW; plik = klucz rysunku (nazwa bez rozszerzenia)"""
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.DictReader(f, delimiter=";") if r.get("id") and not r["id"].startswith("#")]
    for r in rows:
        r["typy"] = [s.strip() for s in (r.get("typy") or "").split(",") if s.strip()]
        r["uchwyty"] = [s.strip() for s in (r.get("uchwyty") or "").split(",") if s.strip()]
        r["mnoznik"] = float((r.get("mnoznik") or "1").replace(",", "."))
    return rows


def zmierz(poz, objs):
    sel = [o for o in objs if (not poz.get("plik") or o["plik"] == poz["plik"]) and o["warstwa"] == poz["warstwa"]
           and (not poz["typy"] or o["typ"] in poz["typy"]) and (not poz["uchwyty"] or o["uchwyt"] in poz["uchwyty"])]
    m = poz["pomiar"].upper()
    if m == "POLE":
        sel = [o for o in sel if o["typ"] in ("HATCH", "REGION") and o["pole"] is not None]; v = sum(o["pole"] for o in sel)
    elif m == "POLE_POLILINII":
        sel = [o for o in sel if o["pole"] is not None]; v = sum(o["pole"] for o in sel)
    elif m == "DLUGOSC":
        sel = [o for o in sel if o["dlugosc"] is not None]; v = sum(o["dlugosc"] for o in sel)
    elif m == "LICZBA":
        v = len(sel)
    elif m == "LICZBA_BLOKOW":
        sel = [o for o in sel if o["typ"] == "INSERT"]; v = len(sel)
    else:
        raise ValueError(f"nieznany pomiar {m} w {poz['id']}")
    return v * poz["mnoznik"], len(sel), [o["uchwyt"] for o in sel]


def sumy_warstw(objs):
    s = collections.OrderedDict()
    for o in sorted(objs, key=lambda o: (o["plik"], o["warstwa"], o["typ"])):
        k = (o["plik"], o["warstwa"], o["typ"])
        a = s.setdefault(k, {"liczba": 0, "dlugosc": 0.0, "pole": 0.0})
        a["liczba"] += 1; a["dlugosc"] += o["dlugosc"] or 0; a["pole"] += o["pole"] or 0
    return s


def cmd_obmiar(a):
    klucz = os.path.splitext(os.path.basename(a.rysunek))[0]
    if a.lisp:
        objs = wczytaj_csv_lisp(a.lisp, klucz)
    else:
        doc = wczytaj_rysunek(a.rysunek)
        objs = obiekty(doc, klucz, os_z_rysunku(doc, a.os) if a.os else None)
    if a.csv:
        with open(a.csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(objs[0].keys()), delimiter=";"); w.writeheader(); w.writerows(objs)
    print(f"{klucz}: {len(objs)} obiektow" + (f", md5 {md5(a.rysunek)}" if os.path.exists(a.rysunek) else ""))
    wiersze = []
    if a.mapa:
        for poz in wczytaj_mape(a.mapa):
            v, n, uch = zmierz(dict(poz, plik=poz.get("plik") if poz.get("plik") in (klucz,) else ""), objs)
            wiersze.append([poz["id"], poz.get("pozycja", ""), poz["warstwa"], poz["pomiar"], round(v, 2), poz.get("jm", ""), n,
                            ",".join(uch[:50]) + (" …" if len(uch) > 50 else "")])
            print(f"  {poz['id']:6} {poz.get('pozycja', '')[:50]:50} {v:12.2f} {poz.get('jm', '')}  ({n} ob.)")
        hdr = ["id", "pozycja", "warstwa", "pomiar", "ilosc", "jm", "obiektow", "uchwyty"]
    else:
        for (p, wl, t), s in sumy_warstw(objs).items():
            wiersze.append([wl, t, s["liczba"], round(s["dlugosc"], 2), round(s["pole"], 2)])
            print(f"  {wl[:40]:40} {t:12} {s['liczba']:6d}  dl {s['dlugosc']:10.2f}  pole {s['pole']:10.2f}")
        hdr = ["warstwa", "typ", "liczba", "dlugosc", "pole"]
    if a.xlsx:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "OBMIAR"; ws.append(hdr)
        for r in wiersze: ws.append(r)
        ws2 = wb.create_sheet("OBIEKTY"); ws2.append(list(objs[0].keys()))
        for o in objs: ws2.append(list(o.values()))
        wb.save(a.xlsx); print("zapisano", a.xlsx)


# ---------------------------------------------------------------------------------------------- PDF: pola
POMIN_KOLORY = {(0.0, 0.0, 0.0), (1.0, 1.0, 1.0)}
ZNACZNIKI = {}


def _pierscienie(items):
    rings, cur, last = [], [], None
    for it in items:
        k = it[0]
        if k == "re":
            r = it[1]; rings.append([(r.x0, r.y0), (r.x1, r.y0), (r.x1, r.y1), (r.x0, r.y1)]); last = None; continue
        if k == "qu":
            q = it[1]; rings.append([(q.ul.x, q.ul.y), (q.ur.x, q.ur.y), (q.lr.x, q.lr.y), (q.ll.x, q.ll.y)]); last = None; continue
        p0 = it[1]; pts = [it[1], it[-1]] if k == "l" else [it[1], it[2], it[3], it[4]]
        if last is None or abs(p0.x - last.x) + abs(p0.y - last.y) > 1e-3:
            if len(cur) >= 3:
                rings.append(cur)
            cur = [(p0.x, p0.y)]
        cur += [(p.x, p.y) for p in pts[1:]]; last = pts[-1]
    if len(cur) >= 3:
        rings.append(cur)
    return rings


def szrafury(pdf, strona=0, pomin=()):
    """{kolor: [(geometria, pierscienie)]}; geometria przycieta do aktywnego obszaru clip (rzutnia)"""
    import pymupdf
    from shapely.geometry import Polygon, box
    from shapely.validation import make_valid
    out = collections.defaultdict(list); clip = None
    for g in pymupdf.open(pdf)[strona].get_drawings(extended=True):
        if g["type"] == "clip":
            s = g.get("scissor"); clip = box(s.x0, s.y0, s.x1, s.y1) if s else None; continue
        if not g.get("fill") or "f" not in g["type"]:
            continue
        col = tuple(round(x, 2) for x in g["fill"])
        if col in POMIN_KOLORY or col in pomin:
            continue
        rs = _pierscienie(g["items"]); geo = None
        for r in rs:
            try:
                p = make_valid(Polygon(r))
            except Exception:
                continue
            geo = p if geo is None else (geo.symmetric_difference(p) if g.get("even_odd") else geo.union(p))
        if geo is None or geo.is_empty:
            continue
        if clip is not None:
            geo = geo.intersection(clip)
        out[col].append((geo, rs))
    # znaczniki: identyczne prostokaty powtorzone >= 3 razy w tym samym kolorze (linie podzialu arkuszy, strzalki)
    for col, v in out.items():
        grupy = collections.defaultdict(list)
        for i, (geo, rs) in enumerate(v):
            if len(rs) == 1 and len(rs[0]) <= 5 and geo.area > 0.98 * geo.minimum_rotated_rectangle.area:
                xy = list(geo.minimum_rotated_rectangle.exterior.coords)
                grupy[tuple(sorted((round(math.dist(xy[0], xy[1]), 0), round(math.dist(xy[1], xy[2]), 0))))].append(i)
        usun = {i for g in grupy.values() if len(g) >= 3 for i in g}
        if usun:
            ZNACZNIKI[col] = ZNACZNIKI.get(col, 0) + len(usun)
            out[col] = [x for i, x in enumerate(v) if i not in usun]
    return out


def _kabsch(A, B):
    import numpy as np
    A, B = np.asarray(A), np.asarray(B); ca, cb = A.mean(0), B.mean(0)
    U, S, Vt = np.linalg.svd((B - cb).T @ (A - ca)); d = np.sign(np.linalg.det(U @ Vt))
    R = (U @ np.diag([1, d]) @ Vt).T
    return R, ca - R @ cb, float(np.abs((B @ R.T + ca - R @ cb) - A).max())


def _klucz(col, rs):
    return (col, tuple(len(r) for r in rs))


def _dopasuj(ra, rb, tol):
    """Kabsch na wierzcholkach; przy pojedynczym konturze probuje przesuniec cyklicznych (inny wierzcholek startowy)"""
    A = [p for r in ra for p in r]
    warianty = [rb] if len(rb) > 1 else [[rb[0][k:] + rb[0][:k]] for k in range(len(rb[0]))]
    for w in warianty:
        R, t, err = _kabsch(A, [p for r in w for p in r])
        if err <= tol:
            return R, t
    return None


def rejestracja(sA, sB, tol=0.5, maks=40):
    """te same szrafury na dwoch arkuszach (kolor, liczba wierzcholkow, pole +-0,1 %) -> transformacja B->A;
    wybiera kandydata, dla ktorego najwiecej szraf B trafia srodkiem w szrafy A (symbole powtarzane na kazdym
    arkuszu, np. strzalka polnocy, daja malo trafien i odpadaja)"""
    import numpy as np
    idxA = collections.defaultdict(list)
    for c, v in sA.items():
        for g, rs in v:
            idxA[_klucz(c, rs)].append((g, rs, g.area, g.centroid.x, g.centroid.y))
    srA = {k: np.array([(x[3], x[4]) for x in v]) for k, v in idxA.items()}
    rowne = lambda a1, a2: abs(a1 - a2) <= 1e-3 * max(a1, 1e-9)
    listaB = []
    for c, v in sB.items():
        for g, rs in v:
            k = _klucz(c, rs)
            if k in idxA and any(rowne(g.area, x[2]) for x in idxA[k]):
                listaB.append((k, g, rs, g.centroid.x, g.centroid.y))
    if not listaB:
        return None
    srB = np.array([(x[3], x[4]) for x in listaB]); kluczeB = [x[0] for x in listaB]
    best = None; n = 0
    for k, gb, rb, _, _ in sorted(listaB, key=lambda x: -x[1].area):
        if sum(len(r) for r in rb) < 5:
            continue
        for ga, ra, area, _, _ in idxA[k]:
            if not rowne(ga.area, gb.area):
                continue
            d = _dopasuj(ra, rb, tol)
            if d is None:
                continue
            R, t = d; n += 1
            P = srB @ R.T + t
            traf = sum(bool((np.abs(srA[kb] - P[i]).sum(1) < 1.0).any()) for i, kb in enumerate(kluczeB))
            if best is None or traf > best[3]:
                best = (R, t, k[0], traf)
            if traf >= 0.5 * len(listaB):
                return best
        if n >= maks:
            break
    return best


def legenda_kolorow(pdf, strona=0):
    """kolor wypelnienia probki -> tekst na prawo od niej (legenda rysunku)"""
    import pymupdf
    p = pymupdf.open(pdf)[strona]; W = p.get_text("words"); out = {}
    for g in p.get_drawings():
        if not g.get("fill"):
            continue
        col = tuple(round(x, 2) for x in g["fill"]); r = g["rect"]
        if col in POMIN_KOLORY or r.width < 10 or r.height < 5:
            continue
        t = [w for w in W if abs((w[1] + w[3]) / 2 - (r.y0 + r.y1) / 2) < max(5, r.height / 2) and r.x1 < w[0] < r.x1 + 250]
        if t:
            out.setdefault(col, " ".join(w[4] for w in sorted(t, key=lambda w: w[0])))
    return out


def pola_pdf(skala, arkusze, pomin=()):
    from shapely import affinity
    from shapely.ops import unary_union
    m2 = (25.4 / 72 / 1000 * skala) ** 2
    baza = szrafury(arkusze[0], pomin=pomin); wsp = {c: [g for g, _ in v] for c, v in baza.items()}; info = []
    for f in arkusze[1:]:
        s = szrafury(f, pomin=pomin); rej = rejestracja(baza, s)
        if not rej or rej[3] < 2:
            info.append(f"{os.path.basename(f)}: brak wspolnych szraf z arkuszem 1 - sumuje bez usuwania duplikatow")
            for c, v in s.items():
                wsp.setdefault(c, []).extend(g for g, _ in v)
            continue
        R, t, col, traf = rej
        info.append(f"{os.path.basename(f)}: dopasowany do arkusza 1 ({traf} wspolnych szraf, obrot "
                    f"{math.degrees(math.atan2(R[1, 0], R[0, 0])):.2f} st.)")
        mat = [R[0, 0], R[0, 1], R[1, 0], R[1, 1], t[0], t[1]]
        for c, v in s.items():
            wsp.setdefault(c, []).extend(affinity.affine_transform(g, mat) for g, _ in v)
    netto = {c: unary_union(v).area * m2 for c, v in wsp.items()}
    brutto = {c: sum(g.area for g in v) * m2 for c, v in wsp.items()}
    return netto, brutto, info


def cmd_pdf_pola(a):
    leg = legenda_kolorow(a.legenda) if a.legenda else {}
    pomin = {tuple(float(x) for x in k.split(",")) for k in (a.pomin or [])}
    netto, brutto, info = pola_pdf(a.skala, a.arkusze, pomin)
    for i in info:
        print(i)
    if ZNACZNIKI:
        print("pominiete znaczniki (powtarzalne prostokaty):", {str(k): v for k, v in ZNACZNIKI.items()})
    wyn = []
    for c in sorted(netto, key=lambda c: -netto[c]):
        if netto[c] < a.min:
            continue
        wyn.append({"kolor": c, "opis": leg.get(c, ""), "m2": round(netto[c], 1), "m2_bez_scalania": round(brutto[c], 1)})
        print(f"{str(c):20} {netto[c]:10.1f} m2  (bez scalania arkuszy {brutto[c]:10.1f})  {leg.get(c, '')}")
    if a.json:
        json.dump({"skala": a.skala, "arkusze": a.arkusze, "info": info, "pola": wyn}, open(a.json, "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)


# ---------------------------------------------------------------------------------------------- PDF: materialy
def tabela_materialow(pdf, zoom=1.6, png=None):
    import pymupdf
    out = []
    for i, page in enumerate(pymupdf.open(pdf)):
        W = page.get_text("words"); obr = page.rotation in (90, 270)
        for z in [w for w in W if w[4].startswith("Zastosowane")]:
            x0, y0 = z[0], z[1]
            if obr:
                reg = [w for w in W if x0 - 160 < w[0] < x0 + 10 and y0 - 60 < w[1] < y0 + 600]
                key = lambda w: round(w[0] / 6); order = lambda w: w[1]
            else:
                reg = [w for w in W if x0 - 200 < w[0] < x0 + 300 and y0 - 10 < w[1] < y0 + 260]
                key = lambda w: round(w[1] / 4); order = lambda w: w[0]
            rows = {}
            for w in reg:
                rows.setdefault(key(w), []).append(w)
            linie = [" ".join(x[4] for x in sorted(v, key=order)) for _, v in sorted(rows.items(), reverse=obr)]
            out.append({"strona": i + 1, "linie": linie})
            if png and not any(ch.isdigit() for l in linie if "Zastosowane" not in l for ch in l):
                page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), clip=pymupdf.Rect(x0 - 170, y0 - 10, x0 + 300, y0 + 260)).save(png)
    return out


def cmd_pdf_materialy(a):
    t = tabela_materialow(a.pdf, png=a.png)
    if not t:
        print("brak tabeli „Zastosowane materialy” jako tekstu - teksty SHX: wyrenderuj fragment i odczytaj wzrokowo")
    for x in t:
        print(f"--- strona {x['strona']}")
        for l in x["linie"]:
            print("   ", l)


# ---------------------------------------------------------------------------------------------- stal
def cmd_stal(a):
    suma = collections.Counter(); tot = 0.0
    with open(a.csv, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            fi = int(float(r["fi"])); m = float(r["szt"]) * float(r["dlugosc_m"].replace(",", ".")) * KG_M[fi]
            suma[fi] += m; tot += m
            print(f"  poz {r.get('poz', ''):5} {r['szt']:>5} x O{fi:<3} L={r['dlugosc_m']:>7} m  {m:9.1f} kg")
    for fi in sorted(suma):
        print(f"O{fi}: {suma[fi]:10.1f} kg")
    print(f"RAZEM {tot:10.1f} kg")


# ---------------------------------------------------------------------------------------------- skaner K5/K6
REF = re.compile(r"(?<![A-Z!$])\$?([A-Z]{1,3})\$?(\d+)")


def skan(xlsx, arkusz=None, pierwsza="G", wiersz_nazw=1, od=10, opis="C"):
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL
    wb = load_workbook(xlsx); ws = wb[arkusz] if arkusz else wb.worksheets[0]
    p, o = CI(pierwsza), CI(opis)
    obj = [c for c in range(p, ws.max_column + 1) if ws.cell(wiersz_nazw, c).value]
    nazwa = lambda c: ws.cell(wiersz_nazw, c).value
    for r in range(od, ws.max_row + 1):
        podpisy = {}
        for c in obj:
            v = ws.cell(r, c).value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            for k, _ in set(REF.findall(v)):
                if CI(k) != c and CI(k) in obj:
                    yield ("K5", nazwa(c), f"{CL(c)}{r}", ws.cell(r, o).value, v, f"z kolumny {nazwa(CI(k))}")
            podpisy[c] = " ".join(sorted({("@" if CI(k) == c else "#") + w for k, w in REF.findall(v)}))
        if len(podpisy) >= 3:
            typowy, ile = collections.Counter(podpisy.values()).most_common(1)[0]
            if typowy and ile > len(podpisy) / 2:
                for c, s in podpisy.items():
                    if s != typowy:
                        yield ("K6", nazwa(c), f"{CL(c)}{r}", ws.cell(r, o).value, ws.cell(r, c).value, f"typowo {typowy}")


def cmd_skaner(a):
    n = 0
    for k in skan(a.xlsx, a.arkusz, a.pierwsza, a.wiersz_nazw, a.od, a.opis):
        n += 1; print(" | ".join(str(x) for x in k))
    print(f"{n} uwag")


# ---------------------------------------------------------------------------------------------- weryfikacja
STATUSY = ["zgodna", "różnica", "brak w przedmiarze", "brak w dokumentacji", "nie do sprawdzenia"]
KOLORY = {"zgodna": "E2F0D9", "różnica": "FCE4D6", "brak w przedmiarze": "F8CBAD", "brak w dokumentacji": "FFF2CC",
          "nie do sprawdzenia": "EDEDED"}


def status(zam, dok, jm, tol=0.01):
    """zgodna przy |roznica| <= 1 % (szt. bez tolerancji)"""
    if dok is None:
        return "nie do sprawdzenia" if zam is not None else ""
    if zam is None:
        return "brak w przedmiarze"
    if (jm or "").strip().lower().startswith("szt"):
        return "zgodna" if abs(dok - zam) < 1e-9 else "różnica"
    return "zgodna" if abs(dok - zam) <= tol * max(abs(zam), 1e-9) else "różnica"


def weryfikacja(cfg_path, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    W = os.path.dirname(os.path.abspath(cfg_path)); CFG = json.load(open(cfg_path, encoding="utf-8"))
    branze = []
    for f in CFG["wyniki"]:
        d = json.load(open(os.path.join(W, f), encoding="utf-8"))
        for b in (d if isinstance(d, list) else [d]):
            for pref, pw in CFG.get("pomin_wiersze", {}).items():
                if b["branza"].startswith(pref):       # wiersze sprawdzone w arkuszu innej branzy - pokaz raz
                    b["pozycje"] = [p for p in b["pozycje"] if not (pw["od"] <= (p.get("wiersz") or 0) <= pw["do"])]
                    b.setdefault("uwagi_ogolne", []).append(pw.get("uwaga", ""))
            branze.append(b)
    thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78"); WR = Alignment(wrap_text=True, vertical="top")
    txt = lambda v: ("wzór " + str(v)) if str(v or "").startswith("=") else ("" if v is None else str(v))

    def num(v):
        if isinstance(v, (int, float)):
            return v
        try:
            return float(str(v).replace(",", ".").replace(" ", ""))
        except Exception:
            return None
    wb = Workbook(); ws0 = wb.active; ws0.title = "PODSUMOWANIE"
    for b in branze:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", b["branza"])[:28])
        ws["A1"] = f"Przedmiar wg dokumentacji – {b['branza']}"; ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Przedmiar Zamawiającego: {b.get('przedmiar_plik', '')}"
        ws["A3"] = "Ilość do wyceny = ilość wg dokumentacji, gdy ustalona; w przeciwnym razie ilość Zamawiającego. Bez cen."
        hdr = ["Lp. przedmiaru", "Kod ST", "Opis pozycji", "j.m.", "Ilość Zamawiającego", "Ilość wg dokumentacji", "Różnica",
               "Różnica %", "Ilość do wyceny", "Status", "Źródło w dokumentacji", "Uwagi"]
        r0 = 5
        for i, h in enumerate(hdr, 1):
            c = ws.cell(r0, i, h); c.font = H; c.fill = HF; c.alignment = WR; c.border = B
        r = r0
        for p in b["pozycje"]:
            r += 1; z, d = num(p.get("ilosc_zam")), num(p.get("ilosc_dok"))
            st = p.get("status") or status(z, d, p.get("jm"))
            lp = p.get("lp"); lp = "" if lp in (None, "null") else lp
            for i, v in enumerate([lp, txt(p.get("st")), txt(p.get("opis")), txt(p.get("jm")), z, d], 1):
                ws.cell(r, i, v)
            ws.cell(r, 7, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),F{r}-E{r},"")')
            ws.cell(r, 8, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r}),E{r}<>0),(F{r}-E{r})/E{r},"")')
            ws.cell(r, 9, f'=IF(ISNUMBER(F{r}),F{r},IF(ISNUMBER(E{r}),E{r},""))')
            ws.cell(r, 10, st); ws.cell(r, 11, txt(p.get("zrodlo"))); ws.cell(r, 12, txt(p.get("uwaga")))
            for i in range(1, 13):
                ws.cell(r, i).border = B; ws.cell(r, i).alignment = WR
            if st in KOLORY:
                ws.cell(r, 10).fill = PatternFill("solid", fgColor=KOLORY[st])
            for i in (5, 6, 7, 9):
                ws.cell(r, i).number_format = "#,##0.00"
            ws.cell(r, 8).number_format = "0.0%"
            p["status"] = st
        for i, w in enumerate([9, 13, 48, 7, 13, 13, 12, 9, 13, 16, 45, 50], 1):
            ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = ws.cell(r0 + 1, 4); ws.auto_filter.ref = f"A{r0}:L{max(r, r0)}"
        b["_ark"], b["_r0"], b["_r1"] = ws.title, r0 + 1, r
        rr = r + 2; ws.cell(rr, 1, "Uwagi ogólne").font = Font(bold=True)
        zm = [f"Zmiana oryginał→zamienny: {z if isinstance(z, str) else json.dumps(z, ensure_ascii=False)}"
              for z in b.get("zmiany_oryginal_zamienny", [])]
        for u in b.get("uwagi_ogolne", []) + zm:
            rr += 1; ws.cell(rr, 1, u); ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=12)
            ws.cell(rr, 1).alignment = WR; ws.row_dimensions[rr].height = 30
    ws = ws0
    ws["A1"] = CFG.get("tytul", "Weryfikacja przedmiaru"); ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = CFG.get("opis", ""); ws["A3"] = f"Opracowanie: {CFG.get('autor', 'Pracownia Projektowa MiD')}, " \
                                             f"{datetime.date.today():%d.%m.%Y}. Arkusz roboczy do wyceny – bez cen."
    ws["A5"] = CFG.get("obiekty", "")
    for c in ("A2", "A3", "A5"):
        ws[c].alignment = WR
    for rr in (2, 5):
        ws.merge_cells(f"A{rr}:H{rr}"); ws.row_dimensions[rr].height = 32
    ws.merge_cells("A1:H1"); ws.merge_cells("A3:H3")
    r = 7
    for i, h in enumerate(["Branża", "Arkusz", "Pozycji", "Zgodne", "Różnice", "Brak w przedmiarze", "Brak w dokumentacji",
                           "Nie do sprawdzenia"], 1):
        c = ws.cell(r, i, h); c.font = H; c.fill = HF; c.border = B; c.alignment = WR
    for b in branze:
        r += 1; a_, r0, r1 = b["_ark"], b["_r0"], b["_r1"]
        ws.cell(r, 1, b["branza"]); ws.cell(r, 2, a_); ws.cell(r, 3, f"=COUNTA('{a_}'!C{r0}:C{r1})")
        for i, s in enumerate(STATUSY, 4):
            ws.cell(r, i, f"=COUNTIF('{a_}'!J{r0}:J{r1},\"{s}\")")
        for i in range(1, 9):
            ws.cell(r, i).border = B
    r += 2; ws.cell(r, 1, "Najważniejsze ustalenia (pełna lista: arkusz KONTROLE)").font = Font(bold=True, size=12)
    for i, (k, t) in enumerate([tuple(x) for x in CFG.get("top", [])], 1):
        r += 1; ws.cell(r, 1, f"{i}. {k}").font = Font(bold=True); ws.cell(r, 1).alignment = WR
        ws.cell(r, 2, t).alignment = WR; ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 32
    for i, w in enumerate([34, 28, 9, 9, 9, 12, 12, 12], 1):
        ws.column_dimensions[L(i)].width = w
    wk = wb.create_sheet("KONTROLE", 1)
    for i, h in enumerate(["Branża", "Lp.", "Opis", "j.m.", "Zamawiający", "Wg dokumentacji", "Różnica", "Status", "Źródło",
                           "Uwagi"], 1):
        c = wk.cell(1, i, h); c.font = H; c.fill = HF; c.border = B; c.alignment = WR
    r = 1
    for b in branze:
        for p in b["pozycje"]:
            if p.get("status") not in ("różnica", "brak w przedmiarze", "brak w dokumentacji"):
                continue
            r += 1
            vals = [b["branza"], "" if p.get("lp") in (None, "null") else p.get("lp"), txt(p.get("opis")), txt(p.get("jm")),
                    num(p.get("ilosc_zam")), num(p.get("ilosc_dok")), f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),F{r}-E{r},"")',
                    p.get("status"), txt(p.get("zrodlo")), txt(p.get("uwaga"))]
            for i, v in enumerate(vals, 1):
                c = wk.cell(r, i, v); c.border = B; c.alignment = WR
            wk.cell(r, 8).fill = PatternFill("solid", fgColor=KOLORY.get(p.get("status"), "FFFFFF"))
            for i in (5, 6, 7):
                wk.cell(r, i).number_format = "#,##0.00"
    for i, w in enumerate([18, 7, 45, 7, 13, 13, 12, 16, 45, 55], 1):
        wk.column_dimensions[L(i)].width = w
    wk.freeze_panes = "C2"; wk.auto_filter.ref = f"A1:J{max(r, 1)}"
    wz = wb.create_sheet("ŹRÓDŁA"); wz["A1"] = CFG.get("zrodla_tytul", "Źródła"); wz["A1"].font = Font(bold=True)
    for i, s in enumerate(CFG.get("zrodla", []), 3):
        wz.cell(i, 1, s).alignment = WR
    wz.column_dimensions["A"].width = 140
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"; w.page_setup.paperSize = 9
        w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0; w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    return [b["branza"] for b in branze]


def cmd_weryfikacja(a):
    print("zapisano", a.wynik, weryfikacja(a.konfiguracja, a.wynik))


def cmd_wzor(a):
    os.makedirs(a.katalog, exist_ok=True)
    cfg = {"tytul": "Weryfikacja przedmiaru – <zadanie>", "opis": "Ryczałt: ilości wg dokumentacji projektowej (PW + zamienne).",
           "obiekty": "Obiekty mostowe: <ile i jakie>", "wyniki": ["wynik_branza1.json"], "pomin_wiersze": {},
           "top": [["Branża", "najważniejsza różnica: ilość Zamawiającego → ilość wg dokumentacji, źródło"]],
           "zrodla_tytul": "Źródła", "zrodla": ["<ścieżka>/<plik> – co z niego wzięto"]}
    br = {"branza": "Branża 1", "przedmiar_plik": "<plik przedmiaru>", "pozycje": [
        {"lp": 1, "st": "D-00.00.00", "opis": "Pozycja przykładowa", "jm": "m2", "ilosc_zam": 100.0, "ilosc_dok": 112.5,
         "zrodlo": "rys. X, tabela Y", "status": "", "uwaga": ""}], "zmiany_oryginal_zamienny": [], "uwagi_ogolne": []}
    json.dump(cfg, open(os.path.join(a.katalog, "konfiguracja.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(br, open(os.path.join(a.katalog, "wynik_branza1.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("zapisano wzor w", a.katalog)


# ---------------------------------------------------------------------------------------------- testy
def cmd_test(a):
    ok = True; tmp = tempfile.mkdtemp(prefix="przedmiar_test_")

    def spr(nazwa, war, opis=""):
        nonlocal ok
        ok &= bool(war); print(f"[{'OK' if war else 'BLAD'}] {nazwa} {opis}")
    # 1. DXF: kreskowanie 10x10 z wyspa 2x2, polilinia zamknieta 4x5, linia 7 m, 3 bloki
    ezdxf = _ezdxf(); doc = ezdxf.new(); ms = doc.modelspace()
    h = ms.add_hatch(dxfattribs={"layer": "NAW"}); h.paths.add_polyline_path([(0, 0), (10, 0), (10, 10), (0, 10)], is_closed=True)
    h.paths.add_polyline_path([(4, 4), (6, 4), (6, 6), (4, 6)], is_closed=True)
    ms.add_lwpolyline([(20, 0), (24, 0), (24, 5), (20, 5)], close=True, dxfattribs={"layer": "CHOD"})
    ms.add_line((0, 20), (7, 20), dxfattribs={"layer": "KRAW"})
    blk = doc.blocks.new("SLUP")
    blk.add_circle((0, 0), 0.2)
    for x in range(3):
        ms.add_blockref("SLUP", (x * 5, 30), dxfattribs={"layer": "OSW"})
    f = os.path.join(tmp, "t.dxf"); doc.saveas(f)
    objs = obiekty(wczytaj_rysunek(f), "t"); s = sumy_warstw(objs)
    spr("obmiar: kreskowanie z wyspa", abs(s[("t", "NAW", "HATCH")]["pole"] - 96) < 1e-6, f"{s[('t', 'NAW', 'HATCH')]['pole']:.2f} (96)")
    spr("obmiar: pole polilinii", abs(s[("t", "CHOD", "LWPOLYLINE")]["pole"] - 20) < 1e-6)
    spr("obmiar: dlugosc linii", abs(s[("t", "KRAW", "LINE")]["dlugosc"] - 7) < 1e-6)
    spr("obmiar: bloki", s[("t", "OSW", "INSERT")]["liczba"] == 3)
    mapa = os.path.join(tmp, "mapa.csv")
    with open(mapa, "w", encoding="utf-8") as fm:
        fm.write("id;plik;warstwa;pomiar;typy;uchwyty;mnoznik;jm;pozycja;cel;uwagi\n1;t;NAW;POLE;;;1;m2;nawierzchnia;;\n"
                 "2;t;KRAW;DLUGOSC;;;1;m;krawężnik;;\n")
    v = [zmierz(p, objs)[0] for p in wczytaj_mape(mapa)]
    spr("obmiar: mapa warstw", abs(v[0] - 96) < 1e-6 and abs(v[1] - 7) < 1e-6)
    # 2. PDF: dwa arkusze 1:500, ten sam obszar na obu (arkusz 2 obrocony o 30 st.) + znacznik
    import pymupdf
    m = 72 / 25.4 * 1000 / 500                      # pt na metr w skali 1:500

    def arkusz(nazwa, rot, dx, dy, dodatkowy):
        d = pymupdf.open(); p = d.new_page(width=1200, height=800); sh = p.new_shape()
        c, s_ = math.cos(math.radians(rot)), math.sin(math.radians(rot))
        T = lambda x, y: pymupdf.Point(dx + (x * c - y * s_) * m, dy + (x * s_ + y * c) * m)
        for poly, col in (([(0, 0), (60, 0), (60, 8), (30, 12), (0, 8)], (0.75, 0.75, 0.75)),
                          ([(0, 9), (60, 9), (60, 11), (0, 11)], (0.8, 0.77, 0.56))):
            sh.draw_polyline([T(*q) for q in poly] + [T(*poly[0])]); sh.finish(fill=col, color=None, closePath=True)
        if dodatkowy:
            poly = [(70, 0), (100, 0), (100, 8), (70, 8)]
            sh.draw_polyline([T(*q) for q in poly] + [T(70, 0)]); sh.finish(fill=(0.75, 0.75, 0.75), color=None)
        for i in range(3):                              # znaczniki: 3 identyczne prostokaty
            sh.draw_rect(pymupdf.Rect(900, 600 + 40 * i, 960, 612 + 40 * i)); sh.finish(fill=(0.75, 0.75, 0.75), color=None)
        sh.commit(); fn = os.path.join(tmp, nazwa); d.save(fn); return fn
    a1, a2 = arkusz("a1.pdf", 0, 100, 100, False), arkusz("a2.pdf", 30, 400, 150, True)
    netto, brutto, info = pola_pdf(500, [a1, a2])
    g = netto[(0.75, 0.75, 0.75)]; k = netto[(0.8, 0.77, 0.56)]
    spr("pdf-pola: duplikat miedzy arkuszami usuniety", abs(g - (60 * 8 + 60 * 4 / 2 + 30 * 8)) < 2,
        f"{g:.1f} m2 (840), bez scalania {brutto[(0.75, 0.75, 0.75)]:.1f}")
    spr("pdf-pola: drugi kolor", abs(k - 120) < 1, f"{k:.1f} m2 (120)")
    # 3. skaner K5/K6
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws["G1"], ws["H1"], ws["I1"] = "OB-1", "OB-2", "OB-3"
    for col in "GHI":
        ws[f"{col}5"] = 10
    ws["C10"] = "beton"; ws["G10"] = "=G5*2"; ws["H10"] = "=G5*2"; ws["I10"] = "=I5*2"
    ws["C11"] = "stal"; ws["G11"] = "=G5*3"; ws["H11"] = "=H5*3"; ws["I11"] = "=I5*3+H5"
    fx = os.path.join(tmp, "p.xlsx"); wb.save(fx)
    u = list(skan(fx))
    spr("skaner: K5 (H10 bierze G5)", any(x[0] == "K5" and x[2] == "H10" for x in u))
    spr("skaner: K6 (I11 inaczej)", any(x[0] == "K6" and x[2] == "I11" for x in u), f"{len(u)} uwag")
    # 4. weryfikacja
    cmd_wzor(argparse.Namespace(katalog=tmp)); wx = os.path.join(tmp, "w.xlsx")
    weryfikacja(os.path.join(tmp, "konfiguracja.json"), wx)
    from openpyxl import load_workbook
    w2 = load_workbook(wx)
    spr("weryfikacja: arkusze", w2.sheetnames[:2] == ["PODSUMOWANIE", "KONTROLE"] and "ŹRÓDŁA" in w2.sheetnames)
    spr("weryfikacja: status auto", w2["Branża 1"]["J6"].value == "różnica")
    spr("status: szt. bez tolerancji", status(10, 10.05, "szt.") == "różnica" and status(100, 100.9, "m") == "zgodna")
    print("WSZYSTKO OK" if ok else "SA BLEDY", "-", tmp)
    return 0 if ok else 1


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("obmiar"); p.add_argument("rysunek"); p.add_argument("--mapa"); p.add_argument("--lisp")
    p.add_argument("--os"); p.add_argument("--csv"); p.add_argument("--xlsx"); p.set_defaults(f=cmd_obmiar)
    p = sp.add_parser("pdf-pola"); p.add_argument("skala", type=float); p.add_argument("arkusze", nargs="+")
    p.add_argument("--legenda"); p.add_argument("--json"); p.add_argument("--min", type=float, default=1.0)
    p.add_argument("--pomin", nargs="*", help="kolory do pominiecia, np. 0,1,1 0,0.59,0.59 (logo)")
    p.set_defaults(f=cmd_pdf_pola)
    p = sp.add_parser("pdf-materialy"); p.add_argument("pdf"); p.add_argument("--png"); p.set_defaults(f=cmd_pdf_materialy)
    p = sp.add_parser("stal"); p.add_argument("csv"); p.set_defaults(f=cmd_stal)
    p = sp.add_parser("skaner"); p.add_argument("xlsx"); p.add_argument("--arkusz"); p.add_argument("--pierwsza", default="G")
    p.add_argument("--wiersz-nazw", type=int, default=1); p.add_argument("--od", type=int, default=10)
    p.add_argument("--opis", default="C"); p.set_defaults(f=cmd_skaner)
    p = sp.add_parser("weryfikacja"); p.add_argument("konfiguracja"); p.add_argument("wynik"); p.set_defaults(f=cmd_weryfikacja)
    p = sp.add_parser("wzor-weryfikacji"); p.add_argument("katalog"); p.set_defaults(f=cmd_wzor)
    p = sp.add_parser("test"); p.set_defaults(f=cmd_test)
    a = ap.parse_args(argv)
    return a.f(a) or 0


if __name__ == "__main__":
    sys.exit(main())
