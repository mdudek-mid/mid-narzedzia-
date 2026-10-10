#!/usr/bin/env python3
"""kadra_tool.py - baza referencji i kadry MiD: dobor uslug i osob do warunkow udzialu, wykazy, zal. 1A.

Baza (prywatna, nie w tym repo): KATALOG/referencje.json i KATALOG/kadra.json.
Katalog: --baza KATALOG albo zmienna MID_BAZA (domyslnie ./baza_mid).

Komendy:
  stan                                  podsumowanie bazy i luki do uzupelnienia
  uslugi [filtry]                       referencje spelniajace warunek, z ocena kazdego kryterium
        --typ most,wiadukt,kladka,tunel,przejscie_podziemne,przepust,estakada,przejscie_dla_zwierzat
        --zakres PB,PW,ZRID,...  (wszystkie musza wystapic)   --dowolny-zakres PB,PW (wystarczy jeden)
        --lat 3 --od RRRR-MM-DD (termin skladania ofert)      --min-wartosc 120000 [--netto]
        --dlugosc-min M --rozpietosc-min M --klasa-drogi G --nad kolej --szukaj TEKST
        --tylko-spolka (bez doswiadczenia JDG, ktore wymaga udostepnienia zasobow)   --max 15
  osoby [--specjalnosc mostowa] [--bez-ograniczen] [--lat-od-uprawnien 5 --od RRRR-MM-DD]
        [--rownowazne] (uprawnienia konstrukcyjno-budowlane traktuj jak mostowe/drogowe wg starych przepisow)
        [--dokumentacje-min N --typ most --lat 10]  [--wszyscy] (takze osoby znane tylko z referencji)
  pokaz ID                              pelny rekord (R... albo K...)
  doswiadczenie ID_OSOBY [--typ most] [--lat 10 --od RRRR-MM-DD] [--form ZAL_1A.docx --wybierz 1,2,3 --out X.docx]
        dokumentacje osoby (kryterium doswiadczenia / zal. 1A), opcjonalnie wpisane w formularz
  wykaz-uslug ID,ID [--form WZOR.docx --out WYNIK.docx]   tresc wierszy wykazu uslug (MD albo wpisana w formularz)
  wykaz-osob ID[:rola],ID[:rola] [--od RRRR-MM-DD] [--form WZOR.docx --out WYNIK.docx]
  xlsx --out Baza_referencji_i_kadry.xlsx   eksport do przegladu w Excelu
"""
import argparse, json, os, re, sys
from datetime import date
from pathlib import Path

_PL = str.maketrans("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ", "acelnoszzACELNOSZZ")
DUZO = "[DO UZUPEŁNIENIA]"


def norm(s):
    return (s or "").translate(_PL).lower()


def zl(v):
    return f"{v:,.2f} zł".replace(",", " ").replace(".", ",") if isinstance(v, (int, float)) else DUZO


def dpl(s):
    """RRRR-MM-DD -> DD.MM.RRRR"""
    if not s:
        return DUZO
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    return f"{m.group(3)}.{m.group(2)}.{m.group(1)}" if m else s


def wczytaj(baza):
    b = Path(baza)
    r = json.load(open(b / "referencje.json", encoding="utf-8"))
    k = json.load(open(b / "kadra.json", encoding="utf-8"))
    return r["referencje"], k["osoby"]


def lata_wstecz(od: date, lat: float) -> date:
    try:
        return od.replace(year=od.year - int(lat))
    except ValueError:  # 29 lutego
        return od.replace(year=od.year - int(lat), day=28)


# ----------------------------------------------------------------- referencje
def opis_obiektow(r):
    out = []
    for o in r.get("obiekty") or []:
        t = o.get("typ") or "obiekt"
        cz = [t + (f" nad {o['nad']}" if o.get("nad") else "")]
        if o.get("dlugosc_calkowita_m"):
            cz.append(f"L={o['dlugosc_calkowita_m']:g} m")
        if o.get("rozpietosc_max_m"):
            cz.append(f"przęsło max {o['rozpietosc_max_m']:g} m")
        if o.get("klasa_obciazenia"):
            cz.append(f"kl. {o['klasa_obciazenia']}")
        out.append(", ".join(cz))
    return "; ".join(out)


def ocen_usluge(r, a, od):
    """Lista (kryterium, wynik ✓/✗/?, komentarz)."""
    oc = []
    obj = r.get("obiekty") or []
    if a.typ:
        typy = set(a.typ.split(","))
        jest = [o for o in obj if o.get("typ") in typy]
        oc.append(("typ", "✓" if jest else "✗", ",".join(sorted({o.get("typ") or "?" for o in obj})) or "brak obiektu"))
    zak = set(r.get("zakres") or [])
    if a.zakres:
        brak = [z for z in a.zakres.split(",") if z not in zak]
        oc.append(("zakres", "✗" if brak else "✓", ("brak: " + ",".join(brak)) if brak else ",".join(sorted(zak))))
    if a.dowolny_zakres:
        jest = [z for z in a.dowolny_zakres.split(",") if z in zak]
        oc.append(("zakres (jeden z)", "✓" if jest else "✗", ",".join(sorted(zak))))
    if a.lat:
        granica = lata_wstecz(od, a.lat)
        dz = r.get("data_zakonczenia")
        if not dz:
            oc.append((f"okres {a.lat:g} lat", "?", "brak daty zakończenia"))
        else:
            oc.append((f"okres {a.lat:g} lat", "✓" if date.fromisoformat(dz) >= granica else "✗", f"zakończono {dpl(dz)} (granica {dpl(granica.isoformat())})"))
    if a.min_wartosc:
        w = r.get("wartosc_netto_pln") if a.netto else r.get("wartosc_brutto_pln")
        if w is None:
            oc.append(("wartość", "?", "brak wartości " + ("netto" if a.netto else "brutto")))
        else:
            oc.append(("wartość", "✓" if w >= a.min_wartosc else "✗", zl(w)))
    if a.dlugosc_min:
        d = max([o.get("dlugosc_calkowita_m") or 0 for o in obj] or [0])
        oc.append(("długość obiektu", "?" if not d else ("✓" if d >= a.dlugosc_min else "✗"), f"{d:g} m" if d else "brak danych"))
    if a.rozpietosc_min:
        d = max([o.get("rozpietosc_max_m") or 0 for o in obj] or [0])
        oc.append(("rozpiętość przęsła", "?" if not d else ("✓" if d >= a.rozpietosc_min else "✗"), f"{d:g} m" if d else "brak danych"))
    if a.klasa_drogi:
        k = (r.get("droga") or {}).get("klasa")
        rank = {"A": 6, "S": 5, "GP": 4, "G": 3, "Z": 2, "L": 1, "D": 0}
        oc.append(("klasa drogi", "?" if not k else ("✓" if rank.get(k, -1) >= rank.get(a.klasa_drogi, 9) else "✗"), k or "brak danych"))
    if a.nad:
        jest = [o for o in obj if norm(o.get("nad")) == norm(a.nad)]
        oc.append((f"nad {a.nad}", "✓" if jest else ("?" if not obj else "✗"), ""))
    p = r.get("podmiot") or ""
    if "JDG" in p:
        oc.append(("podmiot", "!" if not a.tylko_spolka else "✗", "doświadczenie JDG Marcin Dudek - w ofercie spółki tylko jako udostępnienie zasobów (art. 118 Pzp)"))
    return oc


def cmd_uslugi(a):
    R, _ = wczytaj(a.baza)
    od = date.fromisoformat(a.od) if a.od else date.today()
    wyn = []
    for r in R:
        if a.szukaj and norm(a.szukaj) not in norm(json.dumps(r, ensure_ascii=False)):
            continue
        oc = ocen_usluge(r, a, od)
        if any(w == "✗" for _, w, _ in oc):
            continue
        niepewne = sum(1 for _, w, _ in oc if w in "?!")
        wyn.append((niepewne, -(r.get("wartosc_brutto_pln") or 0), r, oc))
    wyn.sort(key=lambda x: (x[0], x[1]))
    print(f"Pasujących referencji: {len(wyn)} (✓ spełnia, ? brak danych w bazie - sprawdź w dowodzie, ! uwaga formalna)\n")
    for _, _, r, oc in wyn[: a.max]:
        flag = " [DO WERYFIKACJI]" if r.get("do_weryfikacji") else ""
        print(f"### {r['id']} {r.get('etykieta','')}{flag}")
        print(f"- {r.get('nazwa')}")
        print(f"- zamawiający: {r.get('zamawiajacy_nazwa')}; zakończenie {dpl(r.get('data_zakonczenia'))}; wartość brutto {zl(r.get('wartosc_brutto_pln'))}")
        if r.get("obiekty"):
            print(f"- obiekty: {opis_obiektow(r)}")
        for k, w, c in oc:
            print(f"  {w} {k}: {c}")
        if r.get("rozbieznosci"):
            print(f"  ! rozbieżności: {'; '.join(r['rozbieznosci'])[:300]}")
        print(f"  dowód: {', '.join(r.get('dowody') or []) or 'brak w bazie'}\n")
    if not wyn:
        print("Brak referencji spełniających wszystkie kryteria. Rozluźnij filtr (bez --min-wartosc / --lat), żeby zobaczyć najbliższe, albo rozważ podmiot udostępniający zasoby / konsorcjum.")


# ----------------------------------------------------------------- osoby
def upr_tekst(u):
    t = f"uprawnienia budowlane {u.get('rodzaj') or 'do projektowania'} {u.get('zakres') or ''} w specjalności {u.get('specjalnosc') or DUZO}"
    return re.sub(r"\s+", " ", t).strip() + f", nr {u.get('numer') or DUZO}, wydane {dpl(u.get('data_wydania'))}"


def lata_od(d, od):
    if not d:
        return None
    d = date.fromisoformat(d)
    return round((od - d).days / 365.25, 1)


def dokumentacje_osoby(p, R, typ=None, lat=None, od=None, tylko_projekty=False):
    """Dokumentacje z pola doswiadczenie osoby i z referencji, gdzie jest w projektantach (bez duplikatow)."""
    nazw = norm(p["imie_nazwisko"])
    out = []
    for d in p.get("doswiadczenie") or []:
        out.append({"zadanie": d.get("zadanie"), "rola": d.get("rola"), "zamawiajacy": d.get("zamawiajacy"), "data": d.get("data"), "zrodlo": "kadra"})
    for r in R:
        for pr in r.get("projektanci") or []:
            if norm(pr.get("osoba")) == nazw:
                out.append({"zadanie": r.get("nazwa"), "rola": pr.get("rola"), "zamawiajacy": r.get("zamawiajacy_nazwa"),
                            "data": r.get("data_zakonczenia"), "zrodlo": r["id"], "obiekty": [o.get("typ") for o in r.get("obiekty") or []]})
    # duplikaty: ten sam projekt z pola doswiadczenie osoby i z referencji - wspolne rzadkie slowo (miejscowosc, nr drogi)
    from collections import Counter
    ogolne = {"rozbudowa", "przebudowa", "opracowanie", "wykonanie", "projekt", "budowa", "remont", "modernizacja", "aktualizacja",
              "generalna", "zarzad", "gmina", "miasta", "miasto", "dyrekcja", "drogi", "drogowych", "krajowych", "autostrad",
              "zadania", "etap", "ekspertyza", "ocena", "pelnienie", "sporzadzenie", "projektowej", "dokumentacji"}
    tok = lambda x: {norm(t) for t in re.findall(r"\b[A-ZĄĆĘŁŃÓŚŹŻ][a-ząćęłńóśźż]{4,}", x or "")} - ogolne | set(re.findall(r"\b(?:dk|dw|dp|s|lk)\s?\d+", norm(x)))
    df = Counter(t for r in R for t in tok(r.get("nazwa")))
    rzadkie = lambda x: {t for t in tok(x) if df[t] <= 1}
    uniq, klucze = [], []
    for d in sorted(out, key=lambda x: x.get("zrodlo") == "kadra"):
        k = rzadkie(d.get("zadanie"))
        rok = (d.get("data") or "")[:4]
        if k and any((k & kk) and (not rok or not rr or abs(int(rok) - int(rr)) <= 1) for kk, rr in klucze):
            continue
        klucze.append((k, rok))
        uniq.append(d)
    if typ:
        typy = typ.split(",")
        uniq = [d for d in uniq if any(t in norm(d.get("zadanie")) or t in (d.get("obiekty") or []) for t in typy)
                or any(t in norm(d.get("zadanie")) for t in {"most": ["most", "wiadukt", "kladk", "obiekt"]}.get(typ, []))]
    if tylko_projekty:  # kryteria "doswiadczenie projektanta" zwykle nie licza ekspertyz, przegladow ani samego nadzoru
        uniq = [d for d in uniq if not re.search(r"ekspertyz|przeglad|ocen\w* stanu|nadzor|orzeczeni|analiz\w* oblicze", norm(d.get("zadanie")))]
    if lat and od:
        g = lata_wstecz(od, lat)
        uniq = [d for d in uniq if d.get("data") and date.fromisoformat(d["data"][:10]) >= g]
    return sorted(uniq, key=lambda d: d.get("data") or "", reverse=True)


def cmd_osoby(a):
    R, K = wczytaj(a.baza)
    od = date.fromisoformat(a.od) if a.od else date.today()
    for p in K:
        if not a.wszyscy and p["kategoria"].startswith("tylko"):
            continue
        ups = p.get("uprawnienia") or []
        if a.specjalnosc:
            sp = norm(a.specjalnosc)
            ok = [u for u in ups if sp in norm(u.get("specjalnosc"))]
            if a.rownowazne and sp in ("mostowa", "drogowa"):
                ok += [u for u in ups if "konstrukcyjno" in norm(u.get("specjalnosc"))]
            ups = ok
        if a.bez_ograniczen:
            ups = [u for u in ups if "bez" in norm(u.get("zakres"))]
        if (a.specjalnosc or a.bez_ograniczen) and not ups:
            continue
        lata = max([lata_od(u.get("data_wydania"), od) or -1 for u in ups] or [-1])
        if a.lat_od_uprawnien and lata < a.lat_od_uprawnien:
            if lata >= 0:
                continue
        docs = dokumentacje_osoby(p, R, a.typ, a.lat, od, a.tylko_projekty)
        if a.dokumentacje_min and len(docs) < a.dokumentacje_min:
            continue
        flag = " [DO WERYFIKACJI]" if p.get("do_weryfikacji") else ""
        print(f"### {p['id']} {p.get('tytul') or ''} {p['imie_nazwisko']} - {p['kategoria']}{flag}")
        for u in ups or p.get("uprawnienia") or []:
            print(f"- {upr_tekst(u)}")
        print(f"- lata od uprawnień (na {dpl(od.isoformat())}): {lata if lata >= 0 else 'brak daty'}; podstawa: {p.get('podstawa_dysponowania') or DUZO}")
        print(f"- dokumentacje{' (' + a.typ + ')' if a.typ else ''}{' z ' + str(a.lat) + ' lat' if a.lat else ''}: {len(docs)}")
        if p.get("uwagi"):
            print(f"  uwagi: {p['uwagi'][:250]}")
        print()


def cmd_doswiadczenie(a):
    R, K = wczytaj(a.baza)
    od = date.fromisoformat(a.od) if a.od else date.today()
    p = next((x for x in K if x["id"] == a.id), None)
    if not p:
        sys.exit(f"Brak osoby {a.id}")
    docs = dokumentacje_osoby(p, R, a.typ, a.lat, od, a.tylko_projekty)
    print(f"{p.get('tytul') or ''} {p['imie_nazwisko']}: {len(docs)} pozycji" + (f" (typ {a.typ})" if a.typ else "") + (f", ostatnie {a.lat} lat" if a.lat else ""))
    print("| # | zadanie | rola | zamawiający | data | źródło |\n|---|---|---|---|---|---|")
    for i, d in enumerate(docs, 1):
        print(f"| {i} | {(d.get('zadanie') or '')[:140]} | {d.get('rola') or ''} | {d.get('zamawiajacy') or ''} | {dpl(d.get('data'))} | {d.get('zrodlo')} |")
    if a.form:
        wyb = [docs[int(x) - 1] for x in a.wybierz.split(",")] if a.wybierz else docs
        imie = f"{p.get('tytul') or ''} {p['imie_nazwisko']}".strip()
        wiersze = []
        for d in wyb:
            r = next((x for x in R if x["id"] == d.get("zrodlo")), None)
            wiersze.append({"imie": imie, "nazwa": d.get("zadanie") or DUZO, "podmiot": d.get("zamawiajacy") or DUZO,
                            "parametry": (opis_obiektow(r) if r else "") or DUZO, "data": dpl(d.get("data"))})
        mapa = wypelnij_docx(a.form, a.out, wiersze, KOL_1A)
        print(f"\nWpisano {len(wiersze)} pozycji do {a.out}; kolumny: {mapa}")
        print("Sprawdź parametry obiektów (DUZO = brak w bazie) i czy pozycje spełniają definicję z SWZ - ten załącznik zwykle NIE podlega uzupełnieniu.")


def cmd_pokaz(a):
    R, K = wczytaj(a.baza)
    for x in R + K:
        if x["id"] == a.id:
            print(json.dumps(x, ensure_ascii=False, indent=1))
            return
    sys.exit(f"Brak rekordu {a.id}")


def cmd_stan(a):
    R, K = wczytaj(a.baza)
    od = date.today()
    print(f"Referencje: {len(R)}; z wartością brutto {sum(1 for r in R if r.get('wartosc_brutto_pln'))}; "
          f"z datą {sum(1 for r in R if r.get('data_zakonczenia'))}; do weryfikacji {sum(1 for r in R if r.get('do_weryfikacji'))}")
    for lat in (3, 5, 10):
        g = lata_wstecz(od, lat).isoformat()
        print(f"  zakończone w ostatnich {lat} latach: {sum(1 for r in R if (r.get('data_zakonczenia') or '') >= g)}")
    print(f"  doświadczenie JDG (wymaga udostępnienia zasobów): {sum(1 for r in R if 'JDG' in (r.get('podmiot') or ''))}")
    print(f"Osoby: {len(K)}; w wykazach MiD {sum(1 for p in K if not p['kategoria'].startswith('tylko'))}; do weryfikacji {sum(1 for p in K if p.get('do_weryfikacji'))}")
    print("\nLuki w referencjach z ostatnich 5 lat (bez wartości albo daty):")
    g = lata_wstecz(od, 5).isoformat()
    for r in R:
        if r.get("pewnosc") != "niska" and (r.get("data_zakonczenia") or "9999") >= g and not (r.get("wartosc_brutto_pln") and r.get("data_zakonczenia")):
            print(f"  {r['id']} {r.get('etykieta','')}: brak {'wartości ' if not r.get('wartosc_brutto_pln') else ''}{'daty' if not r.get('data_zakonczenia') else ''}")
    nis = [r for r in R if r.get("pewnosc") == "niska"]
    print(f"  + {len(nis)} referencji znanych tylko z nazwy pliku (skany bez tekstu) - do odczytu OCR, gdy komputer z OneDrive jest połączony")


# ----------------------------------------------------------------- wykazy i formularze
def wiersz_uslugi(r):
    nazwa = r.get("nazwa") or DUZO
    obj = opis_obiektow(r)
    zak = ", ".join(r.get("zakres") or [])
    przedmiot = nazwa + (f" (obiekty: {obj})" if obj else "")
    zam = (r.get("zamawiajacy_nazwa") or DUZO) + (f", {r['zamawiajacy_adres']}" if r.get("zamawiajacy_adres") else "")
    if r.get("inwestor_koncowy"):
        zam += f" (inwestor: {r['inwestor_koncowy']})"
    okres = (f"{dpl(r.get('data_rozpoczecia'))} – " if r.get("data_rozpoczecia") else "") + dpl(r.get("data_zakonczenia"))
    wyk = "Pracownia Projektowa MiD Marcin Dudek (podmiot udostępniający zasoby)" if "JDG" in (r.get("podmiot") or "") else "Pracownia Projektowa MiD Sp. z o.o."
    return {"przedmiot": przedmiot, "wartosc": zl(r.get("wartosc_brutto_pln")), "data": okres, "podmiot": zam, "nazwa": nazwa,
            "wykonawca": wyk, "miejsce": r.get("miejsce") or DUZO, "parametry": obj or DUZO}


def wiersz_osoby(p, rola, od, R):
    ups = p.get("uprawnienia") or []
    kw = "; ".join(upr_tekst(u) for u in ups) or DUZO
    lata = max([lata_od(u.get("data_wydania"), od) or -1 for u in ups] or [-1])
    docs = dokumentacje_osoby(p, R)[:4]
    dosw = (f"{int(lata)} lat od uzyskania uprawnień" if lata >= 0 else DUZO) + ("; m.in.: " + "; ".join(
        f"{(d.get('zadanie') or '')[:110]} ({d.get('rola') or ''}, {dpl(d.get('data'))})" for d in docs) if docs else "")
    pod = p.get("podstawa_dysponowania") or DUZO
    posr = bool(re.search(r"udostępni|podmiot|art\.? ?118", pod))
    return {"imie": f"{p.get('tytul') or ''} {p['imie_nazwisko']}".strip(), "kwalifikacje": kw, "doswiadczenie": dosw,
            "rola": rola or ", ".join(p.get("funkcje_w_wykazach") or [])[:80] or DUZO, "podstawa": pod,
            "posrednie": pod if posr else "—", "bezposrednie": "—" if posr else pod,
            "firma": p.get("firma_zewnetrzna") or ("nie dotyczy" if not posr else DUZO),
            "wyksztalcenie": "wyższe techniczne" + (f" ({p['tytul']})" if p.get("tytul") else "")}


# kolejnosc ma znaczenie: pierwsze dopasowanie wygrywa (naglowki zamawiajacych sa rozwlekle)
KOL_USLUGI = [("lp", r"^l\.?\s*p|^lp\b|^poz"), ("wartosc", r"^warto"), ("data", r"^data|^termin|^okres|^czas realiz"),
              ("wykonawca", r"^nazwa wykonawc|^wykonawc"), ("podmiot", r"podmiot|zamawiaj|odbiorc|na rzecz|inwestor|zleceniodaw"),
              ("miejsce", r"^miejsce"), ("przedmiot", r"przedmiot|rodzaj|nazwa|opis|zakres|uslug")]
KOL_OSOBY = [("lp", r"^l\.?\s*p|^lp\b"), ("imie", r"imi|nazwisk"), ("kwalifikacje", r"kwalifikac|uprawnie"),
             ("doswiadczenie", r"doswiadcz"), ("wyksztalcenie", r"wyksztalc"), ("bezposrednie", r"bezposredni"),
             ("posrednie", r"posredni"), ("firma", r"konsorcj|nazwe firmy|podmiot"), ("podstawa", r"podstaw|dysponow"),
             ("rola", r"zakres|funkcj|rola|czynno|stanowisk")]
KOL_1A = [("lp", r"^l\.?\s*p|^lp\b"), ("imie", r"imi|nazwisk"), ("parametry", r"parametr"),
          ("podmiot", r"inwestor|zamawiaj|podmiot|na rzecz"), ("data", r"^data|termin"), ("nazwa", r"nazwa|zadani|przedmiot|opis")]


def _vmerge_cont(cell):
    tcpr = cell._tc.tcPr
    if tcpr is None:
        return False
    vm = tcpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vMerge")
    return vm is not None and vm.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") in (None, "continue")


def _wpisz(cell, tekst):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    p.text = tekst
    if tekst:  # wzory maja czesto interlinie 2 i justowanie - wpis robi sie wtedy na kilka stron
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = 0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _mapuj(teksty, kolumny):
    mapa, uzyte = {}, set()
    for ci, tx in enumerate(teksty):
        for pole, rx in kolumny:
            if pole not in uzyte and re.search(rx, tx):
                mapa[ci] = pole
                uzyte.add(pole)
                break
    return mapa, uzyte


def wypelnij_docx(wzor, out, wiersze, kolumny):
    """Wpisuje wiersze w pierwsza tabele, ktorej naglowki da sie rozpoznac. Obsluguje naglowki wielopoziomowe,
    rekordy zlozone z kilku wierszy (komorki scalone pionowo, np. data / miejsce wykonania) i komorki scalone poziomo."""
    import copy
    import docx
    d = docx.Document(wzor)
    for t in d.tables:
        pusty = lambda r: not "".join(c.text for c in r.cells).strip()
        nag = []
        for r in t.rows:
            if pusty(r) or len(nag) == 3:
                break
            nag.append(r)
        if not nag:
            continue
        ncol = max(len(r.cells) for r in nag)
        kol_teksty = lambda r: [norm(" ".join(r.cells[ci].text.split())) if ci < len(r.cells) else "" for ci in range(ncol)]
        teksty = []
        for ci in range(ncol):
            cz = []
            for r in nag:
                x = kol_teksty(r)[ci]
                if x and x not in cz:
                    cz.append(x)
            teksty.append(" ".join(cz))
        mapa, uzyte = _mapuj(teksty, kolumny)
        if len(uzyte - {"lp"}) < 2:
            continue
        dane = [r for r in t.rows[len(nag):] if pusty(r)]
        if not dane:
            dane = [t.rows[-1]]
        # grupy wierszy jednego rekordu (pierwsza komorka scalona pionowo w kolejnych wierszach)
        grupy = []
        for r in dane:
            # python-docx zwraca dla komorki scalonej pionowo ten sam obiekt co w wierszu wyzej
            if grupy and r.cells and (r.cells[0]._tc is grupy[-1][-1].cells[0]._tc or _vmerge_cont(r.cells[0])):
                grupy[-1].append(r)
            else:
                grupy.append([r])
        # pola wierszy dodatkowych w grupie: z kolejnych wierszy naglowka (np. "miejsce wykonania")
        pod = {}
        if len(grupy[0]) > 1 and len(nag) > 1:
            for k in range(1, len(grupy[0])):
                hk = kol_teksty(nag[min(k, len(nag) - 1)])
                h0 = kol_teksty(nag[0])
                for ci, tx in enumerate(hk):
                    if tx == h0[ci]:
                        continue
                    for pole, rx in kolumny:
                        if re.search(rx, tx) and pole != mapa.get(ci):
                            pod[(k, ci)] = pole
                            break
        for i, w in enumerate(wiersze):
            if i < len(grupy):
                g = grupy[i]
            else:
                nowe = []
                for r in grupy[0]:
                    tr = copy.deepcopy(r._tr)
                    t._tbl.append(tr)
                    nowe.append(t.rows[-1])
                for r in nowe:
                    for c in r.cells:
                        _wpisz(c, "")
                g = nowe
            w_grupie = {}  # id -> element; trzymamy referencje, zeby id() proxy lxml nie zostalo uzyte ponownie
            for k, r in enumerate(g):
                wartosci = {}
                for ci, c in enumerate(r.cells):
                    if k > 0 and (id(c._tc) in w_grupie or _vmerge_cont(c)):
                        continue
                    pole = mapa.get(ci) if k == 0 else pod.get((k, ci))
                    if not pole:
                        continue
                    v = str(i + 1) if pole == "lp" else w.get(pole, DUZO)
                    lst = wartosci.setdefault(id(c._tc), (c, []))[1]
                    if v not in lst:
                        lst.append(v)
                for c, lst in wartosci.values():
                    znacz = [v for v in lst if v != "—"]
                    _wpisz(c, "; ".join(znacz) if znacz else "—")
                w_grupie.update({id(c._tc): c._tc for c in r.cells})
        d.save(out)
        return {f"k{k}:{teksty[k][:35]}": v for k, v in mapa.items()} | {f"wiersz{k}+k{ci}": v for (k, ci), v in pod.items()}
    raise SystemExit("Nie znalazłem tabeli z rozpoznawalnymi nagłówkami - wpisz ręcznie z wydruku MD (bez --form).")


def cmd_wykaz_uslug(a):
    R, _ = wczytaj(a.baza)
    ids = a.ids.split(",")
    rr = [next((r for r in R if r["id"] == i), None) for i in ids]
    if None in rr:
        sys.exit(f"Nieznane ID: {[i for i, r in zip(ids, rr) if r is None]}")
    wiersze = [wiersz_uslugi(r) for r in rr]
    if a.form:
        mapa = wypelnij_docx(a.form, a.out, wiersze, KOL_USLUGI)
        print(f"Wpisano {len(wiersze)} usług do {a.out}; kolumny: {mapa}")
    for r, w in zip(rr, wiersze):
        print(f"\n**{r['id']}**\n- przedmiot: {w['przedmiot']}\n- wartość brutto: {w['wartosc']}\n- data: {w['data']}\n- podmiot: {w['podmiot']}\n- dowód: {', '.join(r.get('dowody') or []) or DUZO}")
        if r.get("do_weryfikacji"):
            print(f"- ⚠ DO WERYFIKACJI: {'; '.join(r.get('rozbieznosci') or []) or r.get('podmiot') or 'pewność ' + str(r.get('pewnosc'))}"[:400])


def cmd_wykaz_osob(a):
    R, K = wczytaj(a.baza)
    od = date.fromisoformat(a.od) if a.od else date.today()
    wiersze = []
    for x in a.ids.split(","):
        i, _, rola = x.partition(":")
        p = next((k for k in K if k["id"] == i), None)
        if not p:
            sys.exit(f"Nieznane ID osoby: {i}")
        wiersze.append((p, wiersz_osoby(p, rola, od, R)))
    if a.form:
        mapa = wypelnij_docx(a.form, a.out, [w for _, w in wiersze], KOL_OSOBY)
        print(f"Wpisano {len(wiersze)} osób do {a.out}; kolumny: {mapa}")
    for p, w in wiersze:
        print(f"\n**{p['id']} {w['imie']}** - {w['rola']}\n- {w['kwalifikacje']}\n- doświadczenie: {w['doswiadczenie']}\n- podstawa: {w['podstawa']}")
        if p.get("do_weryfikacji"):
            print("- ⚠ DO WERYFIKACJI (numer/data uprawnień albo rozbieżności w źródłach)")


def cmd_xlsx(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    R, K = wczytaj(a.baza)
    wb = Workbook()
    ws = wb.active
    ws.title = "Referencje"
    nag = ["ID", "Etykieta", "Nazwa", "Zamawiający", "Inwestor końcowy", "Rola MiD", "Podmiot", "Zakres", "Obiekty", "Droga",
           "Wartość brutto", "Data rozpoczęcia", "Data zakończenia", "Projektanci", "Dowody", "Rozbieżności", "Do weryfikacji"]
    ws.append(nag)
    for r in R:
        dr = r.get("droga") or {}
        ws.append([r["id"], r.get("etykieta"), r.get("nazwa"), r.get("zamawiajacy_nazwa"), r.get("inwestor_koncowy"), r.get("rola_mid"),
                   r.get("podmiot"), ", ".join(r.get("zakres") or []), opis_obiektow(r), " ".join(str(dr.get(k) or "") for k in ("numer", "klasa")).strip(),
                   r.get("wartosc_brutto_pln"), r.get("data_rozpoczecia"), r.get("data_zakonczenia"),
                   "; ".join(f"{p.get('osoba')} ({p.get('rola')})" for p in r.get("projektanci") or []), "\n".join(r.get("dowody") or []),
                   "\n".join(r.get("rozbieznosci") or []), "TAK" if r.get("do_weryfikacji") else ""])
    ws2 = wb.create_sheet("Kadra")
    ws2.append(["ID", "Osoba", "Kategoria", "Uprawnienia", "Podstawa dysponowania", "Liczba dokumentacji", "Uwagi", "Do weryfikacji"])
    for p in K:
        ws2.append([p["id"], f"{p.get('tytul') or ''} {p['imie_nazwisko']}".strip(), p["kategoria"], "\n".join(upr_tekst(u) for u in p.get("uprawnienia") or []),
                    p.get("podstawa_dysponowania"), len(dokumentacje_osoby(p, R)), p.get("uwagi"), "TAK" if p.get("do_weryfikacji") else ""])
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        w.freeze_panes = "A2"
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 18
    ws.column_dimensions["C"].width = 70
    ws["K1"].number_format = '#,##0.00 "zł"'
    wb.save(a.out)
    print(f"Zapisano {a.out}: {len(R)} referencji, {len(K)} osób")


def main():
    import signal
    if hasattr(signal, 'SIGPIPE'):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)  # spokojne zakonczenie przy | head
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--baza", default=os.environ.get("MID_BAZA", "baza_mid"))
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("stan")
    s = sub.add_parser("uslugi")
    for x in ("--typ", "--zakres", "--dowolny-zakres", "--od", "--klasa-drogi", "--nad", "--szukaj"):
        s.add_argument(x)
    for x in ("--lat", "--min-wartosc", "--dlugosc-min", "--rozpietosc-min"):
        s.add_argument(x, type=float)
    s.add_argument("--netto", action="store_true"); s.add_argument("--tylko-spolka", action="store_true"); s.add_argument("--max", type=int, default=15)
    s = sub.add_parser("osoby")
    for x in ("--specjalnosc", "--od", "--typ"):
        s.add_argument(x)
    s.add_argument("--bez-ograniczen", action="store_true"); s.add_argument("--rownowazne", action="store_true"); s.add_argument("--wszyscy", action="store_true")
    s.add_argument("--lat-od-uprawnien", type=float); s.add_argument("--dokumentacje-min", type=int); s.add_argument("--lat", type=float)
    s.add_argument("--tylko-projekty", action="store_true")
    s = sub.add_parser("pokaz"); s.add_argument("id")
    s = sub.add_parser("doswiadczenie"); s.add_argument("id"); s.add_argument("--typ"); s.add_argument("--lat", type=float); s.add_argument("--od")
    s.add_argument("--form"); s.add_argument("--out", default="zal_1A.docx"); s.add_argument("--wybierz", help="numery pozycji z listy, np. 1,2,5")
    s.add_argument("--tylko-projekty", action="store_true", help="bez ekspertyz, przegladow i samego nadzoru")
    s = sub.add_parser("wykaz-uslug"); s.add_argument("ids"); s.add_argument("--form"); s.add_argument("--out", default="wykaz_uslug.docx")
    s = sub.add_parser("wykaz-osob"); s.add_argument("ids"); s.add_argument("--od"); s.add_argument("--form"); s.add_argument("--out", default="wykaz_osob.docx")
    s = sub.add_parser("xlsx"); s.add_argument("--out", required=True)
    a = p.parse_args()
    {"stan": cmd_stan, "uslugi": cmd_uslugi, "osoby": cmd_osoby, "pokaz": cmd_pokaz, "doswiadczenie": cmd_doswiadczenie,
     "wykaz-uslug": cmd_wykaz_uslug, "wykaz-osob": cmd_wykaz_osob, "xlsx": cmd_xlsx}[a.cmd](a)


if __name__ == "__main__":
    main()
