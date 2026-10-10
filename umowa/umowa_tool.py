#!/usr/bin/env python3
"""umowa_tool.py - przeglad umow dla Pracowni MiD (projektant: wykonawca, podwykonawca GW albo zlecajacy).

Komendy:
  struktura PLIK [PLIK..] [--json out.json]        klauzule z lokalizacja (§/ust./pkt/lit., subklauzule FIDIC) i strona
  analiza PLIK [PLIK..] [--wynagrodzenie NETTO] [--miesiace N] [--zamawiajacy publiczny|prywatny|duzy]
          [--rola wykonawca|podwykonawca|zlecajacy] [--standardy progi.json] [--md raport.md] [--json analiza.json]
        lista kontrolna: kary (stawki, podstawa, zwloka/opoznienie), limit kar, odszkodowanie uzupelniajace,
        ograniczenie odpowiedzialnosci, platnosci i odbior, platnosci czesciowe (art. 443), waloryzacja (art. 439, 436 pkt 4),
        zabezpieczenie (art. 452-453), prawa autorskie, AI, nadzor autorski, terminy zamawiajacego, decyzje organow,
        odstapienie, ograniczenie zakresu (art. 433 pkt 4), OC, rekojmia/gwarancja, potracenia, cesja
  kary analiza.json --wynagrodzenie NETTO [--dni 7,30,60,90]      ekspozycja na kary i dni do wyczerpania limitu
  komentarze UMOWA.docx --analiza analiza.json --out UMOWA_uwagi_MiD.docx [--poziom srednie|wysokie|info]
        komentarze Worda przy klauzulach (ocena, podstawa, propozycja zapisu)
  standardy --out standardy_umow.json            szablon progow i zasad MiD (do kalibracji, trzymac prywatnie)
  pytania analiza.json [--tryb pzp|negocjacje] [--md out.md] [--docx out.docx]
        pytania do zamawiajacego (przed terminem z art. 135/284 Pzp) albo tabela propozycji zmian do negocjacji

Podstawy prawne sprawdzone w tekstach jednolitych z API ELI (stan 10.2026): Pzp (Dz.U. 2026 poz. 793) art. 433, 436, 439,
443, 452, 453, 455; ustawa o przeciwdzialaniu nadmiernym opoznieniom w transakcjach handlowych (Dz.U. 2023 poz. 711)
art. 7, 8, 9; KC art. 473, 483, 484; prawo autorskie (Dz.U. 2025 poz. 24) art. 16, 41, 46, 53.
"""
import argparse, hashlib, json, os, re, subprocess, sys, unicodedata
from pathlib import Path

PROGI = {  # progi ostroznosciowe narzedzia - do kalibracji przez MiD (plik --standardy)
    "kara_dzienna_srednia": 0.2, "kara_dzienna_wysoka": 0.5,       # % podstawy za dzien
    "limit_kar_sredni": 20.0, "limit_kar_wysoki": 30.0,            # % wynagrodzenia
    "kara_jednorazowa_srednia": 10.0, "kara_jednorazowa_wysoka": 20.0,
    "zabezpieczenie_max": 5.0, "zatrzymanie_rekojmia_max": 30.0,
    "rekojmia_max_mies": 60, "platnosc_publiczny_dni": 30, "platnosc_prywatny_dni": 60, "odbior_max_dni": 30,
    "waloryzacja_limit_niski": 5.0, "nadzor_reakcja_min_dni": 3, "limit_odpowiedzialnosci_proponowany": 100,
}
OCENY = {"wysokie": "🔴", "srednie": "🟠", "niskie": "🟡", "ok": "🟢", "info": "ℹ️"}
OCENY_TXT = {"wysokie": "RYZYKO WYSOKIE", "srednie": "RYZYKO ŚREDNIE", "niskie": "UWAGA", "ok": "OK", "info": "INFORMACJA"}
POZIOM = {"wysokie": 0, "srednie": 1, "niskie": 2, "info": 3, "ok": 4}
CACHE = Path(os.environ.get("XDG_CACHE_HOME", Path.home() / ".cache")) / "umowa_tool"


# ----------------------------------------------------------------------------- tekst
def norm(s):
    s = unicodedata.normalize("NFKD", (s or "").lower())
    s = "".join(c for c in s if not unicodedata.combining(c)).replace("ł", "l")
    return re.sub(r"\s+", " ", s)


def n1(s):
    """normalizacja znak po znaku (ta sama dlugosc co oryginal) - pozycje dopasowan wskazuja oryginalny tekst"""
    out = []
    for c in s or "":
        b = unicodedata.normalize("NFKD", c)[:1].lower() or c
        out.append("l" if b in ("ł", "Ł") else b)
    return "".join(out)


def strony_pliku(p):
    p = Path(p)
    suf = p.suffix.lower()
    if suf in (".docx", ".doc", ".odt", ".rtf"):
        CACHE.mkdir(parents=True, exist_ok=True)
        h = hashlib.sha1(p.read_bytes()).hexdigest()[:16]
        pdf = CACHE / f"{h}.pdf"
        if not pdf.exists():
            out = CACHE / f"tmp_{h}"
            out.mkdir(exist_ok=True)
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(p)],
                           capture_output=True, timeout=300)
            pp = list(out.glob("*.pdf"))
            if not pp:
                sys.exit(f"Konwersja do PDF nieudana: {p}")
            pp[0].rename(pdf)
        p, suf = pdf, ".pdf"
    if suf == ".pdf":
        t = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout
        return t.split("\f")
    t = p.read_text(encoding="utf-8", errors="replace")
    if "[[s. " in t:
        cz = re.split(r"\[\[s\. (\d+)\]\]", t)
        return [cz[i + 1] for i in range(1, len(cz) - 1, 2)] if len(cz) > 2 else [t]
    return t.split("\f")


RX_PAR = re.compile(r"^§\s*(\d+[a-z]?)\b\.?\s*(.*)$")
RX_SUBKL = re.compile(r"^(?:Dodana\s+)?(?:Sub)?[Kk]lauzula\s+(\d{1,2}(?:\.\d{1,2}){0,2})\b\s*(.*)$")
RX_FIDIC = re.compile(r"^(\d{1,2}\.\d{1,2}(?:\.\d{1,2})?)\s+([A-ZŁŚŻŹĆŃÓĘĄ].{2,})$")
RX_UST = re.compile(r"^(\d{1,2})\.\s+(\S.*)$")
RX_UST2 = re.compile(r"^(\d{1,2}\.\d{1,2})\.?\s+(\S.*)$")
RX_PKT = re.compile(r"^(\d{1,2})\)\s*(\S.*)$")
RX_LIT = re.compile(r"^\(?([a-z]{1,2})\)\s*(\S.*)$|^([a-z])\.\s+([A-ZŁŚŻŹĆŃÓĘĄ].*)$")
RX_CAP = re.compile(r"^([A-H])[.)]\s+(\S.*)$")
RX_ROMAN = re.compile(r"^([IVX]{1,4})\.?\s+(\S.*)$")
RX_TIRET = re.compile(r"^[-–•▪◦]\s+(\S.*)$")


def stopki(strony):
    """linie powtarzajace sie na wielu stronach (naglowki/stopki) do pominiecia"""
    from collections import Counter
    c = Counter()
    for s in strony:
        widziane = set()
        for l in s.splitlines():
            k = re.sub(r"\d+", "#", l.strip())
            if 3 <= len(k) <= 160 and k not in widziane:
                widziane.add(k)
                c[k] += 1
    prog = max(3, int(0.3 * len(strony)))
    return {k for k, n in c.items() if n >= prog} | {"strona # z #", "- # -", "#"}


class Akapit:
    __slots__ = ("plik", "strona", "lok", "tekst", "rodzic", "tytul", "nr")

    def __init__(self, **kw):
        for k, v in kw.items():
            setattr(self, k, v)

    def d(self):
        return {k: getattr(self, k) for k in self.__slots__}


def akapity(plik, nazwa=None):
    strony = strony_pliku(plik)
    pomin = stopki(strony)
    nazwa = nazwa or Path(plik).name
    wyn = []
    loc = {"par": None, "sub": None, "ust": None, "pkt": None, "lit": None, "rom": None}
    tytul = ""
    rodzic = ""
    biez = None

    def lok():
        cz = []
        if loc["sub"]:
            cz.append(f"Subkl. {loc['sub']}")
        if loc["par"]:
            cz.append(f"§ {loc['par']}")
        if loc["rom"]:
            cz.append(f"pkt {loc['rom']}")
        if loc["ust"]:
            cz.append(f"ust. {loc['ust']}")
        if loc["pkt"]:
            cz.append(f"pkt {loc['pkt']}")
        if loc["lit"]:
            cz.append(f"lit. {loc['lit']})")
        return " ".join(cz) or "—"

    def nowy(tekst, strona):
        nonlocal biez
        biez = Akapit(plik=nazwa, strona=strona, lok=lok(), tekst=tekst, rodzic=rodzic, tytul=tytul, nr=len(wyn))
        wyn.append(biez)

    for si, s in enumerate(strony, 1):
        for surowa in s.splitlines():
            l = surowa.strip()
            if not l:
                biez = None if biez is None or biez.tekst.rstrip().endswith((".", ":", ";")) else biez
                continue
            if re.sub(r"\d+", "#", l) in pomin or re.fullmatch(r"(strona\s*)?\d+\s*(z|/)\s*\d+", l, re.I):
                continue
            m = RX_PAR.match(l)
            if m:
                loc.update(par=m.group(1), ust=None, pkt=None, lit=None, rom=None)
                tytul = m.group(2).strip()
                rodzic = ""
                if tytul:
                    nowy(l, si)
                else:
                    biez = None
                continue
            kont0 = biez is not None and re.search(
                r"\b(p|pkt|ust|lit|art|nr|poz|par|tj|tzw|ppkt|zob|por)\.?\s*$|[§(\-–]\s*$|\b(i|oraz|lub|albo|w|z|do|od|na|o|ze|zgodnie z|punkt\w*|podpunkt\w*|subklauzul\w*|klauzul\w*|paragraf\w*|ustaw\w*|art\w*)\s*$", biez.tekst, re.I)
            m = None if kont0 else (RX_SUBKL.match(l) or (RX_FIDIC.match(l) if len(l) < 90 and (biez is None or biez.tekst.rstrip().endswith((".", ":", ";"))) else None))
            if m:
                loc.update(sub=m.group(1), par=None, ust=None, pkt=None, lit=None, rom=None)
                tytul = m.group(2).strip()
                rodzic = ""
                nowy(l, si)
                continue
            kontynuacja = biez is not None and re.search(
                r"\b(p|pkt|ust|lit|art|nr|poz|par|tj|tzw|ppkt|zob|por)\.?\s*$|[§(\-–]\s*$|\b(i|oraz|lub|albo|w|z|do|od|na|o|ze|zgodnie z|punkt\w*|podpunkt\w*|subklauzul\w*|klauzul\w*|paragraf\w*|ustaw\w*|art\w*)\s*$", biez.tekst, re.I)
            for rx, klucz in ((RX_ROMAN, "rom"), (RX_UST2, "ust"), (RX_UST, "ust"), (RX_CAP, "pkt"), (RX_PKT, "pkt"), (RX_LIT, "lit")):
                m = None if kontynuacja else rx.match(l)
                if m and not (klucz == "ust" and re.match(r"^\d{1,2}\.\s+\d", l)) and not (klucz == "rom" and not loc["sub"] and not re.match(r"^[IVX]+\.\s", l)):
                    if klucz == "rom":
                        loc.update(rom=m.group(1), ust=None, pkt=None, lit=None)
                    elif klucz == "ust":
                        loc.update(ust=m.group(1), pkt=None, lit=None)
                    elif klucz == "pkt":
                        loc.update(pkt=m.group(1), lit=None)
                    else:
                        loc.update(lit=m.group(1) or m.group(3))
                    nowy(l, si)
                    if klucz in ("ust", "rom") or rx is RX_CAP:
                        rodzic = l
                    break
            else:
                if RX_TIRET.match(l) or biez is None:
                    nowy(l, si)
                elif len(l) < 60 and not re.search(r"[.,;:]$", l) and l[:1].isupper() and biez.tekst.rstrip().endswith((".", ":", ";")):
                    tytul = l
                    nowy(l, si)
                else:
                    biez.tekst += " " + l
                    if biez.lok.endswith(("ust. " + str(loc["ust"]), f"pkt {loc['rom']}")) and rodzic and biez.tekst.startswith(rodzic[:20]):
                        rodzic = biez.tekst
    for a in wyn:
        a.tekst = re.sub(r"\s+", " ", a.tekst).strip()
    return wyn


# ----------------------------------------------------------------------------- liczby
def proc(s):
    return [float(x.replace(",", ".")) for x in re.findall(r"(\d+(?:[,.]\d+)?)\s*%", s)]


def kwota(s):
    m = re.search(r"(\d{1,3}(?:[  .]\d{3})+|\d+)(?:,(\d{2}))?\s*(?:zł|zl|pln|złotych)", s, re.I)
    if not m:
        return None
    return float(re.sub(r"[  .]", "", m.group(1)) + "." + (m.group(2) or "00"))


def dni(s):
    m = re.search(r"(\d{1,3})\s*(dni|dzien|dnia)\s*(robocz\w*)?", norm(s))
    if not m:
        return None
    d = int(m.group(1))
    return d * 1.4 if m.group(3) else d, bool(m.group(3))


def cyt(t, n=320):
    t = re.sub(r"\s+", " ", t).strip()
    return t if len(t) <= n else t[:n - 1] + "…"


# ----------------------------------------------------------------------------- analiza
class Analiza:
    def __init__(self, pliki, a):
        self.ak = []
        for p in pliki:
            self.ak += akapity(p)
        self.a = a
        self.progi = dict(PROGI)
        self.zasady = []
        if a.standardy:
            st = json.loads(Path(a.standardy).read_text(encoding="utf-8"))
            self.progi.update(st.get("progi", {}))
            self.zasady = st.get("zasady", [])
        self.N = [n1(x.tekst) for x in self.ak]
        self.NR = [n1(x.rodzic) for x in self.ak]
        cal = " ".join(self.N)
        self.publiczny = (a.zamawiajacy == "publiczny") or (a.zamawiajacy is None and bool(
            re.search(r"zamowieni\w* publiczn|prawo zamowien|ustaw\w* pzp|\bpzp\b|art\.? ?(433|436|439|455)", cal)))
        self.miesiace = a.miesiace or self._miesiace(cal)
        self.decyzje = bool(re.search(r"\bzrid\b|decyzj\w* o (zezwoleniu|pozwoleniu|srodowiskow)|pozwoleni\w* na budowe|decyzj\w* srodowiskow", cal))
        self.f = []

    def _miesiace(self, cal):
        m = re.findall(r"(?:w terminie|termin\w* (?:wykonania|realizacji|zakonczenia)|do dnia)[^.]{0,80}?(\d{1,3})\s*(miesiec\w*|mies\.?|tygodni|dni)", cal)
        best = None
        for v, j in m:
            v = int(v)
            mies = v if j.startswith("mies") else (v / 4.33 if j.startswith("tyg") else v / 30.4)
            best = max(best or 0, mies)
        return round(best, 1) if best else None

    def dodaj(self, id_, temat, ocena, ak=None, ustalenie="", podstawa="", propozycja="", pytanie="", cytat=None, dane=None):
        self.f.append({"id": id_, "temat": temat, "ocena": ocena, "plik": ak.plik if ak else "", "strona": ak.strona if ak else None,
                       "lok": ak.lok if ak else "—", "cytat": cyt(cytat if cytat is not None else (ak.tekst if ak else "")),
                       "ustalenie": ustalenie, "podstawa": podstawa, "propozycja": propozycja, "pytanie": pytanie,
                       "akapit": ak.nr if ak else None, "dane": dane or {}})

    def gdzie(self, rx, z_rodzicem=False):
        r = re.compile(rx)
        return [i for i, t in enumerate(self.N) if r.search(t) or (z_rodzicem and r.search(self.NR[i]))]

    # --------------------------------------------------------------- kary
    def kary(self):
        kary = []
        for i, t in enumerate(self.N):
            ctx = t + " " + self.NR[i] + " " + norm(self.ak[i].tytul)
            if not re.search(r"\bkar(a|y|e|ami|om|ach)?\b|kar\w* umown", ctx) or "w wysokosci" not in t:
                continue
            if re.search(r"laczn\w*.{0,40}(wysokos|wartos|kwot)\w* kar", t):
                continue
            p = proc(self.ak[i].tekst)
            k = kwota(self.ak[i].tekst) if not p else None
            if not p and k is None:
                continue
            m = re.search(r"w wysokosci\s*(.*)", t)
            reszta = m.group(1) if m else t
            jedn = "dzień" if re.search(r"za kazdy\s*(rozpoczety\s*)?dzien|dziennie|za kazdy dzien", t) else (
                "przypadek" if re.search(r"za kazdy\s*(taki\s*)?(przypadek|stwierdzony|ujawniony)", t) else "jednorazowo")
            przyczyna = re.split(r"\s[–-]\s*w wysokosci|w wysokosci", t)[0]
            podst = "opóźnienie" if re.search(r"opozni", przyczyna + " " + reszta[:120]) and not re.search(r"zwlok", przyczyna + " " + reszta[:120]) else (
                "zwłoka" if re.search(r"zwlok", przyczyna + " " + reszta[:120]) else "")
            placi = "Zamawiający" if re.search(r"wykonawc\w* moze zadac od zamawiajac|zamawiajacy zaplaci wykonawcy", ctx) else "Wykonawca"
            baza = ""
            mb = re.search(r"%\s*(.{0,160}?)(?:,|;|\s–|\s-|\sza kazdy|\.$|$)", t[m.end(0) - len(m.group(1)):] if m else t)
            if mb:
                off = (m.end(0) - len(m.group(1)) if m else 0)
                baza = self.ak[i].tekst[off + mb.start(1): off + mb.end(1)]
            przyczyna_o = self.ak[i].tekst[:len(przyczyna)]
            przyczyna_o = re.sub(r"^\(?[a-z]{1,3}\)\s*|^\d+[.)]\s*", "", przyczyna_o)
            kary.append({"i": i, "stawka_proc": p[0] if p else None, "kwota": k, "jednostka": jedn, "podstawa": podst,
                         "placi": placi, "baza": baza.strip(), "przyczyna": przyczyna_o[:220].strip(), "_p": przyczyna, "_b": n1(baza)})
        return kary

    def sprawdz_kary(self):
        P = self.progi
        lista = self.kary()
        self.lista_kar = []
        for k in lista:
            ak = self.ak[k["i"]]
            bz = k["baza"] if len(k["baza"]) <= 90 else k["baza"][:90].rsplit(" ", 1)[0] + "…"
            opis = f"{fmt(k['stawka_proc'], 2).rstrip('0').rstrip(',')}% {bz}" if k["stawka_proc"] is not None else f"{fmt(k['kwota'], 2)} zł"
            opis += f" / {k['jednostka']}"
            rec = {**k, "lok": ak.lok, "strona": ak.strona, "plik": ak.plik, "opis": opis}
            self.lista_kar.append(rec)
            if k["placi"] != "Wykonawca":
                continue
            if re.search(r"niezaplacon|naleznosci (brutto|podwykonaw)|podwykonawc", k["_b"] + " " + k["_p"]):
                rec["podwykonawcy"] = True
                continue
            ocena, uw, prop = "info", [], ""
            if k["podstawa"] == "opóźnienie":
                ocena = "wysokie" if self.publiczny else "srednie"
                uw.append("kara za opóźnienie (także bez winy Wykonawcy)")
                prop = "Zastąpić „opóźnienie” słowem „zwłoka” (odpowiedzialność tylko za okoliczności, za które Wykonawca odpowiada)."
            if k["stawka_proc"] is not None and k["jednostka"] == "dzień":
                if k["stawka_proc"] >= P["kara_dzienna_wysoka"]:
                    ocena = "wysokie"
                    uw.append(f"stawka {fmt(k['stawka_proc'], 2).rstrip('0').rstrip(',')}% za dzień ≥ {fmt(P['kara_dzienna_wysoka'], 1)}%")
                elif k["stawka_proc"] >= P["kara_dzienna_srednia"]:
                    ocena = min(ocena, "srednie", key=POZIOM.get)
                    uw.append(f"stawka {fmt(k['stawka_proc'], 2).rstrip('0').rstrip(',')}% za dzień ≥ {fmt(P['kara_dzienna_srednia'], 1)}%")
                calosc = re.search(r"calosc|calkowit|umown\w*|zaakceptowan\w* kwot\w* kontraktow|okreslon\w* w § ?\d|wskazan\w* w § ?\d", k["_b"]) and not re.search(r"odcink|etap|czesc|poz\. |pozycj|dokument\w* wykonawc", k["_b"])
                czesc = re.search(r"poszczegoln\w* czesc|etap|czesci|termin\w* (posredni|czastkow)|kamien", k["_p"])
                if calosc and czesc and not re.search(r"za (dan\w* )?(czesc|etap)|wynagrodzeni\w* (za )?(dan\w* )?(czesc|etap)", k["_b"]):
                    ocena = min(ocena, "srednie", key=POZIOM.get)
                    uw.append("kara za zwłokę w części liczona od wynagrodzenia za całość")
                    prop = (prop + " " if prop else "") + "Karę za zwłokę w wykonaniu części liczyć od wynagrodzenia za tę część (etap), a nie od wynagrodzenia całkowitego."
                if re.search(r"usunieci\w* wad|usterek", k["_p"]):
                    uw.append("kara za zwłokę w usunięciu wad")
                    prop = (prop + " " if prop else "") + "Karę za zwłokę w usunięciu wad liczyć od wynagrodzenia za wadliwą część, z terminem usunięcia uzgadnianym ze stronami."
            elif k["stawka_proc"] is not None:
                if k["stawka_proc"] >= P["kara_jednorazowa_wysoka"] and not re.search(r"niezaplacon|naleznosci brutto", k["baza"] + k["_p"]):
                    ocena = min(ocena, "wysokie" if not re.search(r"odstapi", k["_p"]) else "srednie", key=POZIOM.get)
                    uw.append(f"kara jednorazowa {fmt(k['stawka_proc'], 1).rstrip('0').rstrip(',')}%")
                elif k["stawka_proc"] >= P["kara_jednorazowa_srednia"] and not re.search(r"niezaplacon", k["baza"] + k["_p"]):
                    ocena = min(ocena, "srednie", key=POZIOM.get)
                    uw.append(f"kara jednorazowa {fmt(k['stawka_proc'], 1).rstrip('0').rstrip(',')}%")
                if re.search(r"odstapi", k["_p"]):
                    prop = (prop + " " if prop else "") + "Karę za odstąpienie z przyczyn Wykonawcy ograniczyć (np. do 10% wynagrodzenia za niewykonaną część) i wprowadzić symetryczną karę za odstąpienie z przyczyn Zamawiającego."
                if re.search(r"osob\w* o mniejszym doswiadczeniu|zmian\w* (osoby|projektanta|personelu)", k["_p"]):
                    prop = (prop + " " if prop else "") + "Kara za zmianę osoby: dopuścić zmianę na osobę o doświadczeniu nie mniejszym niż wymagane w SWZ (nie w ofercie) lub z kryterium; stawka adekwatna do utraconych punktów."
            if re.search(r"w szczegolnosci|wszelkich obowiazkow|obowiazkow wynikajacych z umowy|niewykonani\w* (przez wykonawce )?obowiazkow", k["_p"]):
                ocena = min(ocena, "srednie", key=POZIOM.get)
                uw.append("otwarty katalog obowiązków zagrożonych karą")
                prop = (prop + " " if prop else "") + "Wskazać zamknięty katalog obowiązków, za które naliczana jest kara, i wymóg wezwania do usunięcia naruszenia z terminem."
            if uw or ocena != "info":
                self.dodaj("kara", "Kara umowna", ocena, ak, ustalenie=f"{opis}; " + "; ".join(uw),
                           podstawa=("art. 433 pkt 1–2 Pzp; " if self.publiczny else "") + "art. 483–484 KC (miarkowanie kary rażąco wygórowanej)",
                           propozycja=prop, pytanie="")
        if not lista:
            self.dodaj("kara", "Kary umowne", "info", None, ustalenie="nie znaleziono kar z określoną wysokością (sprawdź wzory w załącznikach/obrazach)")

    def sprawdz_limit(self):
        P = self.progi
        idx = self.gdzie(r"laczn\w*.{0,60}(maksymaln\w*.{0,20})?(wysokos|wartos|kwot|sum)\w*.{0,40}kar|kar\w* umown\w*.{0,120}nie (moze|moga|przekrocz)\w*.{0,40}%|limit\w* kar")
        if not idx:
            self.dodaj("limit_kar", "Łączny limit kar", "wysokie" if self.publiczny else "srednie", None,
                       ustalenie="nie znaleziono łącznego limitu kar umownych",
                       podstawa="art. 436 pkt 3 Pzp – umowa określa łączną maksymalną wysokość kar" if self.publiczny else "brak limitu – nieograniczona ekspozycja",
                       propozycja=f"Łączna wysokość kar umownych nie przekroczy {P['limit_kar_sredni']:.0f}% wynagrodzenia netto.",
                       pytanie="Prosimy o wskazanie łącznej maksymalnej wysokości kar umownych (art. 436 pkt 3 Pzp).")
            return
        idx = sorted(idx, key=lambda i: ("podwykonaw" in self.N[i], i))
        for i in idx[:4]:
            ak, t = self.ak[i], self.N[i]
            p = proc(ak.tekst)
            if not p:
                continue
            lim = p[0]
            ocena = "wysokie" if lim > P["limit_kar_wysoki"] else ("srednie" if lim > P["limit_kar_sredni"] else "ok")
            wyj = re.search(r"nie dotyczy|z wylaczeniem|nie obejmuje|nie wlicza|poza limitem", t)
            uw = f"limit {fmt(lim, 1).rstrip('0').rstrip(',')}%"
            if wyj:
                litery = re.findall(r"\b([a-z])\)", t[wyj.start():])
                dotyczy = [k for k in getattr(self, "lista_kar", []) if any(k["lok"].endswith(f"lit. {x})") for x in litery)]
                if dotyczy and all(k.get("podwykonawcy") for k in dotyczy):
                    uw += "; poza limitem tylko kary dot. zapłaty podwykonawcom"
                elif re.search(r"podwykonaw|439", t):
                    uw += "; poza limitem kary dot. podwykonawców"
                else:
                    uw += "; są wyjątki od limitu"
                    ocena = min(ocena if ocena != "ok" else "srednie", "srednie", key=POZIOM.get)
            self.limit_kar = lim
            self.dodaj("limit_kar", "Łączny limit kar", ocena, ak, ustalenie=uw,
                       podstawa="art. 436 pkt 3 Pzp" if self.publiczny else "",
                       propozycja=f"Obniżyć łączny limit kar do {P['limit_kar_sredni']:.0f}% wynagrodzenia netto i objąć nim wszystkie kary." if ocena != "ok" else "",
                       pytanie=f"Czy Zamawiający obniży łączny limit kar ({ak.lok}) do {P['limit_kar_sredni']:.0f}% wynagrodzenia netto?" if ocena != "ok" else "",
                       dane={"limit_proc": lim})
            break

    def sprawdz_odpowiedzialnosc(self):
        P = self.progi
        idx = self.gdzie(r"odszkodowani\w*.{0,60}(przewyzszajac|uzupelniajac|na zasadach ogolnych)|przenoszac\w*.{0,30}(wysokos\w*.{0,20})?kar")
        cap = self.gdzie(r"ogranicz\w* odpowiedzialnos\w*.{0,200}(do|nie przekr)|odpowiedzialnos\w*.{0,80}(ograniczon\w*|nie przekroczy|do wysokosci|do kwoty)")
        usun = self.gdzie(r"(1\.15|ograniczeni\w* odpowiedzialnosci).{0,160}(skresla|usuwa|usunieto|skreslono|wykreslono|nie stosuje|wykresla)")
        if idx:
            i = idx[0]
            lucrum = "utraconych korzysci" in self.N[i] or "utracone korzysci" in self.N[i]
            ocena = "wysokie" if (not cap or usun) else "srednie"
            self.dodaj("odszkodowanie", "Odszkodowanie ponad kary", ocena, self.ak[i],
                       ustalenie="odszkodowanie uzupełniające ponad kary" + ("; obejmuje utracone korzyści" if lucrum else "") + ("; brak ograniczenia odpowiedzialności" if not cap else ""),
                       podstawa="art. 484 § 1 KC (odszkodowanie ponad karę tylko gdy strony tak postanowiły)",
                       propozycja=f"Łączna odpowiedzialność Wykonawcy (kary i odszkodowania) ograniczona do {P['limit_odpowiedzialnosci_proponowany']}% wynagrodzenia netto, z wyłączeniem utraconych korzyści; nie dotyczy szkody wyrządzonej umyślnie (art. 473 § 2 KC).",
                       pytanie="Czy Zamawiający wprowadzi ograniczenie łącznej odpowiedzialności Wykonawcy do wysokości wynagrodzenia netto i wyłączy utracone korzyści?")
        if usun:
            self.dodaj("limit_odpowiedzialnosci", "Ograniczenie odpowiedzialności", "wysokie", self.ak[usun[0]],
                       ustalenie="klauzula ograniczenia odpowiedzialności usunięta lub wyłączona",
                       propozycja="Przywrócić ograniczenie odpowiedzialności (np. Subklauzula 1.15 FIDIC) do wysokości wynagrodzenia.")
        elif not cap:
            self.dodaj("limit_odpowiedzialnosci", "Ograniczenie odpowiedzialności", "srednie", None,
                       ustalenie="nie znaleziono ograniczenia łącznej odpowiedzialności Wykonawcy",
                       propozycja=f"Dodać: „Łączna odpowiedzialność Wykonawcy z tytułu umowy nie przekroczy {P['limit_odpowiedzialnosci_proponowany']}% wynagrodzenia netto, z wyłączeniem szkody wyrządzonej umyślnie.”")
        bez_winy = self.gdzie(r"niezaleznie od przyczyn|bez wzgledu na przyczyn|niezaleznie od winy|bez wzgledu na win")
        for i in bez_winy[:2]:
            if "kar" in self.N[i] or "odpowiad" in self.N[i]:
                self.dodaj("bez_winy", "Odpowiedzialność niezależna od winy", "wysokie", self.ak[i],
                           ustalenie="odpowiedzialność/kara niezależnie od przyczyn",
                           podstawa="art. 433 pkt 3 Pzp" if self.publiczny else "art. 473 § 1 KC – rozszerzenie odpowiedzialności",
                           propozycja="Ograniczyć do okoliczności, za które Wykonawca ponosi odpowiedzialność.",
                           pytanie=f"Czy Zamawiający usunie z {self.ak[i].lok} odpowiedzialność Wykonawcy niezależną od przyczyn?")
        for i in self.gdzie(r"\bopoznieni\w*|\bopozni\w* sie"):
            t = self.N[i]
            if "zwlok" in t or re.search(r"odsetk|transakcjach handlowych|spowodowan\w* przez (zamawiajac|wladz|organ)|po stronie zamawiajac|z przyczyn (lezacych po stronie )?zamawiajac|opoznieni\w* w platnosci|zamawiajacy nie ponosi", t):
                continue
            if re.search(r"dozna\w* opozni|uprawnion\w*|przedluzeni\w* czasu|prawo do przedluzenia", t):
                continue
            if re.search(r"wykonawc\w*.{0,80}opozni|opoznieni\w*.{0,60}(w stosunku do terminu|wykonani|realizacj)|z opoznieniem", t) and \
                    re.search(r"kar|odpowiedzial|nienalezyt|odstap|odszkodow|potrac", t + " " + self.NR[i]):
                self.dodaj("opoznienie", "Opóźnienie zamiast zwłoki", "wysokie" if self.publiczny else "srednie", self.ak[i],
                           ustalenie="skutki dla Wykonawcy powiązane z opóźnieniem (bez względu na winę)",
                           podstawa="art. 433 pkt 1 Pzp" if self.publiczny else "art. 476 KC (zwłoka) vs opóźnienie",
                           propozycja="Zastąpić „opóźnienie” słowem „zwłoka”.",
                           pytanie=f"Czy Zamawiający zastąpi w {self.ak[i].lok} „opóźnienie” słowem „zwłoka” (art. 433 pkt 1 Pzp)?" if self.publiczny else "")
                break

    def sprawdz_platnosci(self):
        P = self.progi
        found = False
        for i in self.gdzie(r"(zaplat|platn|przelew|wynagrodzeni)\w*.{0,140}w terminie (do )?\d{1,3} dni.{0,80}(faktur|rachun)|w terminie (do )?\d{1,3} dni.{0,80}(od dnia )?(doreczeni|otrzymani|zlozeni|wplywu)\w*.{0,60}(faktur|rachun)"):
            d = dni(self.ak[i].tekst)
            if not d:
                continue
            found = True
            d0 = d[0]
            lim = P["platnosc_publiczny_dni"] if self.publiczny else P["platnosc_prywatny_dni"]
            ocena = "wysokie" if d0 > lim else "ok"
            od = "od zatwierdzenia/akceptacji" if re.search(r"od (dnia )?(zatwierdzeni|akceptac|odbioru|podpisani\w* protokol)", self.N[i]) else ""
            self.dodaj("platnosc", "Termin zapłaty", ocena, self.ak[i], ustalenie=f"{d0:.0f} dni" + (" roboczych" if d[1] else "") + (f"; liczony {od}" if od else ""),
                       podstawa=("art. 8 ust. 2 ustawy o przeciwdziałaniu nadmiernym opóźnieniom (podmiot publiczny: max 30 dni)" if self.publiczny
                                 else "art. 7 ust. 2–2a ustawy o przeciwdziałaniu nadmiernym opóźnieniom (max 60 dni; duży przedsiębiorca wobec MŚP – bezwzględnie)"),
                       propozycja=f"Termin zapłaty {lim} dni od doręczenia prawidłowo wystawionej faktury." if ocena != "ok" else "")
            break
        if not found:
            self.dodaj("platnosc", "Termin zapłaty", "srednie", None, ustalenie="nie znaleziono terminu zapłaty faktury",
                       podstawa="art. 436 pkt 2 Pzp (warunki zapłaty); ustawowo max 30 dni od doręczenia faktury (art. 8 ust. 2 ustawy o przeciwdziałaniu nadmiernym opóźnieniom)" if self.publiczny else "art. 7 ust. 2–2a ustawy o przeciwdziałaniu nadmiernym opóźnieniom",
                       propozycja="Termin zapłaty 30 dni od doręczenia prawidłowo wystawionej faktury.",
                       pytanie="Prosimy o wskazanie terminu zapłaty faktur (art. 436 pkt 2 Pzp)." if self.publiczny else "")
        for i in self.gdzie(r"(odbior|weryfikac|sprawdzen|akceptac|zatwierdz|zaopiniowa|ocen)\w*.{0,140}w terminie (do )?\d{1,3} dni|w terminie (do )?\d{1,3} dni.{0,100}(odbior|protokol\w* odbior|weryfikac|zatwierdz|zgloszeni\w* uwag)"):
            d = dni(self.ak[i].tekst)
            if not d or not re.search(r"zamawiajac|inzynier|inspektor", self.N[i]):
                continue
            if d[0] > P["odbior_max_dni"]:
                self.dodaj("odbior", "Termin odbioru / weryfikacji przez zamawiającego", "srednie", self.ak[i],
                           ustalenie=f"{d[0]:.0f} dni" + (" (dni robocze przeliczone ×1,4)" if d[1] else ""),
                           podstawa="art. 9 ust. 1 ustawy o przeciwdziałaniu nadmiernym opóźnieniom: badanie usługi ≤ 30 dni",
                           propozycja="Odbiór/weryfikacja w terminie do 14 dni (łącznie max 30 dni); brak uwag w terminie oznacza odbiór bez zastrzeżeń; czas weryfikacji nie wlicza się do terminu Wykonawcy.",
                           pytanie=f"Czy Zamawiający skróci termin odbioru w {self.ak[i].lok} do 30 dni kalendarzowych (art. 9 ust. 1 ustawy o przeciwdziałaniu nadmiernym opóźnieniom)?")
                break
        # platnosci czesciowe
        if self.miesiace and self.miesiace > 12:
            cz = self.gdzie(r"(platnos|faktur|wynagrodzeni|rozliczeni)\w* (czesciow|przejsciow)|przejsciow\w* dokument\w* rozliczeniow|za (poszczegolne|kazdy) etap|po wykonaniu (kazdego )?etapu|zaliczk")
            if not cz:
                self.dodaj("platnosci_czesciowe", "Płatności częściowe", "wysokie" if self.publiczny else "srednie", None,
                           ustalenie=f"umowa na ok. {self.miesiace:g} mies. bez płatności częściowych ani zaliczek",
                           podstawa="art. 443 ust. 1 Pzp (umowy > 12 mies.)" if self.publiczny else "finansowanie prac przez Wykonawcę",
                           propozycja="Płatności częściowe po odbiorze etapów; ostatnia część ≤ 50% wynagrodzenia.")
            for i in self.gdzie(r"(koncow|ostatni)\w* (czesc|faktur|platnos|rat)\w*.{0,100}\d+ ?%|\d+ ?%.{0,60}(koncow|ostatni)\w* (czesc|faktur|platnos)"):
                p = [x for x in proc(self.ak[i].tekst) if x > 50]
                if p:
                    self.dodaj("platnosci_czesciowe", "Ostatnia część wynagrodzenia", "wysokie" if self.publiczny else "srednie", self.ak[i],
                               ustalenie=f"ostatnia część {p[0]:g}% wynagrodzenia", podstawa="art. 443 ust. 2 Pzp: ostatnia część ≤ 50%",
                               propozycja="Ostatnia część wynagrodzenia nie więcej niż 50%.")
                    break

    def sprawdz_waloryzacje(self):
        P = self.progi
        idx = self.gdzie(r"waloryzac|art\.? ?439|zmian\w* cen\w* materialow lub kosztow|wskaznik\w* (zmiany )?cen")
        if self.miesiace and self.miesiace > 6 and not idx:
            self.dodaj("waloryzacja", "Waloryzacja", "wysokie", None, ustalenie=f"umowa ok. {self.miesiace:g} mies. bez klauzuli waloryzacyjnej",
                       podstawa="art. 439 ust. 1 Pzp (umowy > 6 mies.)" if self.publiczny else "ryzyko wzrostu kosztów",
                       propozycja="Waloryzacja wg wskaźnika GUS (np. przeciętne wynagrodzenie w sektorze przedsiębiorstw), od 6. miesiąca, przy zmianie > 3%, limit ≥ 10%.",
                       pytanie="Czy Zamawiający wprowadzi klauzulę waloryzacyjną zgodną z art. 439 Pzp?" if self.publiczny else "")
            return
        if not idx:
            return
        lim = None
        kary_i = {k["i"] for k in getattr(self, "lista_kar", [])}
        for i in idx:
            t = self.N[i]
            if i in kary_i or re.search(r"\bkar[aey]?\b|kar\w* umown", t):
                continue
            if re.search(r"(nie przekroczy|maksymaln\w*|limit\w*|lacznie)", t) and proc(self.ak[i].tekst):
                lim = (i, max(proc(self.ak[i].tekst)))
                break
        ocena = "ok"
        uw = "klauzula waloryzacyjna jest"
        if lim:
            uw += f"; limit zmian {fmt(lim[1], 2).rstrip('0').rstrip(',')}%"
            if lim[1] <= P["waloryzacja_limit_niski"]:
                ocena = "srednie"
        wsk = [self.ak[i].tekst for i in idx if re.search(r"gus|wskaznik", self.N[i])]
        start = [re.search(r"od (\d+)\.? ?(miesiac|mies)", self.N[i]) for i in idx]
        start = [int(m.group(1)) for m in start if m]
        if start:
            uw += f"; od {min(start)}. miesiąca"
        ak = self.ak[lim[0]] if lim else self.ak[idx[0]]
        self.dodaj("waloryzacja", "Waloryzacja", ocena, ak, ustalenie=uw + ("; wskaźniki GUS" if wsk else ""),
                   podstawa="art. 439 ust. 2 pkt 4 Pzp (maksymalna wartość zmiany)" if self.publiczny else "",
                   propozycja="Podnieść limit waloryzacji (np. do 10–15% wynagrodzenia) i dobrać wskaźnik do kosztów pracy biura projektowego (wynagrodzenia)." if ocena != "ok" else "",
                   pytanie=f"Czy Zamawiający zwiększy maksymalną wartość zmiany wynagrodzenia z tytułu waloryzacji ({ak.lok}) do 10% wynagrodzenia netto?" if ocena != "ok" and self.publiczny else "")
        if self.miesiace and self.miesiace > 12 and self.publiczny:
            if not self.gdzie(r"minimaln\w* wynagrodzeni\w* za prace|stawk\w* podatku od towarow|ubezpieczen\w* spoleczn|pracownicz\w* plan\w* kapitalow"):
                self.dodaj("436_4b", "Zmiana wynagrodzenia (VAT, płaca minimalna, ZUS, PPK)", "srednie", None,
                           ustalenie="nie znaleziono zasad zmiany wynagrodzenia przy zmianie VAT, płacy minimalnej, składek ZUS/NFZ, PPK",
                           podstawa="art. 436 pkt 4 lit. b Pzp (umowy > 12 mies.)",
                           pytanie="Prosimy o uzupełnienie wzoru umowy o zasady zmiany wynagrodzenia, o których mowa w art. 436 pkt 4 lit. b Pzp.")

    def sprawdz_zabezpieczenie(self):
        P = self.progi
        for i in self.gdzie(r"zabezpieczeni\w* nalezyt\w*"):
            p = proc(self.ak[i].tekst)
            t = self.N[i]
            if not p:
                continue
            if re.search(r"rekojm|gwarancj", t) and re.search(r"pozostaw|zatrzym|30 ?%|70 ?%", t):
                zat = [x for x in p if x <= 100 and re.search(rf"{str(x).rstrip('0').rstrip('.').replace('.', ',')} ?%", self.ak[i].tekst)]
                if any(x > P["zatrzymanie_rekojmia_max"] and x < 100 and "70" not in str(x) for x in zat):
                    self.dodaj("zabezpieczenie", "Zabezpieczenie – część na rękojmię", "wysokie" if self.publiczny else "srednie", self.ak[i],
                               ustalenie="pozostawiona na rękojmię część zabezpieczenia > 30%", podstawa="art. 453 ust. 2 Pzp")
                continue
            v = p[0]
            if v <= 0 or v > 20:
                continue
            ocena = "wysokie" if v > 10 else ("srednie" if v > P["zabezpieczenie_max"] else "ok")
            self.dodaj("zabezpieczenie", "Zabezpieczenie należytego wykonania", ocena, self.ak[i], ustalenie=f"{v:g}% ceny",
                       podstawa="art. 452 ust. 2–3 Pzp: do 5%, do 10% tylko gdy uzasadnione i opisane w SWZ" if self.publiczny else "",
                       propozycja="Zabezpieczenie 5% ceny; zwrot 70% po odbiorze, 30% po rękojmi." if ocena != "ok" else "")
            break

    def sprawdz_prawa_autorskie(self):
        idx = self.gdzie(r"autorsk\w* praw\w* majatkow|majatkow\w* praw\w* autorsk|przenies\w* .{0,40}praw\w* autorsk")
        if not idx:
            self.dodaj("prawa_autorskie", "Prawa autorskie", "info", None, ustalenie="nie znaleziono postanowień o przeniesieniu autorskich praw majątkowych")
            return
        moment = None
        for i in idx + self.gdzie(r"z chwil\w*.{0,60}(przechodz|przenies|nabyw)|(przechodz|przenies|nabyw|nabyci)\w*.{0,80}(z chwil|nastepuje)", z_rodzicem=True):
            t = self.N[i]
            m = re.search(r"z chwil\w* (\w+ ){0,4}?(zaplat|wyplat|uregulowan|zaplacen|odbior|przekazan|wydani|dostarczen|utrwalen|sporzadzen|podpisani|przyjeci)\w*", t)
            if m:
                moment = (i, m.group(2))
                break
        if moment:
            i, mm = moment
            if mm in ("zaplat", "wyplat", "uregulowan", "zaplacen"):
                self.dodaj("prawa_autorskie", "Prawa autorskie – moment przejścia", "ok", self.ak[i], ustalenie="z chwilą zapłaty")
            else:
                ocena = "srednie" if mm in ("przekazan", "wydani", "dostarczen", "utrwalen", "sporzadzen", "podpisani", "przyjeci") else "niskie"
                self.dodaj("prawa_autorskie", "Prawa autorskie – moment przejścia", ocena, self.ak[i],
                           ustalenie=f"prawa przechodzą z chwilą „{mm}…” – przed zapłatą",
                           propozycja="Przeniesienie autorskich praw majątkowych z chwilą zapłaty wynagrodzenia za daną część; do tego czasu licencja na korzystanie w zakresie niezbędnym do realizacji inwestycji.",
                           pytanie=f"Czy Zamawiający dopuści przejście autorskich praw majątkowych z chwilą zapłaty wynagrodzenia za daną część ({self.ak[i].lok})?" if self.publiczny else "")
        for i in self.gdzie(r"zrzek\w*.{0,60}(osobist|praw\w* autorsk\w* osobist)|zbyw\w*.{0,40}praw\w* osobist"):
            self.dodaj("prawa_osobiste", "Autorskie prawa osobiste", "wysokie", self.ak[i],
                       ustalenie="zrzeczenie się / zbycie praw osobistych", podstawa="art. 16 ustawy o prawie autorskim – prawa osobiste nie podlegają zrzeczeniu ani zbyciu",
                       propozycja="Zastąpić zobowiązaniem do niewykonywania praw osobistych w zakresie niezbędnym do realizacji i utrzymania obiektu, z zachowaniem prawa do autorstwa.")
            break
        for i in self.gdzie(r"(zgod\w*|upowazni\w*|zezwal\w*).{0,60}(wykonywani\w* )?(autorsk\w* )?praw\w* (osobist|zalezn)|praw\w* zalezn\w*.{0,60}(przenos|zezwal|zgod)"):
            self.dodaj("prawa_zalezne", "Prawa zależne / wykonywanie praw osobistych", "niskie", self.ak[i],
                       ustalenie="zgoda na zmiany i opracowania przez osoby trzecie",
                       podstawa="art. 46 ustawy o prawie autorskim",
                       propozycja="Zastrzec, że Wykonawca nie odpowiada za zmiany wprowadzone przez osoby trzecie, a oznaczenie autorstwa obejmuje tylko wersję Wykonawcy.")
            break
        if self.gdzie(r"wszystkich (znanych )?pol\w* eksploatacji") and not self.gdzie(r"pol\w* eksploatacji.{0,40}(obejmuj|w szczegolnosci|:)|na nastepujacych polach"):
            self.dodaj("pola_eksploatacji", "Pola eksploatacji", "niskie", None, ustalenie="„wszystkie pola eksploatacji” bez wyliczenia",
                       podstawa="art. 41 ust. 2 ustawy o prawie autorskim – pola muszą być wyraźnie wymienione")

    def sprawdz_ai(self):
        for i in self.gdzie(r"sztuczn\w* inteligencj|\bsystem\w* ai\b|\bnarzedzi\w* ai\b|\bai\b.{0,40}(generat|wytworz)|wytwor\w* ai|2024/1689"):
            zak = re.search(r"zakaz|nie (moze|wolno)|wylacznie pomocnicz|nie jest wytworem|oswiadcza|gwarant", self.N[i])
            self.dodaj("ai", "Sztuczna inteligencja", "srednie" if zak else "info", self.ak[i],
                       ustalenie="ograniczenia/oświadczenia dotyczące AI" if zak else "wzmianka o AI",
                       propozycja="Dopuścić narzędzia AI wspomagające (obliczenia, redakcja, kontrola) pod nadzorem projektanta z uprawnieniami; oświadczenie o autorstwie dotyczy rozwiązań projektowych.",
                       pytanie=f"Czy Zamawiający potwierdzi, że użycie narzędzi wspomagających (w tym AI) pod nadzorem projektanta nie narusza {self.ak[i].lok}?" if self.publiczny and zak else "")
            break

    def sprawdz_nadzor(self):
        P = self.progi
        idx = self.gdzie(r"nadzor\w* autorsk")
        if not idx:
            return
        cal = " ".join(self.N[i] for i in idx)
        uw, ocena, ak = [], "ok", self.ak[idx[0]]
        if re.search(r"w ramach wynagrodzenia|ryczalt|wynagrodzeni\w* .{0,40}obejmuje", cal) and not re.search(r"\d+ (pobyt|wizyt|nadzor|wyjazd)|liczb\w* (pobyt|wizyt)|za (kazdy )?pobyt|stawk\w* za", cal):
            ocena = "srednie"
            uw.append("nadzór autorski w ryczałcie bez limitu liczby pobytów")
        for i in idx:
            m = re.search(r"w (ciagu|terminie) (do )?(\d{1,3}) (godzin|dni)", self.N[i])
            if m:
                v = int(m.group(3)) / (24 if m.group(4) == "godzin" else 1)
                if v < P["nadzor_reakcja_min_dni"]:
                    ocena = "srednie"
                    uw.append(f"czas reakcji {m.group(3)} {m.group(4)}")
                    ak = self.ak[i]
                break
        if re.search(r"do (czasu|dnia) (zakonczenia|odbioru|uzyskania)\w* .{0,40}(robot|budowy|inwestycji|pozwolenia na uzytkowanie)", cal):
            uw.append("okres nadzoru zależny od robót (bez daty granicznej)")
            ocena = "srednie"
        self.dodaj("nadzor", "Nadzór autorski", ocena, ak, ustalenie="; ".join(uw) or "postanowienia o nadzorze autorskim",
                   propozycja="Limit pobytów w cenie (np. N pobytów) i stawka za każdy kolejny; czas reakcji ≥ 3 dni robocze; nadzór do określonej daty, potem za dodatkowym wynagrodzeniem." if ocena != "ok" else "",
                   pytanie="Prosimy o wskazanie szacowanej liczby pobytów w ramach nadzoru autorskiego i okresu jego pełnienia." if ocena != "ok" and self.publiczny else "")

    def sprawdz_terminy_zamawiajacego(self):
        rev = self.gdzie(r"zamawiajac\w*.{0,160}(uwag|zastrzezen|zatwierdz|zaopiniow|weryfik|odbior)\w*.{0,100}w (terminie|ciagu) (do )?\d{1,3} dni|w terminie (do )?\d{1,3} dni.{0,120}(zglos|przekaz)\w* uwag|inzynier\w*.{0,100}(przeglad|uwag)\w*.{0,60}\d{1,3} dni")
        wylacz = self.gdzie(r"(nie wlicza|wydluza|przedluz)\w*.{0,80}(termin)\w*.{0,120}(weryfik|uwag|zatwierdz|opini|uzgodni)|czas\w* (weryfikacji|sprawdzenia|zatwierdzania).{0,80}nie wlicza")
        if not rev:
            self.dodaj("terminy_zamawiajacego", "Terminy działań zamawiającego", "srednie", None,
                       ustalenie="nie znaleziono terminów na uwagi/zatwierdzenie dokumentacji przez zamawiającego",
                       propozycja="Zamawiający zgłasza uwagi w terminie 14 dni; brak uwag w terminie = akceptacja; czas weryfikacji nie wlicza się do terminów Wykonawcy.",
                       pytanie="Prosimy o określenie terminów weryfikacji dokumentacji przez Zamawiającego i potwierdzenie, że czas weryfikacji nie wlicza się do terminów Wykonawcy." if self.publiczny else "")
        elif not wylacz:
            self.dodaj("terminy_zamawiajacego", "Czas weryfikacji a terminy Wykonawcy", "niskie", self.ak[rev[0]],
                       ustalenie="są terminy weryfikacji, ale nie znaleziono wyłączenia czasu weryfikacji z terminów Wykonawcy",
                       propozycja="Czas weryfikacji przez Zamawiającego nie wlicza się do terminów Wykonawcy.")

    def okno(self, i, n=4):
        return " ".join(self.N[max(0, i - n):i + 1])

    def sprawdz_organy(self):
        if not self.decyzje:
            return
        idx = [i for i in self.gdzie(r"\borgan(u|ow|y|ami|ach|em)?\b|decyzj|uzgodnie|administracyjn|wladz|gestor")
               if re.search(r"przedluz\w*.{0,40}termin|termin\w*.{0,60}(ulega|moze ulec|zostanie) (przedluz|zmian)|wydluz\w*.{0,30}termin", self.okno(i))]
        idx = idx or self.gdzie(r"(przedluz|zmian)\w*.{0,30}termin\w*.{0,250}(organ|decyzj|uzgodni|administracyjn|opini\w*|wladz|osob\w* trzeci|gestor)|(organ|decyzj|wladz|administracyjn)\w*.{0,150}(przedluz|zmian)\w*.{0,30}termin", z_rodzicem=True)
        if idx:
            self.dodaj("organy", "Terminy a postępowania przed organami", "ok", self.ak[idx[0]],
                       ustalenie="przewidziano zmianę terminu przy przewlekłości organów / osób trzecich")
        else:
            self.dodaj("organy", "Terminy a postępowania przed organami", "wysokie", None,
                       ustalenie="zakres obejmuje decyzje (ZRID/PnB/DŚU), a nie znaleziono zmiany terminu przy przewlekłości organów",
                       podstawa="art. 455 ust. 1 pkt 1 Pzp" if self.publiczny else "",
                       propozycja="Termin wykonania ulega przedłużeniu o czas postępowań przed organami i gestorami ponad terminy ustawowe oraz o czas oczekiwania na dane od Zamawiającego.",
                       pytanie="Czy Zamawiający przewiduje przedłużenie terminów o czas postępowań administracyjnych przekraczający terminy ustawowe?" if self.publiczny else "")

    def sprawdz_odstapienie(self):
        idx = self.gdzie(r"odstapi\w*|wypowiedz\w* umow")
        if not idx:
            return
        for i in idx:
            t = self.N[i]
            if re.search(r"(nie przysluguj|bez prawa do|traci prawo do)\w*.{0,40}wynagrodzeni", t):
                self.dodaj("odstapienie", "Odstąpienie – wynagrodzenie", "wysokie", self.ak[i],
                           ustalenie="brak wynagrodzenia za wykonaną część przy odstąpieniu",
                           propozycja="Przy odstąpieniu Wykonawcy przysługuje wynagrodzenie za część wykonaną do dnia odstąpienia.")
                break
        for i in idx:
            m = re.search(r"(zwlok|opozni|przerw|nie (rozpoczal|podjal))\w*.{0,80}(przekracz\w*|dluzsz\w* niz|powyzej) (\d{1,3}) dni", self.N[i])
            if m and int(m.group(4)) < 14:
                self.dodaj("odstapienie", "Odstąpienie – krótki próg", "srednie", self.ak[i], ustalenie=f"odstąpienie już po {m.group(4)} dniach",
                           propozycja="Odstąpienie po bezskutecznym wezwaniu z terminem min. 14 dni.")
                break
        if not self.gdzie(r"wykonawc\w* (moze|ma prawo|przysluguje|jest uprawnion\w*).{0,60}(odstap|wypowiedz)|wykonawcy przysluguje prawo odstapienia"):
            self.dodaj("odstapienie_wykonawcy", "Odstąpienie przez Wykonawcę", "niskie", None,
                       ustalenie="nie znaleziono prawa Wykonawcy do odstąpienia (np. przy zwłoce w zapłacie, braku współdziałania)",
                       propozycja="Prawo odstąpienia przez Wykonawcę przy zwłoce w zapłacie > 30 dni lub braku współdziałania Zamawiającego mimo wezwania; wynagrodzenie za wykonaną część i kara symetryczna.")

    def sprawdz_zakres(self):
        idx = self.gdzie(r"(ogranicz|zmniejsz|rezygn|wylacz)\w*.{0,60}(zakres|czesc|przedmiot)\w*.{0,80}(zamowieni|umowy|prac|przedmiotu)|rzeczywiscie wykonan\w* (zakres|prac)|za (rzeczywiscie )?wykonany zakres")
        if not idx:
            return
        minimum = [i for i in self.gdzie(r"minimaln\w* (wartos|wielkos|zakres|wynagrodzeni\w* (wykonawcy|umown))|nie mniej niz \d+ ?%|co najmniej \d+ ?% (wynagrodzeni|wartosci|zakresu)")
                   if not re.search(r"minimaln\w* wynagrodzeni\w* za prace|minimalnej stawki godzinowej", self.N[i])]
        jawne = [i for i in idx if re.search(r"(ogranicz|zmniejsz|rezygn|wylacz)", self.N[i])]
        if minimum:
            self.dodaj("ograniczenie_zakresu", "Ograniczenie zakresu", "ok", self.ak[minimum[0]], ustalenie="określono minimalny zakres / wartość")
        else:
            i0 = jawne[0] if jawne else idx[0]
            self.dodaj("ograniczenie_zakresu", "Ograniczenie zakresu", ("wysokie" if self.publiczny else "srednie") if jawne else "srednie", self.ak[i0],
                       ustalenie=("możliwe ograniczenie zakresu bez minimalnej wartości świadczenia" if jawne else
                                  "wynagrodzenie za rzeczywiście wykonany zakres (do kwoty) bez minimalnej wartości – Zamawiający może zlecić mniej"),
                       podstawa="art. 433 pkt 4 Pzp" if self.publiczny else "",
                       propozycja="Wskazać minimalną wartość świadczenia (np. 80% wynagrodzenia) albo pozycje, z których Zamawiający może zrezygnować.",
                       pytanie=f"Prosimy o wskazanie minimalnej wartości lub wielkości świadczenia w związku z {self.ak[i0].lok} (art. 433 pkt 4 Pzp)." if self.publiczny else "")

    def sprawdz_inne(self):
        P = self.progi
        for i in self.gdzie(r"potrac\w*"):
            if re.search(r"kar|naleznos|wynagrodzeni", self.N[i]):
                zgoda = re.search(r"za zgod|uznan|bezsporn|prawomocn", self.N[i])
                self.dodaj("potracenie", "Potrącanie kar z wynagrodzenia", "ok" if zgoda else "niskie", self.ak[i],
                           ustalenie="potrącenie kar " + ("ograniczone do uznanych/bezspornych" if zgoda else "bez ograniczeń"),
                           propozycja="" if zgoda else "Potrącenie wyłącznie kar uznanych przez Wykonawcę lub stwierdzonych prawomocnie.")
                break
        for i in self.gdzie(r"ubezpiecz\w*.{0,60}(odpowiedzialnosci cywilnej|oc\b)"):
            t = self.ak[i].tekst
            frag = [m.group(0) for m in re.finditer(r"[^.;:]{0,60}?(\d[\d \u00a0.,]*\s*(zł|PLN|mln)|\d+(,\d+)? ?%)[^.;]{0,30}", t)][:3]
            if frag:
                self.dodaj("oc", "Ubezpieczenie OC", "info", self.ak[i],
                           ustalenie="; ".join(cyt(x.strip(), 110) for x in frag),
                           propozycja="Porównać z polisą OC MiD (suma, podlimity, franszyza); OC projektowe zwykle osobno.")
                break
        kand = self.gdzie(r"okres\w* (rekojmi|gwarancji)\w*.{0,100}(\d{1,3}) ?(lat|miesiec\w*|mies)") or self.gdzie(r"(rekojm|gwarancj)\w*.{0,160}(\d{1,3}) ?(lat|miesiec\w*|mies)")
        for i in kand:
            if "zabezpieczeni" in self.N[i]:
                continue
            m = re.search(r"(rekojm|gwarancj)\w*.{0,160}?(\d{1,3}) ?(lat|miesiec\w*|mies)", self.N[i])
            if not m:
                continue
            mies = int(m.group(2)) * (12 if m.group(3) == "lat" else 1)
            zwiaz = re.search(r"(od dnia|od daty).{0,80}(odbioru robot|zakonczenia robot|swiadectw\w* przejeci|oddania obiektu|pozwolenia na uzytkowanie)", self.N[i])
            ocena = "srednie" if mies > P["rekojmia_max_mies"] or zwiaz else "info"
            self.dodaj("rekojmia", "Rękojmia / gwarancja za dokumentację", ocena, self.ak[i],
                       ustalenie=f"{mies} mies." + ("; liczona od zakończenia robót (okres nieznany z góry)" if zwiaz else ""),
                       propozycja="Rękojmia za dokumentację liczona od odbioru dokumentacji, z datą graniczną (np. 60 mies.)." if ocena != "info" else "")
            break
        for i in self.gdzie(r"(przelew|cesj|przeniesieni\w*)\w*.{0,40}wierzytelnos\w*.{0,80}(zgod|bez zgody)"):
            self.dodaj("cesja", "Cesja wierzytelności", "info", self.ak[i], ustalenie="przelew wierzytelności wymaga zgody",
                       podstawa="" if self.publiczny else "art. 9a ustawy o przeciwdziałaniu nadmiernym opóźnieniom (duży przedsiębiorca wobec MŚP)")
            break
        for i in self.gdzie(r"umow\w* o prace"):
            if "kar" in self.N[i] or "kar" in self.NR[i]:
                self.dodaj("umowa_o_prace", "Wymóg zatrudnienia na umowę o pracę", "info", self.ak[i], ustalenie="kary za niespełnienie wymogu zatrudnienia (art. 95 Pzp)")
                break

    def podsumowanie(self):
        self.sprawdz_kary(); self.sprawdz_limit(); self.sprawdz_odpowiedzialnosc(); self.sprawdz_platnosci()
        self.sprawdz_waloryzacje(); self.sprawdz_zabezpieczenie(); self.sprawdz_prawa_autorskie(); self.sprawdz_ai()
        self.sprawdz_nadzor(); self.sprawdz_terminy_zamawiajacego(); self.sprawdz_organy(); self.sprawdz_odstapienie()
        self.sprawdz_zakres(); self.sprawdz_inne()
        self.f.sort(key=lambda d: (POZIOM[d["ocena"]], d["plik"], d["strona"] or 0))
        return self.f


# ----------------------------------------------------------------------------- ekspozycja na kary
def ekspozycja(an, W=None, dni_lista=(7, 30, 60, 90)):
    kary = [k for k in an.get("kary", []) if k["placi"] == "Wykonawca" and not k.get("podwykonawcy")]
    lim = next((f["dane"].get("limit_proc") for f in an["ustalenia"] if f["id"] == "limit_kar" and f["dane"].get("limit_proc")), None)
    wiersze = []
    for k in kary:
        if k["stawka_proc"] is None or k["jednostka"] != "dzień":
            continue
        r = {"lok": k["lok"], "opis": k["opis"], "przyczyna": k["przyczyna"][:90], "stawka": k["stawka_proc"]}
        for d in dni_lista:
            v = k["stawka_proc"] * d
            r[f"{d} dni"] = min(v, lim) if lim else v
        r["dni_do_limitu"] = (lim / k["stawka_proc"]) if lim else None
        wiersze.append(r)
    return {"limit_proc": lim, "wynagrodzenie": W, "dni": list(dni_lista), "wiersze": wiersze,
            "jednorazowe": [k for k in kary if k["stawka_proc"] is not None and k["jednostka"] != "dzień"]}


def fmt(v, n=1):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:,.{n}f}".replace(",", " ").replace(".", ",")
    return str(v)


# ----------------------------------------------------------------------------- raport
def raport_md(an, pliki):
    L = [f"# Przegląd umowy — {', '.join(Path(p).name for p in pliki)}", ""]
    m = an["meta"]
    L += [f"Perspektywa: **{m['rola']}**; zamawiający: **{'publiczny (Pzp)' if m['publiczny'] else 'prywatny'}**; "
          f"okres umowy: {fmt(m['miesiace'], 0) + ' mies.' if m['miesiace'] else 'nie ustalono'}; zakres z decyzjami: {'tak' if m['decyzje'] else 'nie'}.", ""]
    licz = {k: sum(1 for f in an["ustalenia"] if f["ocena"] == k) for k in OCENY}
    L += [f"Wynik: {licz['wysokie']} 🔴, {licz['srednie']} 🟠, {licz['niskie']} 🟡, {licz['ok']} 🟢, {licz['info']} ℹ️. "
          "Lista kontrolna jest mechaniczna – każdą pozycję sprawdź w tekście (cytat i lokalizacja).", ""]
    L += ["## Ustalenia", "", "| | Temat | Gdzie | Ustalenie | Podstawa | Propozycja |", "|---|---|---|---|---|---|"]
    for f in an["ustalenia"]:
        gdzie = f"{f['lok']}, s. {f['strona']}" if f["strona"] else "—"
        if len(an["meta"]["pliki"]) > 1 and f["plik"]:
            gdzie = f"{f['plik']}: {gdzie}"
        L.append(f"| {OCENY[f['ocena']]} | {f['temat']} | {gdzie} | {f['ustalenie']} | {f['podstawa']} | {f['propozycja']} |")
    if an.get("kary"):
        L += ["", "## Kary umowne (wyciąg)", "", "| Gdzie | Płaci | Wysokość | Za co | Podstawa |", "|---|---|---|---|---|"]
        for k in an["kary"]:
            L.append(f"| {k['lok']}, s. {k['strona']} | {k['placi']} | {k['opis']} | {k['przyczyna'][:140]} | {k['podstawa'] or '—'} |")
    e = an.get("ekspozycja")
    if e and e["wiersze"]:
        L += ["", "## Ekspozycja na kary dzienne (% wynagrodzenia, z limitem łącznym)", "",
              f"Limit łączny: {fmt(e['limit_proc'], 0) + '%' if e['limit_proc'] else 'brak'}" + (f"; wynagrodzenie {fmt(e['wynagrodzenie'], 0)} zł netto" if e["wynagrodzenie"] else ""), "",
              "| Gdzie | Stawka/dzień | " + " | ".join(f"{d} dni" for d in e["dni"]) + " | Dni do limitu |", "|---|---|" + "---|" * (len(e["dni"]) + 1)]
        for r in e["wiersze"]:
            kom = []
            for d in e["dni"]:
                v = r[f"{d} dni"]
                kom.append(f"{fmt(v, 1)}%" + (f" ({fmt(v / 100 * e['wynagrodzenie'], 0)} zł)" if e["wynagrodzenie"] else ""))
            L.append(f"| {r['lok']} | {fmt(r['stawka'], 2)}% | " + " | ".join(kom) + f" | {fmt(r['dni_do_limitu'], 0)} |")
        L += ["", "Uwaga: stawka liczona od podstawy wskazanej w umowie; jeśli podstawą jest wynagrodzenie całkowite, a zwłoka dotyczy części, ekspozycja względem tej części jest wielokrotnie wyższa."]
    if an["meta"].get("zasady"):
        L += ["", "## Zasady MiD (z pliku standardów)", ""] + [f"- {z}" for z in an["meta"]["zasady"]]
    L += ["", "## Podstawy prawne", "",
          "Pzp (t.j. Dz.U. 2026 poz. 793): art. 433 (zakazane postanowienia), 436 (obowiązkowe postanowienia, limit kar), 439 (waloryzacja), "
          "443 (płatności częściowe), 452–453 (zabezpieczenie), 455 (zmiany). Ustawa o przeciwdziałaniu nadmiernym opóźnieniom "
          "w transakcjach handlowych (t.j. Dz.U. 2023 poz. 711): art. 7–9. KC (t.j. Dz.U. 2026 poz. 795): art. 473, 483–484. "
          "Ustawa o prawie autorskim (t.j. Dz.U. 2025 poz. 24): art. 16, 41, 46, 53.",
          "", "Progi ocen (stawki kar, limit, zabezpieczenie) to ustawienia ostrożnościowe narzędzia – do potwierdzenia przez Marcina (plik `--standardy`)."]
    return "\n".join(L)


# ----------------------------------------------------------------------------- komendy
def cmd_struktura(a):
    wyn = []
    for p in a.pliki:
        wyn += [x.d() for x in akapity(p)]
    if a.json:
        Path(a.json).write_text(json.dumps(wyn, ensure_ascii=False, indent=1), encoding="utf-8")
    for x in wyn[: a.limit]:
        print(f"[s.{x['strona']}] {x['lok']:<28} {x['tekst'][:110]}")
    print(f"\n{len(wyn)} akapitów.")


def cmd_analiza(a):
    an = Analiza(a.pliki, a)
    f = an.podsumowanie()
    wynik = {"meta": {"pliki": [str(p) for p in a.pliki], "rola": a.rola, "publiczny": an.publiczny, "miesiace": an.miesiace,
                      "decyzje": an.decyzje, "progi": an.progi, "zasady": an.zasady},
             "ustalenia": f, "kary": getattr(an, "lista_kar", [])}
    wynik["ekspozycja"] = ekspozycja(wynik, a.wynagrodzenie)
    if a.rola == "podwykonawca":
        wynik["przeniesienie"] = [k for k in wynik["kary"] if k["placi"] == "Wykonawca"
                                  and re.search(r"dokument\w* wykonawc|zrid|decyzj\w*|dokumentacj\w*|projektant|wniosk\w* o|kamie\w* milow|projekt\w* (budowlan|wykonawcz|techniczn)", k["_p"] + " " + k["_b"])
                                  and not re.search(r"umow\w* o podwykonawstwo|organizacji ruchu|bhp|bezpieczenstwa pracy", k["_p"])]
    t = raport_md(wynik, a.pliki)
    if a.rola == "podwykonawca" and wynik.get("przeniesienie"):
        t += "\n\n## Do umowy z Generalnym Wykonawcą (flow-down)\n\nKary zamawiającego związane z projektowaniem – GW zwykle przenosi je na projektanta. W umowie MiD–GW: tylko za zwłokę zawinioną przez projektanta, od wynagrodzenia projektanta za daną część, z limitem łącznym i z wyłączeniem czasu przeglądów Inżyniera/Zamawiającego oraz oczekiwania na dane wejściowe.\n\n| Gdzie | Wysokość | Za co |\n|---|---|---|\n"
        t += "\n".join(f"| {k['lok']}, s. {k['strona']} | {k['opis']} | {k['przyczyna'][:160]} |" for k in wynik["przeniesienie"])
    if a.json:
        Path(a.json).write_text(json.dumps(wynik, ensure_ascii=False, indent=1), encoding="utf-8")
    if a.md:
        Path(a.md).write_text(t + "\n", encoding="utf-8")
    print(t)


def cmd_kary(a):
    an = json.loads(Path(a.analiza).read_text(encoding="utf-8"))
    e = ekspozycja(an, a.wynagrodzenie, [int(x) for x in a.dni.split(",")])
    an["ekspozycja"] = e
    t = raport_md(an, an["meta"]["pliki"])
    if "## Ekspozycja" in t:
        print("## Ekspozycja" + t.split("## Ekspozycja")[1].split("## Podstawy")[0].split("## Zasady")[0])
    else:
        print("Brak kar dziennych wyrażonych w % – ekspozycji nie policzono.")
    for k in e["jednorazowe"]:
        v = k["stawka_proc"]
        print(f"- jednorazowa {k['lok']}: {fmt(v, 2)}%" + (f" = {fmt(v / 100 * a.wynagrodzenie, 0)} zł" if a.wynagrodzenie else "") + f" – {k['przyczyna'][:90]}")


def cmd_standardy(a):
    Path(a.out).write_text(json.dumps({"_opis": "Progi ocen przeglądu umów MiD – ustala Marcin; plik trzymać prywatnie (baza_mid)",
                                       "progi": PROGI, "zasady": ["[zasada negocjacyjna MiD, np. limit kar ≤ …% wynagrodzenia]"]},
                                      ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Zapisano {a.out}")


def cmd_komentarze(a):
    import docx
    an = json.loads(Path(a.analiza).read_text(encoding="utf-8"))
    d = docx.Document(a.docx)
    if not hasattr(d, "add_comment"):
        sys.exit("Wymagany python-docx >= 1.2 (pip install -U python-docx).")
    pars = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                pars += c.paragraphs
    ntxt = [norm(p.text) for p in pars]
    prog = POZIOM[a.poziom]
    dodane, brak = 0, []
    nazwa = Path(a.docx).name
    for f in an["ustalenia"]:
        if POZIOM[f["ocena"]] > prog or not f["cytat"] or (f["plik"] and Path(f["plik"]).stem != Path(nazwa).stem and f["plik"] != nazwa):
            continue
        c = norm(f["cytat"]).replace("…", "")
        c = re.sub(r"^(§ ?\d+|\d+[.)]|\(?[a-z]{1,3}\)|[a-z]\.|[ivx]+\.|subklauzula [\d.]+)\s*", "", c)
        cel = None
        for dl in (70, 45, 28):
            frag = c[:dl].strip()
            if len(frag) < 15:
                continue
            for i, t in enumerate(ntxt):
                if frag and frag in t:
                    cel = i
                    break
            if cel is not None:
                break
        if cel is None:
            brak.append(f"{f['temat']} ({f['lok']})")
            continue
        p = pars[cel]
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            continue
        tekst = f"[{OCENY_TXT[f['ocena']]}] {f['temat']}: {f['ustalenie']}"
        if f["podstawa"]:
            tekst += f"\nPodstawa: {f['podstawa']}"
        if f["propozycja"]:
            tekst += f"\nPropozycja MiD: {f['propozycja']}"
        d.add_comment(runs, text=tekst, author=a.autor, initials="MiD")
        dodane += 1
    d.save(a.out)
    print(f"Zapisano {a.out}: {dodane} komentarzy.")
    if brak:
        print("Nie znaleziono akapitu dla: " + "; ".join(brak))


def cmd_pytania(a):
    an = json.loads(Path(a.analiza).read_text(encoding="utf-8"))
    f = [x for x in an["ustalenia"] if x["ocena"] in ("wysokie", "srednie")]
    L = []
    if a.tryb == "pzp":
        L += ["# Wnioski o wyjaśnienie treści SWZ – projektowane postanowienia umowy", "",
              "Termin na pytania: art. 135 ust. 2 / art. 284 ust. 2 Pzp (skill analiza-swz, `terminy`).", ""]
        n = 0
        ids = {x["id"] for x in f}
        for x in f:
            if x["id"] == "limit_odpowiedzialnosci" and "odszkodowanie" in ids:
                continue
            q = x["pytanie"] if x["pytanie"] and not x["pytanie"].startswith("Czy Zamawiający dopuści zmianę zapisu") else (
                f"Wykonawca wnosi o zmianę {x['lok']} zgodnie z propozycją: {x['propozycja']} Czy Zamawiający wyrazi zgodę na taką zmianę?"
                if x["propozycja"] and x["lok"] != "—" else (
                    f"Wykonawca wnosi o uzupełnienie wzoru umowy: {x['propozycja']} Czy Zamawiający wyrazi na to zgodę?" if x["propozycja"] else x["pytanie"]))
            if not q:
                continue
            n += 1
            gdzie = f"{x['lok']}{', s. ' + str(x['strona']) if x['strona'] else ''}" if x["lok"] != "—" else "wzór umowy – brak postanowienia"
            L += [f"**Pytanie {n}** (dotyczy: {gdzie}{'; ' + x['plik'] if len(an['meta']['pliki']) > 1 and x['plik'] else ''}):", "", q, ""]
    else:
        L += ["# Propozycje zmian do umowy – Pracownia Projektowa MiD", "", "| Lp. | Zapis | Obecnie | Propozycja MiD | Uzasadnienie |", "|---|---|---|---|---|"]
        for i, x in enumerate([x for x in f if x["propozycja"]], 1):
            L.append(f"| {i} | {x['lok']} | {x['cytat'][:180]} | {x['propozycja']} | {x['ustalenie']}{'; ' + x['podstawa'] if x['podstawa'] else ''} |")
    t = "\n".join(L)
    if a.md:
        Path(a.md).write_text(t + "\n", encoding="utf-8")
    if a.docx:
        import docx
        from docx.shared import Pt
        d = docx.Document()
        st = d.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(10.5)
        for l in L:
            if l.startswith("# "):
                d.add_heading(l[2:], level=1)
            elif l.startswith("|") and not l.startswith("|---"):
                cells = [c.strip() for c in l.strip("|").split("|")]
                if not hasattr(cmd_pytania, "_t") or cmd_pytania._t is None or len(cmd_pytania._t.columns) != len(cells):
                    cmd_pytania._t = d.add_table(rows=0, cols=len(cells))
                    cmd_pytania._t.style = "Table Grid"
                row = cmd_pytania._t.add_row().cells
                for c, v in zip(row, cells):
                    c.text = v
            elif l.startswith("**"):
                cmd_pytania._t = None
                p = d.add_paragraph()
                p.add_run(l.strip("*").replace("**", "")).bold = True
            elif l:
                cmd_pytania._t = None
                d.add_paragraph(l)
        d.save(a.docx)
    print(t)


def main():
    import signal
    if hasattr(signal, "SIGPIPE"):
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("struktura"); s.add_argument("pliki", nargs="+"); s.add_argument("--json"); s.add_argument("--limit", type=int, default=60)
    s = sub.add_parser("analiza"); s.add_argument("pliki", nargs="+"); s.add_argument("--wynagrodzenie", type=float)
    s.add_argument("--miesiace", type=float); s.add_argument("--zamawiajacy", choices=["publiczny", "prywatny", "duzy"])
    s.add_argument("--rola", default="wykonawca", choices=["wykonawca", "podwykonawca", "zlecajacy"]); s.add_argument("--standardy")
    s.add_argument("--md"); s.add_argument("--json")
    s = sub.add_parser("kary"); s.add_argument("analiza"); s.add_argument("--wynagrodzenie", type=float); s.add_argument("--dni", default="7,30,60,90")
    s = sub.add_parser("standardy"); s.add_argument("--out", default="standardy_umow.json")
    s = sub.add_parser("komentarze"); s.add_argument("docx"); s.add_argument("--analiza", required=True); s.add_argument("--out", required=True)
    s.add_argument("--poziom", default="niskie", choices=["wysokie", "srednie", "niskie", "info"]); s.add_argument("--autor", default="Pracownia Projektowa MiD")
    s = sub.add_parser("pytania"); s.add_argument("analiza"); s.add_argument("--tryb", default="pzp", choices=["pzp", "negocjacje"])
    s.add_argument("--md"); s.add_argument("--docx")
    a = ap.parse_args()
    {"struktura": cmd_struktura, "analiza": cmd_analiza, "kary": cmd_kary, "komentarze": cmd_komentarze, "pytania": cmd_pytania,
     "standardy": cmd_standardy}[a.cmd](a)


if __name__ == "__main__":
    main()
